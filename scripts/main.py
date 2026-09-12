import sys, os
sys.path.insert(0, os.path.abspath("."))
# -*- coding: utf-8 -*-
import openpyxl
import os, time
from scripts.styles import get_workbook_styles
from scripts.infra_sheets import (
    build_vol1_cover,
    build_rates_master,
    build_global_factors,
    build_labour_productivity,
    build_sundries_reference
)
from scripts.trade_builder import build_standard_trade
from scripts.trade_builder_carr import build_carriage_trade
from scripts.trade_builder_earth import build_earthwork_trade
from scripts.trade_configs import get_all_trade_configs
from scripts.scope_inputs import merge_scope_metadata
from scripts.cross_volume import build_resolved_cross_volume
from scripts.paths import WB_VOL1_FILE, WB_VOL1_LATEST_FILE, WB_VOL1_TEMPLATE_FILE

def generate_full_workbook():
    t0 = time.time()
    print("================================================================")
    print("CPWD DAR 2019 VOLUME 1 CUSTOM RATE ANALYSIS WORKBOOK GENERATOR")
    print("Strict MS Excel 2016 Compatibility & excel-estimator-design Standard")
    print("================================================================")
    
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    
    styles = get_workbook_styles()
    
    # 1. Cover sheet (first tab)
    print("--- Generating Cover Sheet ---")
    build_vol1_cover(wb, styles)

    # 2. Infrastructure Sheets
    print("--- Generating 5 Shared Infrastructure Sheets ---")
    build_rates_master(wb, styles)
    build_global_factors(wb, styles)
    build_labour_productivity(wb, styles)
    build_sundries_reference(wb, styles)
    build_resolved_cross_volume(wb, styles)
    
    # Remove default sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)
        
    # 3. Trade Builder Sheets (01 to 12)
    print("--- Generating 12 Dedicated Trade Builder Sheets ---")
    configs = merge_scope_metadata(get_all_trade_configs())
    
    # 01 Carriage
    build_carriage_trade(wb, configs['01_Carriage_of_Materials'], styles)
    
    # 02 Earth Work (dedicated analytical simulator & first-principles engine)
    build_earthwork_trade(wb, configs['02_Earth_Work'], styles)
    
    # 03 to 12 Standard Trade Builders
    trade_keys = [
        '03_Mortars',
        '04_Concrete_Work',
        '05_RCC_Work',
        '06_Masonry_Work',
        '07_Stone_Work',
        '08_Cladding_Work',
        '09_Wood_and_PVC_Work',
        '10_Steel_Work',
        '11_Flooring',
        '12_Roofing'
    ]
    
    for key in trade_keys:
        build_standard_trade(wb, configs[key], styles)
        
    primary_out = WB_VOL1_FILE
    print(f"Saving complete workbook to {primary_out}...")
    wb.save(primary_out)
    
    import shutil
    for copy_target in [
        WB_VOL1_LATEST_FILE,
        WB_VOL1_TEMPLATE_FILE
    ]:
        print(f"Synchronizing byte-identical copy to {copy_target}...")
        shutil.copyfile(primary_out, copy_target)

    elapsed = time.time() - t0
    file_size_mb = os.path.getsize(primary_out) / (1024 * 1024)
    print(f"SUCCESS! Generated {len(wb.sheetnames)} sheets in {elapsed:.2f}s ({file_size_mb:.2f} MB)")
    print(f"Sheet names: {wb.sheetnames}")
    print("================================================================")

if __name__ == '__main__':
    generate_full_workbook()
