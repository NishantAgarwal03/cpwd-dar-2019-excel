import openpyxl, re, json

wb = openpyxl.load_workbook('CivilDAR_2019_Vol_1_Converted.xlsx', data_only=True)
sheets = [
    '01_Carriage_of_Materials', '02_Earth_Work', '03_Mortars', '04_Concrete_Work',
    '05_RCC_Work', '06_Masonry_Work', '07_Stone_Work', '08_Cladding_Work',
    '09_Wood_and_PVC_Work', '10_Steel_Work', '11_Flooring', '12_Roofing'
]

productivity_records = []
sundries_records = []

for sh_name in sheets:
    s = wb[sh_name]
    current_item_no = ''
    current_item_desc = ''
    current_unit_basis = ''
    
    for r in range(1, s.max_row + 1):
        col2 = str(s.cell(r, 2).value or '').strip()
        col3 = str(s.cell(r, 3).value or '').strip()
        col4 = str(s.cell(r, 4).value or '').strip()
        col5 = str(s.cell(r, 5).value or '').strip()
        col6 = str(s.cell(r, 6).value or '').strip()
        col7 = str(s.cell(r, 7).value or '').strip()
        
        # Check if item header
        # Pattern like 1.1, 1.1.1, 4.1.2, 6.1.1 etc
        if re.match(r'^\d+\.\d+(\.\d+)?$', col2):
            current_item_no = col2
            if col3:
                current_item_desc = col3
                # look for basis
                m_basis = re.search(r'Details of cost for (.*?)(\.|\s+MATERIAL|\s+LABOUR|$)', col3, re.IGNORECASE)
                if m_basis:
                    current_unit_basis = m_basis.group(1).strip()
                else:
                    current_unit_basis = '1 unit'
        elif 'details of cost for' in col3.lower():
            m_basis = re.search(r'Details of cost for (.*?)(\.|\s+MATERIAL|\s+LABOUR|$)', col3, re.IGNORECASE)
            if m_basis:
                current_unit_basis = m_basis.group(1).strip()
        
        # Check if Labour or Machinery
        if re.match(r'^(01\d{2}|00\d{2})$', col2):
            try:
                qty_val = float(col5.replace(',', ''))
            except:
                qty_val = 0.0
            try:
                rate_val = float(col6.replace(',', ''))
            except:
                rate_val = 0.0
                
            productivity_records.append({
                'subhead': sh_name,
                'item_no': current_item_no,
                'item_desc': current_item_desc[:100],
                'basis': current_unit_basis,
                'code': col2,
                'description': col3,
                'unit': col4,
                'coefficient': qty_val,
                'rate': rate_val,
                'type': 'Labour' if col2.startswith('01') else 'Machinery'
            })
            
        # Check if Sundries
        if col2 == '9999' or 'sundries' in col3.lower():
            try:
                qty_val = float(col5.replace(',', ''))
            except:
                qty_val = 0.0
            try:
                amt_val = float(col7.replace(',', ''))
            except:
                amt_val = 0.0
                
            sundries_records.append({
                'subhead': sh_name,
                'item_no': current_item_no,
                'item_desc': current_item_desc[:100],
                'basis': current_unit_basis,
                'description': col3 if col3 else 'Sundries',
                'base_ls': qty_val,
                'multiplier': 2.0,
                'amount': amt_val
            })

print(f'Mined {len(productivity_records)} productivity records.')
print(f'Mined {len(sundries_records)} sundries records.')

with open('labour_productivity.json', 'w', encoding='utf-8') as f:
    json.dump(productivity_records, f, indent=2)

with open('sundries_reference.json', 'w', encoding='utf-8') as f:
    json.dump(sundries_records, f, indent=2)

print('Saved reference data successfully.')
