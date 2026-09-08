import openpyxl, pypdf, re, json

print("Starting extraction...")

# 1. Parse PDF rates for all codes 0001-2399
reader = pypdf.PdfReader('CivilDAR_2019_Vol_1.pdf')
pdf_rates = {}
current_code = None

for p in range(11, 73):
    text = reader.pages[p].extract_text()
    for line in text.split('\n'):
        line = line.strip()
        m = re.match(r'^(\d{4})\s+(.*?)([\d,]+\.\d{2})$', line)
        if m:
            code = m.group(1)
            try:
                pdf_rates[code] = float(m.group(3).replace(',', ''))
            except:
                pass
            current_code = None
        else:
            m_start = re.match(r'^(\d{4})\s+(.*)$', line)
            if m_start:
                current_code = m_start.group(1)
            elif current_code is not None:
                m2 = re.search(r'([\d,]+\.\d{2})$', line)
                if m2:
                    try:
                        pdf_rates[current_code] = float(m2.group(1).replace(',', ''))
                    except:
                        pass
                    current_code = None

print(f"Extracted {len(pdf_rates)} rates from PDF.")

# 2. Parse Excel 00_Basic_Rates
wb = openpyxl.load_workbook('CivilDAR_2019_Vol_1_Converted.xlsx', data_only=True)
s_rates = wb['00_Basic_Rates']

items = []
current = None

for r in range(134, s_rates.max_row + 1):
    code = str(s_rates.cell(r, 2).value or '').strip()
    desc = str(s_rates.cell(r, 3).value or '').strip()
    unit = str(s_rates.cell(r, 4).value or '').strip()
    rate = str(s_rates.cell(r, 6).value or '').strip()
    
    if re.match(r'^\d{4}$', code):
        if current:
            items.append(current)
        r_val = None
        if rate not in ['', 'None']:
            try:
                r_val = float(rate.replace(',', ''))
            except:
                pass
        current = {'code': code, 'desc': desc, 'unit': unit, 'rate': r_val}
    elif current:
        if desc and not code and rate in ['', 'None']:
            current['desc'] += ' ' + desc
        elif rate not in ['', 'None'] and current['rate'] is None:
            try:
                current['rate'] = float(rate.replace(',', ''))
            except:
                pass
        if unit and not current['unit']:
            current['unit'] = unit

if current:
    items.append(current)

# Fill gaps using PDF
filled_count = 0
for it in items:
    if it['rate'] is None and it['code'] in pdf_rates:
        it['rate'] = pdf_rates[it['code']]
        filled_count += 1

print(f"Total basic rates from Excel: {len(items)}, filled from PDF: {filled_count}")

# Categorize items
for it in items:
    code_int = int(it['code'])
    if code_int < 100:
        it['category'] = 'Hire Charges of Plants & Machinery'
    elif code_int < 200:
        it['category'] = 'Labour Wages'
    elif 2200 <= code_int <= 2399:
        it['category'] = 'Carriage of Materials'
    else:
        it['category'] = 'Building Materials'

# Save cleaned rates
with open('rates_master_clean.json', 'w', encoding='utf-8') as f:
    json.dump(items, f, indent=2)

print("Saved rates_master_clean.json successfully.")
