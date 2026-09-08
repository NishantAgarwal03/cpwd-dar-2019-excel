# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Protection
from openpyxl.worksheet.datavalidation import DataValidation

from scripts.trade_layout import (
    R_TITLE, R_LEGEND, R_GUIDE,
    R_P1_HEAD, R_P1_FIRST, R_P1_LAST,
    R_META, R_NOMEN,
    R_P2_FIRST, R_P2_LAST, R_KEY,
    R_AUDIT,
    R_MAT_HEAD, R_MAT_COLS, R_MAT_FIRST, R_MAT_LAST, R_MAT_SUB, R_A_TOTAL,
    R_LAB_HEAD, R_LAB_COLS, R_LAB_FIRST, R_LAB_LAST, R_LAB_SUB,
    R_SUN_HEAD, R_SUN,
    R_MU_HEAD, R_MU_COLS, R_W, R_X1, R_X, R_Y1, R_Y, R_Z1, R_Z, R_Z2,
    R_COST, R_RATE, R_SAY,
    R_SAY_NOTE,
    R_LIB_HEAD, R_LIB_NOTE, R_LIB_COLS, R_LIB_FIRST, R_LIB_LAST,
    LAST_COL, SAY_RULE_NOTE,
    section_bar, col_headers, border_row, build_guide,
    build_scope_panel, build_meta_and_nomenclature, build_driver_panel,
)

TBL_HEADERS = ['Line', 'Code / Source', 'Description / Specification', 'Unit', 'Quantity / Coeff',
               'Basic Rate (Rs)', 'Amount (Rs)', 'Source Reference / Remarks', 'Markup Tag']

MU_HEADERS = ['Step', 'Component / Markup Description', 'Apply? (YES/NO)', 'Basis Applied',
              'Base Amount (Rs)', 'Factor / %', 'Amount (Rs)', 'CPWD Statutory Rule & Guidance Note',
              'A-Rule']

LIB_HEADERS = ['Item Code', 'Item Nomenclature / Specification', 'Unit', 'Output Basis',
               'Direct Cost W (Rs)', 'Markups Applied', 'Unit Rate (Rs)', 'Say Rate (Rs)', 'A-Total (Rs)']


def add_trade_header_and_legend(ws, config, styles):
    ws.merge_cells(start_row=R_TITLE, start_column=1, end_row=R_TITLE, end_column=LAST_COL)
    c_title = ws.cell(row=R_TITLE, column=1)
    c_title.value = config['trade_title']
    c_title.font = styles['font_title']
    c_title.fill = styles['fill_title']
    c_title.alignment = styles['align_center']
    ws.row_dimensions[R_TITLE].height = 30

    ws.cell(row=R_LEGEND, column=1, value='LEGEND:').font = styles['font_legend']
    ws.cell(row=R_LEGEND, column=1).alignment = styles['align_center']
    legend_items = [
        (2, 'INPUT - you type it', styles['fill_input'], styles['font_bold']),
        (3, 'OVERRIDE - optional', styles['fill_override'], styles['font_bold']),
        (4, 'LOOKUP - fetched', styles['fill_lookup'], styles['font_regular']),
        (5, 'DERIVED - calculated', styles['fill_calc'], styles['font_regular']),
        (6, 'RESULT - subtotal', styles['fill_result'], styles['font_bold']),
        (7, 'SAY - final rate', styles['fill_say'], styles['font_bold']),
        (8, 'CPWD rule / note', styles['fill_note'], styles['font_note']),
        (9, 'Not editable', styles['fill_subtotal'], styles['font_regular']),
    ]
    for col, text, fill, font in legend_items:
        c = ws.cell(row=R_LEGEND, column=col, value=text)
        c.fill = fill
        c.font = font
        c.border = styles['border_thin']
        c.alignment = styles['align_center']
    ws.row_dimensions[R_LEGEND].height = 20

    ws.merge_cells(start_row=R_GUIDE, start_column=1, end_row=R_GUIDE, end_column=LAST_COL)
    c_guide = ws.cell(row=R_GUIDE, column=1)
    c_guide.value = config['trade_guidance']
    c_guide.font = styles['font_note']
    c_guide.fill = styles['fill_note']
    c_guide.alignment = styles['align_wrap']
    c_guide.border = styles['border_thin']
    ws.row_dimensions[R_GUIDE].height = 30


def _resource_row(ws, r, idx, spec, styles, kind, dv_codes):
    """Emit one MATERIAL or LABOUR line."""
    fallback = 'Custom / Intermediate Material' if kind == 'material' else 'Custom Labour / Plant'

    ws.cell(row=r, column=1, value=idx).alignment = styles['align_center']

    c_code = ws.cell(row=r, column=2, value=(spec['code'] if spec else ''))
    c_code.alignment = styles['align_center']
    c_code.font = styles['font_bold']
    c_code.fill = styles['fill_input']
    if dv_codes is not None:
        dv_codes.add(c_code)

    c_desc = ws.cell(row=r, column=3)
    if spec and spec.get('custom_desc'):
        c_desc.value = spec['custom_desc']
    else:
        c_desc.value = (f'=IF(B{r}="","", IFERROR(INDEX(Rates_Master!$C:$C, '
                        f'MATCH(B{r}, Rates_Master!$A:$A, 0)), "{fallback}"))')
    c_desc.alignment = styles['align_left']
    c_desc.fill = styles['fill_lookup']

    c_unit = ws.cell(row=r, column=4)
    if spec and spec.get('custom_unit'):
        c_unit.value = spec['custom_unit']
    else:
        c_unit.value = (f'=IF(B{r}="","", IFERROR(INDEX(Rates_Master!$D:$D, '
                        f'MATCH(B{r}, Rates_Master!$A:$A, 0)), ""))')
    c_unit.alignment = styles['align_center']
    c_unit.fill = styles['fill_lookup']

    c_qty = ws.cell(row=r, column=5, value=(spec['coeff'] if spec else None))
    c_qty.alignment = styles['align_right']
    c_qty.number_format = styles['fmt_qty']
    c_qty.fill = styles['fill_input']

    c_rate = ws.cell(row=r, column=6)
    if spec and spec.get('custom_rate_formula'):
        c_rate.value = spec['custom_rate_formula']
    elif spec and spec.get('custom_rate') is not None:
        c_rate.value = spec['custom_rate']
    else:
        c_rate.value = (f'=IF(B{r}="","", IFERROR(INDEX(Rates_Master!$E:$E, '
                        f'MATCH(B{r}, Rates_Master!$A:$A, 0)), 0))')
    c_rate.alignment = styles['align_right']
    c_rate.number_format = styles['fmt_currency']
    c_rate.fill = styles['fill_lookup']

    c_amt = ws.cell(row=r, column=7, value=f'=IF(OR(B{r}="", E{r}=""), 0, ROUND(E{r} * F{r}, 2))')
    c_amt.alignment = styles['align_right']
    c_amt.number_format = styles['fmt_currency']
    c_amt.fill = styles['fill_calc']

    c_rem = ws.cell(row=r, column=8, value=(spec['note'] if spec else ''))
    c_rem.alignment = styles['align_wrap']
    c_rem.font = styles['font_note']

    c_tag = ws.cell(row=r, column=9, value=('A' if (spec and spec.get('a_tag')) else ''))
    c_tag.alignment = styles['align_center']
    c_tag.font = styles['font_bold']
    c_tag.fill = styles['fill_say'] if (spec and spec.get('a_tag')) else styles['fill_input']

    border_row(ws, r, styles)
    ws.row_dimensions[r].height = 20


def build_standard_trade(wb, config, styles):
    ws = wb.create_sheet(title=config['sheet_name'])
    ws.views.sheetView[0].showGridLines = True

    add_trade_header_and_legend(ws, config, styles)
    build_guide(ws, styles)

    dv_yesno = DataValidation(type='list', formula1='"YES,NO"', allow_blank=False)
    ws.add_data_validation(dv_yesno)
    dv_atag = DataValidation(type='list', formula1='"A,"', allow_blank=True,
                             prompt='Enter A only if this line imports a rate that already includes '
                                    'Water, GST, CPOH and Cess. A-tagged lines are excluded from every '
                                    'markup base per the CPWD (W-A) convention.',
                             promptTitle='CPWD A-tag')
    ws.add_data_validation(dv_atag)
    dv_codes = DataValidation(type='list', formula1='=Master_Codes', allow_blank=True,
                              showErrorMessage=False,
                              prompt='Pick a Rates_Master code, or type an intermediate source such as '
                                     '03_Mortars or a Resolved_Cross_Volume_Items item number.',
                              promptTitle='Component picker')
    ws.add_data_validation(dv_codes)

    # --- Panels ---------------------------------------------------------
    dv_registry = []
    build_scope_panel(ws, config, styles, dv_registry)
    build_meta_and_nomenclature(ws, config, styles)
    build_driver_panel(ws, config, styles, dv_registry)

    for options, cell in dv_registry:
        formula = '"' + ','.join(options) + '"'
        if len(formula) <= 255:
            dv = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv)
            dv.add(cell)

    # --- Audit bar ------------------------------------------------------
    ws.cell(row=R_AUDIT, column=1, value='AUDIT STATUS:').font = styles['font_white_bold']
    ws.cell(row=R_AUDIT, column=1).fill = styles['fill_header']
    ws.cell(row=R_AUDIT, column=1).alignment = styles['align_center']

    b = ws.cell(row=R_AUDIT, column=2,
                value=f'=IF(AND(D{R_AUDIT}="OK", F{R_AUDIT}="OK", H{R_AUDIT}="OK"), '
                      f'"[PASS] ALL CHECKS OK", "[ALERT] CHECKS FAILED")')
    b.font = styles['font_result']
    b.fill = styles['fill_result']
    b.alignment = styles['align_center']

    ws.cell(row=R_AUDIT, column=3, value='Output Qty Check:').font = styles['font_note']
    ws.cell(row=R_AUDIT, column=3).alignment = styles['align_right']
    ws.cell(row=R_AUDIT, column=4, value=f'=IF(D{R_META}>0, "OK", "ERR: Qty<=0")').alignment = styles['align_center']

    ws.cell(row=R_AUDIT, column=5, value='Direct Cost Check:').font = styles['font_note']
    ws.cell(row=R_AUDIT, column=5).alignment = styles['align_right']
    ws.cell(row=R_AUDIT, column=6, value=f'=IF(G{R_W}>0, "OK", "ERR: W<=0")').alignment = styles['align_center']

    ws.cell(row=R_AUDIT, column=7, value='Markup Flow Check:').font = styles['font_note']
    ws.cell(row=R_AUDIT, column=7).alignment = styles['align_right']
    ws.cell(row=R_AUDIT, column=8,
            value=f'=IF(AND(G{R_W}<=G{R_X}, G{R_X}<=G{R_Y}, G{R_Y}<=G{R_Z}, G{R_A_TOTAL}<=G{R_W}), '
                  f'"OK", "ERR: Markups Broken")').alignment = styles['align_center']

    ws.cell(row=R_AUDIT, column=9, value=f'=IF(G{R_A_TOTAL}>0, "A-rule active", "No A imports")')
    ws.cell(row=R_AUDIT, column=9).alignment = styles['align_center']
    ws.cell(row=R_AUDIT, column=9).font = styles['font_note']

    for c in (4, 6, 8):
        ws.cell(row=R_AUDIT, column=c).font = styles['font_bold']
    border_row(ws, R_AUDIT, styles)
    ws.row_dimensions[R_AUDIT].height = 22

    # --- Section 3: materials -------------------------------------------
    section_bar(ws, R_MAT_HEAD,
                config.get('material_section_title',
                           '3. MATERIAL COMPONENT BUILD-UP  (tag column I with "A" for any line whose rate '
                           'already includes statutory markups)'),
                styles)
    col_headers(ws, R_MAT_COLS, TBL_HEADERS, styles)

    materials = config.get('default_materials', [])
    for i in range(R_MAT_LAST - R_MAT_FIRST + 1):
        r = R_MAT_FIRST + i
        spec = materials[i] if i < len(materials) else None
        _resource_row(ws, r, i + 1, spec, styles, 'material', dv_codes)
        dv_atag.add(ws.cell(row=r, column=9))

    ws.merge_cells(start_row=R_MAT_SUB, start_column=1, end_row=R_MAT_SUB, end_column=6)
    ws.cell(row=R_MAT_SUB, column=1, value='Subtotal Materials (Rs):').font = styles['font_bold']
    ws.cell(row=R_MAT_SUB, column=1).alignment = styles['align_right']
    cm = ws.cell(row=R_MAT_SUB, column=7, value=f'=SUM(G{R_MAT_FIRST}:G{R_MAT_LAST})')
    cm.font = styles['font_bold']
    cm.alignment = styles['align_right']
    cm.number_format = styles['fmt_currency']
    for c in range(1, LAST_COL + 1):
        ws.cell(row=R_MAT_SUB, column=c).fill = styles['fill_subtotal']
        ws.cell(row=R_MAT_SUB, column=c).border = styles['border_double_bottom']
    ws.row_dimensions[R_MAT_SUB].height = 22

    ws.merge_cells(start_row=R_A_TOTAL, start_column=1, end_row=R_A_TOTAL, end_column=6)
    ws.cell(row=R_A_TOTAL, column=1,
            value='of which "A" - imported lines already carrying statutory markups:').font = styles['font_bold']
    ws.cell(row=R_A_TOTAL, column=1).alignment = styles['align_right']
    ca = ws.cell(row=R_A_TOTAL, column=7,
                 value=f'=ROUND(SUMIF($I${R_MAT_FIRST}:$I${R_MAT_LAST}, "A", '
                       f'$G${R_MAT_FIRST}:$G${R_MAT_LAST}), 2)')
    ca.font = styles['font_bold']
    ca.alignment = styles['align_right']
    ca.number_format = styles['fmt_currency']
    ws.cell(row=R_A_TOTAL, column=8,
            value='CPWD (W-A) convention: this amount stays inside W and inside every subtotal, but is '
                  'subtracted from each markup BASE so an already-marked-up import is not marked up twice. '
                  'Verified against DAR 2019 items 10.1, 18.78 and 18.79.').font = styles['font_note']
    ws.cell(row=R_A_TOTAL, column=8).alignment = styles['align_wrap']
    for c in range(1, LAST_COL + 1):
        ws.cell(row=R_A_TOTAL, column=c).fill = styles['fill_subtotal']
    border_row(ws, R_A_TOTAL, styles)
    ws.row_dimensions[R_A_TOTAL].height = 30

    # --- Section 4: labour & machinery ----------------------------------
    section_bar(ws, R_LAB_HEAD,
                '4. LABOUR & MACHINERY COMPONENT BUILD-UP  (crew suggestions: see the work-type summary on '
                'Labour_Machinery_Productivity. Plant codes 0001-0083 already include operator, fuel and '
                'lubricants for an 8-hour shift per CPWD Note 1)', styles)
    col_headers(ws, R_LAB_COLS, TBL_HEADERS, styles)

    labour = config.get('default_labour', [])
    for i in range(R_LAB_LAST - R_LAB_FIRST + 1):
        r = R_LAB_FIRST + i
        spec = labour[i] if i < len(labour) else None
        _resource_row(ws, r, i + 1, spec, styles, 'labour', dv_codes)
        ws.cell(row=r, column=9).fill = styles['fill_calc']

    ws.merge_cells(start_row=R_LAB_SUB, start_column=1, end_row=R_LAB_SUB, end_column=6)
    ws.cell(row=R_LAB_SUB, column=1, value='Subtotal Labour & Plant Hire (Rs):').font = styles['font_bold']
    ws.cell(row=R_LAB_SUB, column=1).alignment = styles['align_right']
    cl = ws.cell(row=R_LAB_SUB, column=7, value=f'=SUM(G{R_LAB_FIRST}:G{R_LAB_LAST})')
    cl.font = styles['font_bold']
    cl.alignment = styles['align_right']
    cl.number_format = styles['fmt_currency']
    for c in range(1, LAST_COL + 1):
        ws.cell(row=R_LAB_SUB, column=c).fill = styles['fill_subtotal']
        ws.cell(row=R_LAB_SUB, column=c).border = styles['border_double_bottom']
    ws.row_dimensions[R_LAB_SUB].height = 22

    # --- Section 5: sundries --------------------------------------------
    section_bar(ws, R_SUN_HEAD, '5. SUNDRIES & LUMP-SUM ALLOWANCES  '
                                '(manual entry - the DAR gives no reproducible formula; use Sundries_Reference '
                                'to find a comparable item)', styles)
    ws.cell(row=R_SUN, column=1, value='1').alignment = styles['align_center']
    ws.cell(row=R_SUN, column=2, value='9999').alignment = styles['align_center']
    ws.cell(row=R_SUN, column=2).font = styles['font_bold']
    ws.cell(row=R_SUN, column=3, value='Sundries (guided by Sundries_Reference sheet)')
    ws.cell(row=R_SUN, column=4, value='L.S.').alignment = styles['align_center']
    cs = ws.cell(row=R_SUN, column=5, value=config.get('default_sundries_base', 0.0))
    cs.alignment = styles['align_right']
    cs.number_format = '0.00'
    cs.fill = styles['fill_input']
    cf = ws.cell(row=R_SUN, column=6, value='=Factor_Sundries')
    cf.alignment = styles['align_right']
    cf.number_format = '0.00'
    cf.fill = styles['fill_lookup']
    cg = ws.cell(row=R_SUN, column=7, value=f'=ROUND(E{R_SUN} * F{R_SUN}, 2)')
    cg.alignment = styles['align_right']
    cg.number_format = styles['fmt_currency']
    cg.font = styles['font_bold']
    ws.cell(row=R_SUN, column=8,
            value='Base L.S. allowance x Cost Index Multiplier (2.00) from Global_Factors').font = styles['font_note']
    border_row(ws, R_SUN, styles)
    ws.row_dimensions[R_SUN].height = 22

    # --- Section 6: markup chain ----------------------------------------
    section_bar(ws, R_MU_HEAD,
                '6. STATUTORY MARKUPS & FINAL RATE DERIVATION (CPWD DAR METHOD, WITH THE (W-A) EXCLUSION)',
                styles)
    col_headers(ws, R_MU_COLS, MU_HEADERS, styles)

    toggles = config.get('toggles', {'water': 'YES', 'gst': 'YES', 'cpoh': 'YES', 'cess': 'YES'})
    notes = config.get('toggle_notes', {})

    rows = [
        (R_W, 'W', 'Base Direct Production Cost (Materials + Labour + Sundries)', None, 'Direct Sum',
         None, None, f'=G{R_MAT_SUB} + G{R_LAB_SUB} + G{R_SUN}',
         'Total direct cost of production (W), A-tagged imports included, exactly as the book totals it.',
         'includes A'),
        (R_X1, 'X1', 'Add Water Charges', toggles['water'], 'On (W - A)',
         f'=G{R_W} - G{R_A_TOTAL}', '=Factor_Water',
         f'=IF(C{R_X1}="YES", ROUND(E{R_X1} * F{R_X1}, 2), 0)',
         notes.get('water', 'CPWD Standard: 1% for curing & site water. Toggle NO if dry item or water supplied free.'),
         'net of A'),
        (R_X, 'X', 'Subtotal "X" (W + Water Charges)', None, 'W + Water',
         None, None, f'=G{R_W} + G{R_X1}', 'Compounded base for GST calculation', 'includes A'),
        (R_Y1, 'Y1', 'Add GST on Works Contract', toggles['gst'], 'On (X - A)',
         f'=G{R_X} - G{R_A_TOTAL}', '=Factor_GST',
         f'=IF(C{R_Y1}="YES", ROUND(E{R_Y1} * F{R_Y1}, 2), 0)',
         notes.get('gst', 'CPWD DAR 2019 factor 0.1405 (works contract). Toggle NO if tax exempt.'),
         'net of A'),
        (R_Y, 'Y', 'Subtotal "Y" (X + GST)', None, 'X + GST',
         None, None, f'=G{R_X} + G{R_Y1}', 'Compounded base for Contractor Profit & Overheads', 'includes A'),
        (R_Z1, 'Z1', 'Add Contractor Profit & Overheads (15% CPOH)', toggles['cpoh'], 'On (Y - A)',
         f'=G{R_Y} - G{R_A_TOTAL}', '=Factor_CPOH',
         f'=IF(C{R_Z1}="YES", ROUND(E{R_Z1} * F{R_Z1}, 2), 0)',
         notes.get('cpoh', 'Standard CPWD 15% allowance for site overheads, head office costs & margin.'),
         'net of A'),
        (R_Z, 'Z', 'Subtotal "Z" (Y + CPOH)', None, 'Y + CPOH',
         None, None, f'=G{R_Y} + G{R_Z1}', 'Compounded base for Labour Welfare Cess', 'includes A'),
        (R_Z2, 'Z2', 'Add Labour Welfare Cess (1% BOCW Cess)', toggles['cess'], 'On (Z - A)',
         f'=G{R_Z} - G{R_A_TOTAL}', '=Factor_Cess',
         f'=IF(C{R_Z2}="YES", ROUND(E{R_Z2} * F{R_Z2}, 2), 0)',
         notes.get('cess', 'Statutory 1% Building and Other Construction Workers Welfare Cess.'),
         'net of A'),
        (R_COST, 'Cost', f'="Total Cost of Batch (" & TEXT(D{R_META}, "0.00") & " " & F{R_META} & "):"',
         None, 'Z + Cess', None, None, f'=G{R_Z} + G{R_Z2}',
         'Total evaluated cost for the batch size in the Panel 1 header row', '-'),
        (R_RATE, 'Rate', f'="Analyzed Unit Rate per 1.00 " & F{R_META} & ":"',
         None, 'Cost / Batch Qty', None, None, f'=ROUND(G{R_COST} / D{R_META}, 2)',
         'Calculated cost per standard unit of measurement', '-'),
        (R_SAY, 'SAY', f'="OFFICIAL SAY RATE (Rs per " & F{R_META} & "):"',
         None, 'MROUND to Rs 0.05', None, None, f'=MROUND(G{R_RATE}, 0.05)',
         'OFFICIAL CPWD ROUNDED RATE FOR TENDER SCHEDULES & ESTIMATES. CPWD quotes Say rates to '
         'the nearest 5 paise - see the note directly below this row.', '-'),
    ]

    for r, step, desc, toggle, basis, base_f, factor_f, amount_f, note, arule in rows:
        ws.cell(row=r, column=1, value=step).font = styles['font_bold']
        ws.cell(row=r, column=1).alignment = styles['align_center']
        ws.cell(row=r, column=2, value=desc).alignment = styles['align_left']
        ws.cell(row=r, column=3, value=(toggle if toggle else '-')).alignment = styles['align_center']
        ws.cell(row=r, column=4, value=basis).alignment = styles['align_center']

        cb = ws.cell(row=r, column=5, value=(base_f if base_f else '-'))
        cb.alignment = styles['align_right']
        if base_f:
            cb.number_format = styles['fmt_currency']

        cfa = ws.cell(row=r, column=6, value=(factor_f if factor_f else '-'))
        cfa.alignment = styles['align_right']
        if factor_f:
            cfa.number_format = styles['fmt_percent']

        cv = ws.cell(row=r, column=7, value=amount_f)
        cv.alignment = styles['align_right']
        cv.number_format = styles['fmt_currency']
        cv.font = styles['font_bold']

        cn = ws.cell(row=r, column=8, value=note)
        cn.alignment = styles['align_wrap']
        cn.font = styles['font_note']

        ct = ws.cell(row=r, column=9, value=arule)
        ct.alignment = styles['align_center']
        ct.font = styles['font_note']

        border_row(ws, r, styles)

        if r in (R_W, R_X, R_Y, R_Z):
            for c in range(1, LAST_COL + 1):
                ws.cell(row=r, column=c).fill = styles['fill_subtotal']
        elif r in (R_X1, R_Y1, R_Z1, R_Z2):
            ws.cell(row=r, column=3).fill = styles['fill_input']
            ws.cell(row=r, column=3).font = styles['font_bold']
            dv_yesno.add(ws.cell(row=r, column=3))
        elif r == R_COST:
            for c in range(1, LAST_COL + 1):
                ws.cell(row=r, column=c).fill = styles['fill_result']
        elif r == R_RATE:
            for c in range(1, LAST_COL + 1):
                ws.cell(row=r, column=c).fill = styles['fill_subtotal']
        elif r == R_SAY:
            for c in range(1, LAST_COL + 1):
                ws.cell(row=r, column=c).fill = styles['fill_say']
            ws.cell(row=r, column=2).font = styles['font_say']
            ws.cell(row=r, column=7).font = styles['font_say']

        ws.row_dimensions[r].height = 30 if r == R_SAY else 24

    # Why the Say rate is rounded to 5 paise - stated on the sheet itself.
    ws.merge_cells(start_row=R_SAY_NOTE, start_column=1, end_row=R_SAY_NOTE, end_column=LAST_COL)
    sn = ws.cell(row=R_SAY_NOTE, column=1)
    sn.value = SAY_RULE_NOTE
    sn.font = styles['font_note']
    sn.fill = styles['fill_note']
    sn.alignment = styles['align_wrap']
    border_row(ws, R_SAY_NOTE, styles)
    ws.row_dimensions[R_SAY_NOTE].height = 44

    # --- Section 7: running library --------------------------------------
    section_bar(ws, R_LIB_HEAD, '7. RUNNING CUSTOM NON-DSR ITEMS LIBRARY FOR THIS TRADE', styles)
    ws.merge_cells(start_row=R_LIB_NOTE, start_column=1, end_row=R_LIB_NOTE, end_column=LAST_COL)
    ci = ws.cell(row=R_LIB_NOTE, column=1)
    ci.value = ('Log all completed custom rate analyses for this trade below. Items logged here maintain an '
                'internal project audit trail and can be referenced from other builder sheets.')
    ci.font = styles['font_note']
    ci.fill = styles['fill_note']
    ci.alignment = styles['align_wrap']
    ws.row_dimensions[R_LIB_NOTE].height = 20

    col_headers(ws, R_LIB_COLS, LIB_HEADERS, styles)

    sample_lib = config.get('sample_library', [])
    for i in range(R_LIB_LAST - R_LIB_FIRST + 1):
        r = R_LIB_FIRST + i
        item = sample_lib[i] if i < len(sample_lib) else None
        ws.cell(row=r, column=1,
                value=(item['code'] if item else f"C-{config['sheet_name'][:2]}.{i+1:02d}")).alignment = styles['align_center']
        ws.cell(row=r, column=2,
                value=(item['desc'] if item else '(Available for new custom item)')).alignment = styles['align_left']
        ws.cell(row=r, column=3,
                value=(item['unit'] if item else config['default_basis_unit'])).alignment = styles['align_center']
        ws.cell(row=r, column=4,
                value=(item['basis'] if item else config['default_basis_qty'])).alignment = styles['align_center']
        cw = ws.cell(row=r, column=5, value=(item['w'] if item else None))
        cw.alignment = styles['align_right']
        cw.number_format = styles['fmt_currency']
        ws.cell(row=r, column=6,
                value=(item['markups'] if item else 'Full W->X->Y->Z')).alignment = styles['align_center']
        cr = ws.cell(row=r, column=7, value=(item['rate'] if item else None))
        cr.alignment = styles['align_right']
        cr.number_format = styles['fmt_currency']
        cy = ws.cell(row=r, column=8, value=(item['say'] if item else None))
        cy.alignment = styles['align_right']
        cy.number_format = styles['fmt_currency']
        cy.font = styles['font_bold']
        cat = ws.cell(row=r, column=9, value=(item.get('a_total') if item else None))
        cat.alignment = styles['align_right']
        cat.number_format = styles['fmt_currency']

        row_fill = styles['fill_subtotal'] if r % 2 == 0 else styles['fill_calc']
        for c in range(1, LAST_COL + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = row_fill
            cell.border = styles['border_thin']
            if c not in (1, 8):
                cell.font = styles['font_regular']
        ws.row_dimensions[r].height = 20

    # --- Column widths ---------------------------------------------------
    for col, w in {'A': 34, 'B': 26, 'C': 46, 'D': 16, 'E': 18,
                   'F': 18, 'G': 20, 'H': 44, 'I': 14}.items():
        ws.column_dimensions[col].width = w

    # --- Protection ------------------------------------------------------
    for r in range(1, ws.max_row + 1):
        for c in range(1, LAST_COL + 1):
            ws.cell(row=r, column=c).protection = Protection(locked=True)

    unlocked = [f'B{R_META}', f'D{R_META}', f'F{R_META}', f'I{R_NOMEN}', f'E{R_SUN}',
                f'C{R_X1}', f'C{R_Y1}', f'C{R_Z1}', f'C{R_Z2}']
    for ref in unlocked:
        ws[ref].protection = Protection(locked=False)
    for r in range(R_P1_FIRST, R_P1_LAST + 1):
        ws.cell(row=r, column=2).protection = Protection(locked=False)
    for r in range(R_P2_FIRST, R_P2_LAST + 1):
        ws.cell(row=r, column=2).protection = Protection(locked=False)
    for r in range(R_MAT_FIRST, R_MAT_LAST + 1):
        for c in (2, 5, 8, 9):
            ws.cell(row=r, column=c).protection = Protection(locked=False)
    for r in range(R_LAB_FIRST, R_LAB_LAST + 1):
        for c in (2, 5, 8):
            ws.cell(row=r, column=c).protection = Protection(locked=False)
    for r in range(R_LIB_FIRST, R_LIB_LAST + 1):
        for c in range(1, LAST_COL + 1):
            ws.cell(row=r, column=c).protection = Protection(locked=False)

    ws.protection.sheet = True
    ws.freeze_panes = 'A4'
    print(f"Built standard trade sheet: {config['sheet_name']}")
