"""CPWD Automated Item and Gang Parser.

Ingests CPWD rate analysis data across sub-heads into:
1. Gang Composition Registry (pure physical resource mix)
2. Published Productivity Registry (transparent productivity derivations)

Starting with complete 02_Earth_Work (84 items) and providing generic sheet parsing
for 03_Mortars, 04_Concrete_Work, 05_RCC_Work, 06_Masonry_Work, etc.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import openpyxl

from cpwd_gang_registry import GANG_REGISTRY, GangResource, ItemGang
from cpwd_productivity_registry import PRODUCTIVITY_REGISTRY, ProductivityEntry, SHIFT_HOURS
import support_builder_earth_v2


_GANG_EFFORT_MAP: dict[tuple[str, str], tuple[float, float]] = {
    # (item_code, resource_code) -> (gang_size, task_hours)
    ("2.1.1", "0114"): (4.0, 13.60),
    ("2.1.1", "0115"): (4.0, 11.20),
    ("2.1.2", "0114"): (4.0, 17.20),
    ("2.1.2", "0115"): (4.0, 14.20),
    ("2.2.1", "0114"): (4.0, 11.80),
    ("2.2.1", "0115"): (3.0, 9.60),
    ("2.2.1", "0101"): (1.0, 3.20),
    ("2.2.1", "0113"): (1.0, 0.064),
    ("2.2", "0114"): (4.0, 11.80),
    ("2.2", "0115"): (3.0, 9.60),
    ("2.2", "0101"): (1.0, 3.20),
    ("2.2", "0113"): (1.0, 0.064),
    ("2.3.1", "0114"): (2.0, 8.80),
    ("2.3.1", "0115"): (3.0, 9.60),
    ("2.3.1", "0101"): (1.0, 3.20),
    ("2.3.1", "0113"): (1.0, 0.064),
    ("2.4", "0113"): (1.0, 0.064),
    ("2.5", "0101"): (1.0, 3.20),
    ("2.25", "0114"): (1.0, 6.40),
    ("2.25", "0115"): (1.0, 6.80),
    ("2.25", "0101"): (1.0, 1.60),
    ("2.32", "0114"): (2.0, 3.20),
    ("2.32", "0115"): (2.0, 2.68),
}

_MACHINERY_HOURS_MAP: dict[tuple[str, str], float] = {
    # (item_code, resource_code) -> machine_hours
    ("2.2", "0003"): 0.064,
    ("2.2", "EQUIP-0.5T"): 8.80,
    ("2.2.1", "0003"): 0.064,
    ("2.2.1", "EQUIP-0.5T"): 8.80,
    ("2.3.1", "0003"): 0.064,
    ("2.3.1", "EQUIP-0.5T"): 8.80,
    ("2.4", "0003"): -0.064,
    ("2.25", "EQUIP-DURMAT"): 6.40,
    ("2.6.1", "0020"): 0.328,
    ("2.6.1", "0018"): 0.328,
    ("2.6.2", "0020"): 0.504,
    ("2.6.2", "0018"): 0.504,
    ("2.7.1", "0020"): 0.504,
    ("2.7.1", "0017"): 0.504,
    ("2.7.2", "0020"): 1.000,
    ("2.7.2", "0017"): 1.000,
    ("2.7.3", "0020"): 1.000,
    ("2.7.3", "0017"): 1.000,
}


def ingest_earthwork_items() -> int:
    """Ingest all 84 Earth Work items from support_builder_earth_v2 into the registries."""
    count = 0
    for raw in support_builder_earth_v2.ITEMS:
        item_code = raw["id"]
        desc = raw.get("desc", "")
        unit = raw.get("unit", "cum")
        base_qty = float(raw.get("base_qty", 10.0))
        dsr_rate = raw.get("dsr_rate")
        notes = raw.get("notes", "")

        machinery_list: list[GangResource] = []
        labour_list: list[GangResource] = []
        material_list: list[GangResource] = []
        sundries_list: list[GangResource] = []

        has_0003_roller = False

        for section_name, rows in raw.get("sections", []):
            for r in rows:
                code = str(r[0]).strip()
                r_desc = str(r[1]).strip()
                r_unit = str(r[2]).strip()
                coeff = float(r[3])
                rate = float(r[4])

                # Determine machine hours or labour gang hours
                hours = _MACHINERY_HOURS_MAP.get((item_code, code))
                if hours is None and section_name == "MACHINERY":
                    hours = round(coeff * SHIFT_HOURS, 3)

                gang_size, task_hours = _GANG_EFFORT_MAP.get((item_code, code), (None, None))
                if task_hours is None and section_name == "LABOUR":
                    task_hours = round(coeff * SHIFT_HOURS, 2)

                res = GangResource(
                    code=code,
                    name=r_desc,
                    category="Machinery" if section_name == "MACHINERY" else ("Labour" if section_name == "LABOUR" else ("Material" if section_name == "MATERIAL" else "Sundries")),
                    coefficient=coeff,
                    unit=r_unit,
                    rate=rate,
                    gang_count=gang_size,
                    task_hours=hours if section_name == "MACHINERY" else task_hours,
                )

                if section_name == "MACHINERY":
                    machinery_list.append(res)
                    if code == "0003":
                        has_0003_roller = True
                elif section_name == "LABOUR":
                    labour_list.append(res)
                elif section_name == "MATERIAL":
                    if code == "9999" or "sundries" in r_desc.lower():
                        res.category = "Sundries"
                        sundries_list.append(res)
                    else:
                        material_list.append(res)
                else:
                    sundries_list.append(res)

        # For banking/rough excavation items that roll with 8-t roller and 1/2-t roller
        if has_0003_roller and item_code in ("2.2.1", "2.2.2", "2.3.1", "2.3.2"):
            machinery_list.append(
                GangResource(
                    code="EQUIP-0.5T",
                    name="½-tonne hand roller / wooden or steel rammers (durmats)",
                    category="Machinery",
                    coefficient=1.100,
                    unit="day",
                    rate=0.00,
                    task_hours=8.80,
                )
            )

        # For trench/plinth filling item 2.25
        if item_code == "2.25":
            machinery_list.append(
                GangResource(
                    code="EQUIP-DURMAT",
                    name="Wooden or steel rammers (durmats) for layer compaction",
                    category="Machinery",
                    coefficient=0.800,
                    unit="day",
                    rate=0.00,
                    task_hours=6.40,
                )
            )

        item_gang = ItemGang(
            item_code=item_code,
            description=desc,
            batch_quantity=base_qty,
            batch_unit=unit,
            dsr_rate=dsr_rate,
            machinery=machinery_list,
            labour=labour_list,
            materials=material_list,
            sundries=sundries_list,
            notes=notes or "CPWD DAR 2019 Specification",
        )

        GANG_REGISTRY._items[item_code] = item_gang

        # Register in Productivity Registry as well
        for res in [*machinery_list, *labour_list]:
            if res.coefficient != 0:
                raw_bench = {
                    "subhead": "02_Earth_Work",
                    "item_code": item_code,
                    "activity": desc,
                    "resource_code": res.code,
                    "resource_name": res.name,
                    "resource_type": res.category,
                    "coefficient": res.coefficient,
                    "unit": res.unit,
                    "batch_quantity": base_qty,
                    "batch_unit": unit,
                    "machine_hours": res.task_hours if res.category == "Machinery" else None,
                    "task_hours": res.task_hours if res.category == "Labour" else None,
                    "gang_size": res.gang_count,
                    "source_citation": f"CPWD DAR 2019 Item {item_code}",
                }
                entry = PRODUCTIVITY_REGISTRY._build_entry(raw_bench)
                key = (entry.subhead, entry.item_code, entry.resource_code)
                PRODUCTIVITY_REGISTRY._entries[key] = entry

        count += 1

    return count


TRADE_SHEETS: tuple[tuple[str, str], ...] = (
    ("03_Mortars", "03_Mortars"),
    ("04_Concrete_Work", "04_Concrete_Work"),
    ("05_RCC_Work", "05_RCC_Work"),
    ("06_Masonry_Work", "06_Masonry_Work"),
    ("07_Stone_Work", "07_Stone_Work"),
    ("08_Cladding_Work", "08_Cladding_Work"),
    ("09_Wood_and_PVC_Work", "09_Wood_and_PVC_Work"),
    ("10_Steel_Work", "10_Steel_Work"),
    ("11_Flooring", "11_Flooring"),
    ("12_Roofing", "12_Roofing"),
)


def _parse_basis_string(text: str) -> tuple[float | None, str | None]:
    if not text:
        return None, None
    m = re.search(r"([0-9.]+)\s*([a-zA-Z]+)", text)
    if m:
        try:
            return float(m.group(1)), m.group(2).lower()
        except ValueError:
            pass
    word_map = {"one": 1.0, "each": 1.0, "ten": 10.0, "two": 2.0, "three": 3.0}
    parts = text.strip().split()
    if parts and parts[0].lower() in word_map:
        unit = parts[1].lower() if len(parts) > 1 else "unit"
        return word_map[parts[0].lower()], unit
    return None, None


def ingest_all_trades_items(
    converted_xlsx_path: str | Path = "data/converted_xlsx/CivilDAR_2019_Vol_1_Converted.xlsx",
    rates_master_json_path: str | Path = "data/reference_json/rates_master_clean.json",
    labour_prod_json_path: str | Path = "data/reference_json/labour_productivity.json",
) -> tuple[int, int]:
    """Ingest complete CPWD items from Sub-head 02 through 12 into GANG_REGISTRY and PRODUCTIVITY_REGISTRY.

    Returns:
        (total_items_ingested, total_productivity_entries_registered)
    """
    # 1. Start with Sub-head 02 Earth Work (ground truth pilot & learning baseline)
    ingest_earthwork_items()

    # 2. Load Rates Master reference for authoritative categories and basic rates
    rates_master: dict[str, dict[str, Any]] = {}
    rm_path = Path(rates_master_json_path)
    if rm_path.exists():
        with open(rm_path, "r", encoding="utf-8", errors="replace") as f:
            for rec in json.load(f):
                code = str(rec.get("code") or "").strip()
                if code:
                    rates_master[code] = rec

    # 3. Load labour_productivity.json basis cache for fallback batch sizes
    lp_basis_cache: dict[tuple[str, str], str] = {}
    lp_path = Path(labour_prod_json_path)
    if lp_path.exists():
        with open(lp_path, "r", encoding="utf-8", errors="replace") as f:
            for rec in json.load(f):
                sh = rec.get("subhead") or ""
                it = str(rec.get("item_no") or "").strip()
                b = rec.get("basis") or ""
                if sh and it and b:
                    lp_basis_cache[(sh, it)] = b

    # 4. Open converted workbook for trades 03 to 12
    wb_path = Path(converted_xlsx_path)
    if not wb_path.exists():
        return len(GANG_REGISTRY._items), len(PRODUCTIVITY_REGISTRY._entries)

    wb = openpyxl.load_workbook(wb_path, read_only=True)
    total_trade_items = 0

    for sheet_name, subhead in TRADE_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        current_item: dict[str, Any] | None = None

        def _commit_item(item_data: dict[str, Any]) -> None:
            nonlocal total_trade_items
            item_code = item_data["item_code"]
            desc = item_data["description"]
            batch_qty = item_data["batch_quantity"]
            batch_unit = item_data["batch_unit"]

            machinery_list: list[GangResource] = []
            labour_list: list[GangResource] = []
            material_list: list[GangResource] = []
            sundries_list: list[GangResource] = []

            for r in item_data["resources"]:
                code = r["code"]
                name = r["name"]
                cat = r["category"]
                coeff = r["coefficient"]
                unit = r["unit"]
                rate = r["rate"]

                if cat == "Machinery":
                    task_hours = round(abs(coeff) * SHIFT_HOURS, 4)
                    machinery_list.append(
                        GangResource(
                            code=code,
                            name=name,
                            category="Machinery",
                            coefficient=coeff,
                            unit=unit,
                            rate=rate,
                            task_hours=task_hours,
                        )
                    )
                elif cat == "Labour":
                    task_hours = round(abs(coeff) * SHIFT_HOURS, 4)
                    labour_list.append(
                        GangResource(
                            code=code,
                            name=name,
                            category="Labour",
                            coefficient=coeff,
                            unit=unit,
                            rate=rate,
                            gang_count=1.0,
                            task_hours=task_hours,
                        )
                    )
                elif cat == "Sundries":
                    sundries_list.append(
                        GangResource(
                            code=code,
                            name=name,
                            category="Sundries",
                            coefficient=coeff,
                            unit=unit,
                            rate=rate,
                        )
                    )
                else:
                    material_list.append(
                        GangResource(
                            code=code,
                            name=name,
                            category="Material",
                            coefficient=coeff,
                            unit=unit,
                            rate=rate,
                        )
                    )

            item_gang = ItemGang(
                item_code=item_code,
                description=desc,
                batch_quantity=batch_qty,
                batch_unit=batch_unit,
                dsr_rate=None,
                machinery=machinery_list,
                labour=labour_list,
                materials=material_list,
                sundries=sundries_list,
                notes=f"CPWD DAR 2019 {subhead} specification",
            )
            GANG_REGISTRY._items[item_code] = item_gang

            for res in [*machinery_list, *labour_list]:
                if res.coefficient != 0:
                    raw_bench = {
                        "subhead": subhead,
                        "item_code": item_code,
                        "activity": desc,
                        "resource_code": res.code,
                        "resource_name": res.name,
                        "resource_type": res.category,
                        "coefficient": res.coefficient,
                        "unit": res.unit,
                        "batch_quantity": batch_qty,
                        "batch_unit": batch_unit,
                        "machine_hours": res.task_hours if res.category == "Machinery" else None,
                        "task_hours": res.task_hours if res.category == "Labour" else None,
                        "gang_size": res.gang_count,
                        "source_citation": f"CPWD DAR 2019 {subhead} Item {item_code}",
                    }
                    entry = PRODUCTIVITY_REGISTRY._build_entry(raw_bench)
                    key = (entry.subhead, entry.item_code, entry.resource_code)
                    PRODUCTIVITY_REGISTRY._entries[key] = entry

            total_trade_items += 1

        for row in ws.iter_rows(values_only=True):
            page, code_raw, desc_raw, unit_raw, qty_raw, rate_raw, amt_raw = row[:7]
            code = str(code_raw or "").strip()
            desc = str(desc_raw or "").strip()

            if "." in code and code[0].isdigit() and qty_raw is None and rate_raw is None:
                if current_item and current_item["resources"]:
                    _commit_item(current_item)

                b_qty, b_unit = _parse_basis_string(desc)
                if not b_qty:
                    lp_b = lp_basis_cache.get((subhead, code))
                    b_qty, b_unit = _parse_basis_string(lp_b or "")
                if not b_qty:
                    b_qty, b_unit = 1.0, "unit"

                current_item = {
                    "subhead": subhead,
                    "item_code": code,
                    "description": desc,
                    "batch_quantity": b_qty,
                    "batch_unit": b_unit,
                    "resources": [],
                }
                continue

            if current_item:
                if code == "TOTAL" or "Add " in desc or "Cost for" in desc:
                    continue
                if code and len(code) == 4 and code.isdigit() and qty_raw is not None:
                    try:
                        qty = float(qty_raw)
                        rm_entry = rates_master.get(code)
                        if rm_entry:
                            cat_raw = rm_entry.get("category", "")
                            if "Plant" in cat_raw or "Machinery" in cat_raw:
                                cat = "Machinery"
                            elif "Labour" in cat_raw:
                                cat = "Labour"
                            elif code == "9999":
                                cat = "Sundries"
                            else:
                                cat = "Material"
                            rate = float(rm_entry.get("rate", rate_raw or 0.0))
                            name = rm_entry.get("desc", desc)
                        else:
                            if code.startswith("00"):
                                cat = "Machinery"
                            elif code.startswith("01"):
                                cat = "Labour"
                            elif code == "9999":
                                cat = "Sundries"
                            else:
                                cat = "Material"
                            rate = float(rate_raw) if rate_raw is not None else 0.0
                            name = desc

                        current_item["resources"].append({
                            "code": code,
                            "name": name,
                            "category": cat,
                            "coefficient": qty,
                            "unit": unit_raw or "unit",
                            "rate": rate,
                        })
                    except (ValueError, TypeError):
                        pass

        if current_item and current_item["resources"]:
            _commit_item(current_item)

    wb.close()
    return len(GANG_REGISTRY._items), len(PRODUCTIVITY_REGISTRY._entries)


def update_registry_exports() -> tuple[Path, Path, Path, Path]:
    """Re-export base registries to INI and TXT files."""
    g_txt = GANG_REGISTRY.export_to_txt("data/gang_registry.txt")
    g_ini = GANG_REGISTRY.export_to_ini("data/gang_registry.ini")
    p_txt = PRODUCTIVITY_REGISTRY.export_to_txt("data/productivity_registry.txt")
    p_ini = PRODUCTIVITY_REGISTRY.export_to_ini("data/productivity_registry.ini")
    return g_txt, g_ini, p_txt, p_ini


def export_all_trades_registries(
    output_dir: str | Path = "data",
) -> tuple[Path, Path, Path, Path]:
    """Export all trades (Sub-heads 02 to 12) to dedicated INI and TXT registry files."""
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    g_txt = GANG_REGISTRY.export_to_txt(out / "gang_registry_all_trades.txt")
    g_ini = GANG_REGISTRY.export_to_ini(out / "gang_registry_all_trades.ini")
    p_txt = PRODUCTIVITY_REGISTRY.export_to_txt(out / "productivity_registry_all_trades.txt")
    p_ini = PRODUCTIVITY_REGISTRY.export_to_ini(out / "productivity_registry_all_trades.ini")
    return g_txt, g_ini, p_txt, p_ini


if __name__ == "__main__":
    import sys
    print("Ingesting all trades items (Sub-heads 02 to 12)...")
    total_items, total_prod = ingest_all_trades_items()
    print(f"Successfully ingested {total_items} items across all trades with {total_prod} productivity entries.")

    g_txt, g_ini, p_txt, p_ini = update_registry_exports()
    print(f"Base exports updated:\n - {g_txt}\n - {g_ini}\n - {p_txt}\n - {p_ini}")

    at_gtxt, at_gini, at_ptxt, at_pini = export_all_trades_registries()
    print(f"All-trades exports generated:\n - {at_gtxt}\n - {at_gini}\n - {at_ptxt}\n - {at_pini}")


