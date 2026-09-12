# -*- coding: utf-8 -*-
"""
generate_audit_workbook.py
Builds CPWD_DAR_2019_Audit_Test_Cases.xlsx
One sheet per sub-head section, listing every item the audit engine
processed with its PDF Say rate, shadow-calculated rate, and status.
"""
import sys, os, json
sys.path.insert(0, os.path.abspath('.'))

import openpyxl
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter

from scripts.audit_parser import parse_all
from scripts.shadow_calc   import shadow_calculate
from scripts.paths import VOL1_CONVERTED_XLSX, VOL2_CONVERTED_XLSX, RATES_MASTER_JSON, AUDIT_TEST_CASES_XLSX

# ── Files ────────────────────────────────────────────────────────────
VOL1_CONV   = VOL1_CONVERTED_XLSX
VOL2_CONV   = VOL2_CONVERTED_XLSX
RATES_JSON  = RATES_MASTER_JSON
OUT_FILE    = AUDIT_TEST_CASES_XLSX

# ── Sheet name mapping (parser key → display name, max 31 chars) ─────
SHEET_MAP = {
    '01_Carriage_of_Materials':    'Carriage of Materials',
    '02_Earth_Work':               'Earth Work',
    '03_Mortars':                  'Mortars',
    '04_Concrete_Work':            'Concrete Work',
    '05_RCC_Work':                 'Reinforced Cement Concrete',
    '06_Masonry_Work':             'Masonry Work',
    '07_Stone_Work':               'Stone Work',
    '08_Cladding_Work':            'Cladding Work',
    '09_Wood_and_PVC_Work':        'Wood and PVC Work',
    '10_Steel_Work':               'Steel Work',
    '11_Flooring':                 'Flooring',
    '12_Roofing':                  'Roofing',
    '13_Finishing':                'Finishing',
    '14_Repairs_to_Buildings':     'Repairs to Buildings',
    '15_Dismantling_Demolishing':  'Dismantling and Demolishing',
    '16_Road_Work':                'Road Work',
    '17_Sanitary_Installations':   'Sanitary Installations',
    '18_Water_Supply':             'Water Supply',
    '19_Drainage':                 'Drainage',
    '20_Pile_Work':                'Pile Work',
    '21_Aluminium_Work':           'Aluminium Work',
    '22_Water_Proofing':           'Water Proofing',
    '23_Rain_Water_Harvesting':    'Rain Water Harvesting',     # 21 chars
    '24_Heritage_Buildings':       'Conservation Heritage Bldgs', # 27 chars
    '25_Structural_Glazing':       'Structural Glazing ACP',
    '26_New_Technologies':         'New Technologies & Materials', # 28 chars
}

# ── Colours ──────────────────────────────────────────────────────────
C_TITLE_BG   = '1F3864'   # dark navy
C_TITLE_FG   = 'FFFFFF'
C_HDR_BG     = '2E75B6'   # CPWD blue
C_HDR_FG     = 'FFFFFF'
C_PASS       = 'E2EFDA'   # light green
C_FAIL       = 'FCE4D6'   # light salmon
C_SKIP       = 'FFF2CC'   # light yellow
C_MISMATCH   = 'FDEBD0'   # light orange
C_MISSING    = 'E8D5FA'   # light purple
C_ALT        = 'F2F2F2'   # alternating row grey
C_NOTE       = 'FFC000'   # amber for notes

STATUS_FILL = {
    'PASS':         PatternFill('solid', fgColor=C_PASS),
    'FAIL':         PatternFill('solid', fgColor=C_FAIL),
    'SKIP':         PatternFill('solid', fgColor=C_SKIP),
    'RATE_MISMATCH':PatternFill('solid', fgColor=C_MISMATCH),
    'MISSING_CODE': PatternFill('solid', fgColor=C_MISSING),
    'ERROR':        PatternFill('solid', fgColor='FF0000'),
}
STATUS_LABEL = {
    'PASS':         'PASS',
    'FAIL':         'FAIL  (basis qty)',
    'SKIP':         'SKIP  (composite)',
    'RATE_MISMATCH':'RATE MISMATCH',
    'MISSING_CODE': 'MISSING CODE',
    'ERROR':        'ERROR',
}

def _thin():
    s = Side(style='thin', color='BFBFBF')
    return Border(left=s, right=s, top=s, bottom=s)

def _fmt_rate(v):
    return round(v, 2) if v is not None else ''

def build_legend(ws, row):
    """Write a small colour-coded legend block."""
    ws.cell(row=row, column=1, value='Legend:').font = Font(bold=True, size=9)
    items = [
        ('PASS',          C_PASS,    'Shadow calc = PDF rate (within ±0.06)'),
        ('FAIL',          C_FAIL,    'Basis qty parsing issue — math is correct, divisor wrong'),
        ('SKIP',          C_SKIP,    'Composite item references other items; cannot shadow-calc'),
        ('RATE MISMATCH', C_MISMATCH,'Code found in Rates_Master but rate differs from PDF'),
        ('MISSING CODE',  C_MISSING, 'Code used in PDF analysis is absent from Rates_Master'),
    ]
    for i, (label, colour, note) in enumerate(items):
        c1 = ws.cell(row=row+i+1, column=1, value=label)
        c1.fill      = PatternFill('solid', fgColor=colour)
        c1.font      = Font(size=9, bold=True)
        c1.alignment = Alignment(horizontal='center', vertical='center')
        c1.border    = _thin()
        c2 = ws.cell(row=row+i+1, column=2, value=note)
        c2.font      = Font(size=9)
        c2.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.row_dimensions[row+i+1].height = 18
    return row + len(items) + 2


def build_sheet(ws, section_key, display_name, results, vol):
    """Populate one worksheet with the audit results for this section."""
    ws.views.sheetView[0].showGridLines = False

    # ── Title row ────────────────────────────────────────────────────
    ws.merge_cells('A1:I1')
    t = ws['A1']
    t.value     = f'CPWD DAR 2019 — {display_name}  |  Rate Audit Test Cases'
    t.font      = Font(name='Calibri', bold=True, size=12, color=C_TITLE_FG)
    t.fill      = PatternFill('solid', fgColor=C_TITLE_BG)
    t.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    # ── Sub-title ─────────────────────────────────────────────────────
    ws.merge_cells('A2:I2')
    s = ws['A2']
    s.value     = (f'Source: CivilDAR_2019_Vol_{vol}_Converted.xlsx / {section_key}   '
                   f'|   Rates from: rates_master_clean.json   '
                   f'|   Markup chain: W + 1% Water + 14.05% GST + 15% CPOH + 1% Cess')
    s.font      = Font(name='Calibri', italic=True, size=9, color='595959')
    s.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[2].height = 16

    ws.row_dimensions[3].height = 6  # spacer

    # ── Column headers ────────────────────────────────────────────────
    HEADERS = [
        ('Item\nCode',       9),
        ('Item Description', 52),
        ('Basis\nQty',       7),
        ('Basis\nUnit',      7),
        ('Markup\nChain',    8),
        ('PDF\nSay Rate',    11),
        ('Calc\nSay Rate',   11),
        ('Diff\n(Calc-PDF)', 11),
        ('Audit Status',     18),
    ]
    for col, (hdr, width) in enumerate(HEADERS, 1):
        c = ws.cell(row=4, column=col, value=hdr)
        c.font      = Font(name='Calibri', bold=True, size=10, color=C_HDR_FG)
        c.fill      = PatternFill('solid', fgColor=C_HDR_BG)
        c.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)
        c.border    = _thin()
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[4].height = 30

    # ── Data rows ─────────────────────────────────────────────────────
    for idx, r in enumerate(results):
        row = 5 + idx
        alt  = PatternFill('solid', fgColor=C_ALT) if idx % 2 else None
        sfill = STATUS_FILL.get(r.status, alt)

        pdf_say  = _fmt_rate(r.pdf_say)
        calc_say = _fmt_rate(r.shadow_say)
        diff     = round(r.shadow_say - r.pdf_say, 2) if (
                     r.shadow_say is not None and r.pdf_say is not None
                     and r.status not in ('SKIP',)) else ''

        chain_label = 'W+Water+GST+CPOH+Cess' if r.has_markup else 'Sum only'

        notes = r.skip_reason or ', '.join(r.missing_codes) or ', '.join(r.rate_mismatches)

        row_data = [
            r.item_code,
            r.description,
            r.basis_qty  if r.basis_qty != 1.0 else '',
            r.basis_unit if r.basis_unit != 'unit' else '',
            chain_label,
            pdf_say,
            calc_say if r.status != 'SKIP' else '',
            diff,
            STATUS_LABEL.get(r.status, r.status),
        ]

        for col, val in enumerate(row_data, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font      = Font(name='Calibri', size=9)
            c.alignment = Alignment(
                horizontal='center' if col not in (2,) else 'left',
                vertical='center', wrap_text=(col == 2)
            )
            c.border = _thin()
            # Apply status colour to status cell; alt fill to others
            if col == 9:
                c.fill = sfill
                c.font = Font(name='Calibri', size=9, bold=True)
            elif alt:
                c.fill = alt

        # Notes below if there are missing codes or mismatch details
        if notes and r.status in ('MISSING_CODE', 'RATE_MISMATCH'):
            ws.cell(row=row, column=2).value += f'  [!] {notes}'

        ws.row_dimensions[row].height = 16 if len(str(r.description)) < 80 else 28

    # Freeze header
    ws.freeze_panes = 'A5'

    # ── Summary block below data ──────────────────────────────────────
    last = 5 + len(results) + 1
    total   = len(results)
    by_stat = {}
    for r in results:
        by_stat[r.status] = by_stat.get(r.status, 0) + 1

    audited = total - by_stat.get('SKIP', 0)
    passed  = by_stat.get('PASS', 0)

    ws.merge_cells(f'A{last}:B{last}')
    ws[f'A{last}'].value = f'Total items: {total}   |   Audited: {audited}   |   PASS: {passed} ({100*passed/audited:.0f}%)' if audited else f'Total: {total}'
    ws[f'A{last}'].font      = Font(bold=True, size=9, color='1F3864')
    ws[f'A{last}'].alignment = Alignment(horizontal='left')

    build_legend(ws, last + 1)


def main():
    print('Loading rates master...')
    with open(RATES_JSON, encoding='utf-8') as f:
        rates_raw = json.load(f)
    rates = {item['code']: item for item in rates_raw}

    print('Parsing converted Excel files...')
    all_items = parse_all(VOL1_CONV, VOL2_CONV)

    # Group by section
    from collections import defaultdict
    by_section = defaultdict(list)
    for item in all_items:
        by_section[item['sheet']].append(item)

    print('Running shadow calculations...')
    by_section_results = {}
    for section, items in by_section.items():
        by_section_results[section] = [shadow_calculate(i, rates) for i in items]

    print('Building audit workbook...')
    wb = openpyxl.Workbook()
    wb.remove(wb.active)   # remove default sheet

    # ── Summary sheet ─────────────────────────────────────────────────
    ws_sum = wb.create_sheet(title='SUMMARY')
    ws_sum.views.sheetView[0].showGridLines = False
    ws_sum.merge_cells('A1:G1')
    ws_sum['A1'].value     = 'CPWD DAR 2019  |  Rate Audit Summary'
    ws_sum['A1'].font      = Font(bold=True, size=13, color=C_TITLE_FG)
    ws_sum['A1'].fill      = PatternFill('solid', fgColor=C_TITLE_BG)
    ws_sum['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws_sum.row_dimensions[1].height = 30

    sum_hdrs = ['Sub-Head', 'Section Name', 'Total Items', 'SKIP', 'Audited', 'PASS', 'Pass %']
    for col, h in enumerate(sum_hdrs, 1):
        c = ws_sum.cell(row=3, column=col, value=h)
        c.font      = Font(bold=True, size=10, color=C_HDR_FG)
        c.fill      = PatternFill('solid', fgColor=C_HDR_BG)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border    = _thin()
    ws_sum.row_dimensions[3].height = 20

    for col, w in zip(range(1, 8), [12, 32, 12, 8, 10, 8, 10]):
        ws_sum.column_dimensions[get_column_letter(col)].width = w

    # ── One sheet per section in SHEET_MAP order ──────────────────────
    sum_row = 4
    for section_key, display_name in SHEET_MAP.items():
        results = by_section_results.get(section_key, [])
        vol = '1' if section_key <= '12_Roofing' else '2'

        # Sub-head number from key
        sub_num = section_key.split('_')[0]

        # Summary row
        total   = len(results)
        skipped = sum(1 for r in results if r.status == 'SKIP')
        audited = total - skipped
        passed  = sum(1 for r in results if r.status == 'PASS')
        pct     = f'{100*passed/audited:.0f}%' if audited else '—'
        fill    = PatternFill('solid', fgColor=C_PASS if audited and passed == audited
                              else (C_FAIL if audited and passed == 0 else C_ALT))

        row_vals = [sub_num, display_name, total, skipped, audited, passed, pct]
        for col, v in enumerate(row_vals, 1):
            c = ws_sum.cell(row=sum_row, column=col, value=v)
            c.font      = Font(size=9)
            c.alignment = Alignment(horizontal='center' if col != 2 else 'left',
                                    vertical='center')
            c.border    = _thin()
            if col == 7:
                c.fill = fill
        ws_sum.row_dimensions[sum_row].height = 15
        sum_row += 1

        if not results:
            print(f'  {section_key}: no items found, skipping sheet')
            continue

        ws = wb.create_sheet(title=display_name[:31])
        build_sheet(ws, section_key, display_name, results, vol)
        print(f'  Built: {display_name[:31]}  ({total} items, {passed}/{audited} PASS)')

    # ── Save ──────────────────────────────────────────────────────────
    print(f'\nSaving to {OUT_FILE} ...')
    wb.save(OUT_FILE)
    size_kb = os.path.getsize(OUT_FILE) / 1024
    print(f'Done. {size_kb:.0f} KB — {len(wb.sheetnames)} sheets')


if __name__ == '__main__':
    main()
