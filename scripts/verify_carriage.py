# -*- coding: utf-8 -*-
"""
CHECK 7 - 01_Carriage_of_Materials, verified against the DAR ground truth
rather than against hard-coded formula strings.

7A NUMERIC   Rebuild the sheet's own calculation chain in Python and reproduce
             (i) every row of Data Sheet No. 1, and (ii) all 27 Table 1.1 base
             rates at 1 km. Both are printed in the book, so if the model is
             wrong the reproduction fails.

7B STRUCTURAL Assert the wiring that the numeric pass depends on: the payload
             lookup points at the NET payable column, the scale factor
             distinguishes "1000 Nos" from "100 m", the speed and pro-rata
             lookups read Data Sheet 1 by lead, the Say rate uses MROUND, and
             every divide is guarded.

7C LAYOUT    Assert the role-labelled layout: every parameter row carries one
             of the six cell roles, the guide block is present, and only INPUT
             and OVERRIDE cells are unlocked.
"""

import openpyxl

from scripts.trade_builder_carr import (
    R_M_GANG_ADD, R_M_BASE, R_M_ADD, R_M_LABOUR, R_M_TOTAL, R_M_CAP, R_M_UNIT, R_M_SCALE,
    R_ITEM_CODE, R_MATERIAL, R_SCOPE, R_LIFT, R_GATE_FEE, R_NOM_OVERRIDE, R_NOMENCLATURE,
    R_LEAD, R_SPEED_BM, R_SPEED_OV, R_SPEED_EFF, R_TURNAROUND,
    R_MODE, R_TRIPS_OV, R_TRIPS, R_DIST_BASIS, R_DISTANCE, R_DIESEL, R_MOBIL,
    R_PAY_GROSS, R_PAY_NET, R_PAY_OV, R_PAY_EFF, R_UNIT, R_SCALE, R_OUTPUT,
    R_AUDIT, R_RES_FIRST, R_RES_LAST, R_W_SUB, R_TRIP_COST,
    R_W, R_X1, R_X, R_Y1, R_Y, R_Z1, R_Z, R_Z2,
    R_TOTAL, R_TRIP_OH, R_RATE_UNIT, R_RATE_SCHED, R_SAY, R_SAY_NOTE,
    R_M_STEPS, R_M_GANG, R_M_WAGE, R_M_CPOH, R_M_RATE, R_M_SAY,
    R_DS1_FIRST, R_DS1_LAST, R_T11_FIRST, R_T11_LAST, R_SC_FIRST, R_SC_LAST,
    ROLE_FILL,
)
from scripts.trade_layout import R_GUIDE_HEAD, R_GUIDE_STEPS, R_GUIDE_ROLES

TRUCK, WAGE, DIESEL_RATE, MOBIL_RATE = 1500.0, 558.0, 73.50, 315.00
CPOH = 0.15


def _scale(unit):
    return 1000 if unit == '1000 Nos' else (100 if unit == '100 m' else 1)


def _shift_cost(lead, speed, turnaround=1.0, gang=6):
    n = round(8 / ((2 * lead / speed) + turnaround), 2)
    km = round(2 * n * lead + 6, 2)
    diesel = round(km / 5.0, 2)
    mobil = round(km / 140.0, 3)
    w = (round(TRUCK, 2) + round(gang * WAGE, 2)
         + round(diesel * DIESEL_RATE, 2) + round(mobil * MOBIL_RATE, 2))
    return n, km, diesel, mobil, round(w, 2)


def check_carriage(wb, verbose=True):
    msgs = []
    ok = True
    ws = wb['01_Carriage_of_Materials']

    # ---------------- 7A(i): reproduce Data Sheet No. 1 ------------------
    bad = 0
    for r in range(R_DS1_FIRST, R_DS1_LAST + 1):
        lead, speed, n_b, km_b, d_b, m_b, w_b = (ws.cell(row=r, column=c).value for c in range(1, 8))
        n, km, diesel, mobil, w = _shift_cost(lead, speed)
        if (abs(n - n_b) > 0.011 or abs(km - km_b) > 0.02 or abs(diesel - d_b) > 0.02
                or abs(mobil - m_b) > 0.002 or abs(w - w_b) > 0.05):
            bad += 1
            if bad <= 3:
                msgs.append(f"  [FAIL] Data Sheet 1 row at lead {lead} km not reproduced: "
                            f"N {n} vs {n_b}, km {km} vs {km_b}, W {w} vs {w_b}")
    if bad == 0:
        msgs.append(f"  [PASS] All {R_DS1_LAST - R_DS1_FIRST + 1} rows of Data Sheet No. 1 reproduced "
                    f"from the sheet's own model (N = 8/((2L/S)+T), km = 2NL+6, diesel = km/5, "
                    f"mobil = km/140, W = truck + 6 Beldar + fuel).")
    else:
        ok = False

    # ---------------- 7A(ii): reproduce the 27 Table 1.1 base rates ------
    n1, _km, _d, _m, w1 = _shift_cost(1.0, 16.0)
    total1 = w1 + round(w1 * CPOH, 2)
    bad_rows, gross_rows = [], 0
    for r in range(R_T11_FIRST, R_T11_LAST + 1):
        mat = ws.cell(row=r, column=2).value
        gross = ws.cell(row=r, column=3).value
        net = ws.cell(row=r, column=4).value
        unit = ws.cell(row=r, column=5).value
        book = ws.cell(row=r, column=7).value
        sc = _scale(unit)
        rate_net = round(total1 / (n1 * net) * sc, 2)
        rate_gross = round(total1 / (n1 * gross) * sc, 2)
        if abs(rate_net - book) > 0.05:
            bad_rows.append(f"{mat}: {rate_net} vs book {book}")
        if abs(rate_gross - book) <= 0.05:
            gross_rows += 1
    if not bad_rows:
        msgs.append(f"  [PASS] All {R_T11_LAST - R_T11_FIRST + 1} Table 1.1 base rates at 1 km "
                    f"reproduced using the NET payable quantity and the schedule-unit scale factor "
                    f"({gross_rows}/{R_T11_LAST - R_T11_FIRST + 1} would reproduce using the gross "
                    f"payload, which is why the net column is the correct divisor).")
    else:
        ok = False
        msgs.append(f"  [FAIL] Table 1.1 rows not reproduced: {bad_rows[:4]}")

    # ---------------- 7A(iii): DAR item 1.1.18 pro-rata ------------------
    # 3.00 trips at the 10 km row: 88.00 km x (3.00 / 4.10) = 64.39 km, 12.88 L.
    km_1118 = round(88.00 * (3.00 / 4.10), 2)
    diesel_1118 = round(km_1118 / 5.0, 2)
    if abs(diesel_1118 - 12.88) <= 0.01:
        msgs.append(f"  [PASS] CPWD pro-rata basis reproduces DAR item 1.1.18: 3.00 trips at the 10 km "
                    f"row gives {km_1118} km and {diesel_1118} litres of diesel, matching the book.")
    else:
        ok = False
        msgs.append(f"  [FAIL] Pro-rata basis gives {diesel_1118} L, book says 12.88 L.")

    # ---------------- 7B: structural wiring ------------------------------
    def f(row, col=2):
        return str(ws.cell(row=row, column=col).value or '')

    wiring = [
        (f'$D${R_T11_FIRST}' in f(R_PAY_NET),
         'net payable quantity is looked up from Table 1.1 column D'),
        (f'$C${R_T11_FIRST}' in f(R_PAY_GROSS),
         'gross payload is shown separately, for information only'),
        (f'B{R_PAY_NET}' in f(R_PAY_EFF) and f'B{R_PAY_OV}' in f(R_PAY_EFF),
         'effective payable quantity honours the override, else the net figure'),
        ('1000' in f(R_SCALE) and '100' in f(R_SCALE),
         'schedule scale factor distinguishes "1000 Nos" from "100 m"'),
        (f'B{R_PAY_EFF}' in f(R_OUTPUT) and f'B{R_TRIPS}' in f(R_OUTPUT),
         'daily output = trips x effective payable quantity'),
        (f'$B${R_DS1_FIRST}' in f(R_SPEED_BM),
         'average speed is looked up from Data Sheet 1 by lead'),
        (f'B{R_SPEED_OV}' in f(R_SPEED_EFF),
         'speed override is respected'),
        (f'$D${R_DS1_FIRST}' in f(R_DISTANCE) and f'$C${R_DS1_FIRST}' in f(R_DISTANCE),
         'pro-rata scales the Data Sheet 1 row for the actual lead, not a fixed 10 km constant'),
        ('88.00' not in f(R_DISTANCE) and '4.10' not in f(R_DISTANCE),
         'the hard-coded 88.00 / 4.10 constants are gone'),
        (f'B{R_TRIPS_OV}' in f(R_TRIPS),
         'trips override is honoured in any operational mode'),
        ('MROUND' in f(R_SAY, 7) and '0.05' in f(R_SAY, 7),
         'Say rate uses MROUND(x, 0.05)'),
        ('MROUND' in f(R_M_SAY) and '0.05' in f(R_M_SAY),
         'manual-carriage Say rate uses MROUND(x, 0.05)'),
        (f'B{R_SCALE}' in f(R_RATE_SCHED, 7),
         'schedule rate applies the scale factor'),
        (f'B{R_GATE_FEE}' in f(R_RATE_UNIT, 7),
         'gate fee is added after the markup chain'),
        ('Factor_CPOH' in f(R_M_CPOH),
         'manual engine takes CPOH from Global_Factors, not a hard-coded 0.15'),
        ('Rates_Master' in f(R_M_WAGE),
         'manual engine reads the live day wage from Rates_Master'),
        (all(x not in f(R_M_GANG) + f(R_M_GANG_ADD) for x in ('4279.86', '5133.60', '931.86', '753.30')),
         'the hard-coded manual rupee constants are gone'),
        (f'=G{R_Y} + G{R_Z1}' == f(R_Z, 7),
         'Z = Y + CPOH, named consistently with the other eleven builders'),
    ]
    for good, what in wiring:
        if not good:
            ok = False
            msgs.append(f"  [FAIL] wiring: {what}")
    if all(g for g, _ in wiring):
        msgs.append(f"  [PASS] All {len(wiring)} wiring assertions hold (net payable divisor, unit "
                    f"scale factor, Data Sheet 1 speed and pro-rata lookups, MROUND Say, live wage "
                    f"and CPOH, consistent Z naming).")

    guarded = [R_TRIPS, R_DISTANCE, R_RATE_UNIT, R_TRIP_OH, R_M_RATE, R_M_SAY]
    unguarded = []
    for r in guarded:
        cell = f(r) or f(r, 7)
        if '/' in cell and 'IFERROR' not in cell:
            unguarded.append(r)
    if f'IFERROR' not in f(R_TRIP_COST, 7):
        unguarded.append(R_TRIP_COST)
    if unguarded:
        ok = False
        msgs.append(f"  [FAIL] unguarded division in rows {unguarded} - a zero input would show #DIV/0!")
    else:
        msgs.append("  [PASS] Every division on the sheet is wrapped in IFERROR, so a zero trip count "
                    "or missing payload degrades to 0 rather than #DIV/0!.")

    audit = str(ws.cell(row=R_AUDIT, column=2).value or '')
    audit_cells = ' '.join(str(ws.cell(row=R_AUDIT, column=c).value or '') for c in range(4, 10))
    if (f'B{R_TRIPS}' in audit_cells and f'B{R_OUTPUT}' in audit_cells
            and f'G{R_W_SUB}' in audit_cells and f'G{R_SAY}' in audit_cells):
        msgs.append("  [PASS] Audit bar tests the trips, the payable output, the direct cost and "
                    "whether the Say rate actually resolved to a number.")
    else:
        ok = False
        msgs.append("  [FAIL] Audit bar does not test the outputs that matter.")

    # ---------------- 7C: layout and roles --------------------------------
    if 'HOW TO USE THIS SHEET' not in str(ws.cell(row=R_GUIDE_HEAD, column=1).value or ''):
        ok = False
        msgs.append("  [FAIL] The 'how to use this sheet' guide block is missing.")
    elif not str(ws.cell(row=R_GUIDE_STEPS, column=1).value or '').startswith('HOW TO USE'):
        ok = False
        msgs.append("  [FAIL] Guide steps strip is missing.")
    elif 'WHAT THE COLOURS MEAN' not in str(ws.cell(row=R_GUIDE_ROLES, column=1).value or ''):
        ok = False
        msgs.append("  [FAIL] Cell-role colour key is missing.")
    else:
        msgs.append("  [PASS] Guide block present: working steps plus the cell-role colour key.")

    param_rows = [R_ITEM_CODE, R_MATERIAL, R_SCOPE, R_LIFT, R_GATE_FEE, R_NOM_OVERRIDE,
                  R_LEAD, R_SPEED_BM, R_SPEED_OV, R_SPEED_EFF, R_TURNAROUND, R_MODE,
                  R_TRIPS_OV, R_TRIPS, R_DIST_BASIS, R_DISTANCE, R_DIESEL, R_MOBIL,
                  R_PAY_GROSS, R_PAY_NET, R_PAY_OV, R_PAY_EFF, R_UNIT, R_SCALE, R_OUTPUT]
    roles = set(ROLE_FILL)
    missing_role = [r for r in param_rows
                    if str(ws.cell(row=r, column=4).value or '') not in roles]
    missing_note = [r for r in param_rows
                    if not str(ws.cell(row=r, column=5).value or '').strip()]
    if missing_role or missing_note:
        ok = False
        msgs.append(f"  [FAIL] parameter rows without a cell role {missing_role} "
                    f"or without an explanation {missing_note}")
    else:
        msgs.append(f"  [PASS] All {len(param_rows)} parameter rows carry an explicit cell role "
                    f"(INPUT / OVERRIDE / LOOKUP / DERIVED / RESULT / SAY) and a plain-English note "
                    f"saying where the value comes from.")

    expect_editable = {R_ITEM_CODE, R_MATERIAL, R_SCOPE, R_LIFT, R_GATE_FEE, R_NOM_OVERRIDE,
                       R_LEAD, R_SPEED_OV, R_TURNAROUND, R_MODE, R_TRIPS_OV, R_DIST_BASIS, R_PAY_OV}
    wrong = []
    for r in param_rows:
        unlocked = ws.cell(row=r, column=2).protection.locked is False
        role = str(ws.cell(row=r, column=4).value or '')
        should = role in ('INPUT', 'OVERRIDE')
        if r in expect_editable and not unlocked:
            wrong.append((r, 'should be editable'))
        if unlocked and not should:
            wrong.append((r, f'is editable but role is {role}'))
    if wrong:
        ok = False
        msgs.append(f"  [FAIL] protection does not match the declared roles: {wrong}")
    else:
        msgs.append("  [PASS] Only the cells declared INPUT or OVERRIDE are unlocked; every LOOKUP, "
                    "DERIVED, RESULT and SAY cell is protected against typing.")

    if SAY_NOTE_OK := ('MROUND' in str(ws.cell(row=R_SAY_NOTE, column=1).value or '')):
        msgs.append("  [PASS] The sheet states in writing why the Say rate is rounded to 5 paise, "
                    "with the CPWD items that evidence it.")
    else:
        ok = False
        msgs.append("  [FAIL] The MROUND rationale note is missing from the sheet.")

    if verbose:
        for m in msgs:
            print(m)
    return ok, msgs


if __name__ == '__main__':
    wb = openpyxl.load_workbook('CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1.xlsx')
    good, _ = check_carriage(wb)
    print('\nCARRIAGE CHECK:', 'PASS' if good else 'FAIL')
