"""
vol2_chapter_builder.py
Rebuilds any Vol 2 chapter worksheet from its extracted ch{NN}_items.json.
Follows the same layout as support_builder_finishing.py (Ch.13).

Usage:
    python scripts/vol2_chapter_builder.py --chapter 14
    python scripts/vol2_chapter_builder.py --all
"""

import sys, json, argparse, re
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from scripts.paths import WB_VOL2_FILE

CHAPTER_SHEETS = {
    13: "13_Finishing",
    14: "14_Repairs_to_Buildings",
    15: "15_Dismantling_Demolishing",
    16: "16_Road_Work",
    17: "17_Sanitary_Installations",
    18: "18_Water_Supply",
    19: "19_Drainage",
    20: "20_Pile_Work",
    21: "21_Aluminium_Work",
    22: "22_Water_Proofing",
    23: "23_Rain_Water_Harvesting",
    24: "24_Heritage_Buildings",
    25: "25_Structural_Glazing",
    26: "26_New_Technologies",
}

CHAPTER_TITLES = {
    13: "SUB-HEAD 13: FINISHING WORKS",
    14: "SUB-HEAD 14: REPAIRS TO BUILDINGS",
    15: "SUB-HEAD 15: DISMANTLING AND DEMOLISHING",
    16: "SUB-HEAD 16: ROAD WORK",
    17: "SUB-HEAD 17: SANITARY INSTALLATIONS",
    18: "SUB-HEAD 18: WATER SUPPLY",
    19: "SUB-HEAD 19: DRAINAGE",
    20: "SUB-HEAD 20: PILE WORK",
    21: "SUB-HEAD 21: ALUMINIUM WORK",
    22: "SUB-HEAD 22: WATER PROOFING",
    23: "SUB-HEAD 23: RAIN WATER HARVESTING",
    24: "SUB-HEAD 24: CONSERVATION OF HERITAGE BUILDINGS",
    25: "SUB-HEAD 25: STRUCTURAL GLAZING / ACP",
    26: "SUB-HEAD 26: NEW TECHNOLOGIES AND MATERIALS",
}

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "reference_json"

C_TITLE_BG  = "1A1A2E"
C_PARENT_BG = "1F4E79"
C_ITEM_BG   = "2E75B6"
C_TH_BG     = "1F4E79"
C_ALT_ROW   = "F2F7FC"
C_SAY_BG    = "E2EFDA"
C_WHITE     = "FFFFFF"

def _fill(h):
    return PatternFill(fill_type="solid", fgColor=h) if h else PatternFill(fill_type=None)

def _tb():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

AL_C = Alignment(horizontal="center", vertical="center")
AL_L = Alignment(horizontal="left",   vertical="center", wrap_text=True)
AL_R = Alignment(horizontal="right",  vertical="center")

def _cell(ws, r, c, val, fnt=None, aln=None, fl=None, brd=None, num=None):
    cell = ws.cell(row=r, column=c, value=val)
    if fnt: cell.font = fnt
    if aln: cell.alignment = aln
    if fl:  cell.fill = fl
    if brd: cell.border = brd
    if num: cell.number_format = num
    return cell


def build_chapter_sheet(wb, ch: int) -> int:
    sn = CHAPTER_SHEETS[ch]
    json_path = DATA_DIR / f"ch{ch:02d}_items.json"
    if not json_path.exists():
        raise FileNotFoundError(f"No item data found: {json_path}")

    with open(json_path, encoding="utf-8") as f:
        items = json.load(f)

    if sn in wb.sheetnames:
        del wb[sn]
    ws = wb.create_sheet(sn)

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 54
    ws.column_dimensions["D"].width = 9
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 13
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 40

    row = 1
    ws.merge_cells(f"A{row}:H{row}")
    _cell(ws, row, 1,
          f"CPWD DAR 2019 – {CHAPTER_TITLES[ch]}  |  Resource & Rate Analysis  ({len(items)} items)",
          Font(name="Calibri", size=13, bold=True, color=C_WHITE),
          AL_C, _fill(C_TITLE_BG))
    ws.row_dimensions[row].height = 22
    row += 1

    ws.merge_cells(f"A{row}:H{row}")
    _cell(ws, row, 1,
          "PDF source: CivilDAR_2019_Vol_2.pdf  |  "
          "SAY rates reproduced verbatim from the book; markup chain W->X->Y->Z per CPWD DAR convention.",
          Font(name="Calibri", size=9, italic=True, color="444444"),
          AL_L, _fill("F0F0F0"))
    row += 1

    last_parent = None
    for it in items:
        code   = it["code"]
        desc   = it.get("desc", "")
        unit   = it["unit"]
        basis  = float(it["basis"])
        say    = float(it["say"])
        resources = [(r["code"], r["desc"], r["unit"], float(r["qty"]), float(r["rate"]))
                     for r in it.get("resources", [])]

        parts  = code.split(".")
        parent = ".".join(parts[:2]) if len(parts) > 2 else code

        if parent != last_parent:
            last_parent = parent
            row += 1
            ws.merge_cells(f"A{row}:H{row}")
            _cell(ws, row, 1, f"Item group {parent}",
                  Font(name="Calibri", size=10, bold=True, color=C_WHITE),
                  AL_L, _fill(C_PARENT_BG), _tb())
            ws.row_dimensions[row].height = 18
            row += 1

        ws.merge_cells(f"A{row}:H{row}")
        _cell(ws, row, 1, f"{code}  {desc}  [{basis:.0f} {unit} basis]",
              Font(name="Calibri", size=9, bold=True, color=C_WHITE),
              AL_L, _fill(C_ITEM_BG), _tb())
        ws.row_dimensions[row].height = 16
        row += 1

        hdrs = ["#", "Code", "Description / Specification", "Unit",
                "Qty / Coeff", "Basic Rate (Rs)", "Amount (Rs)", "Remarks"]
        for ci, h in enumerate(hdrs, 1):
            _cell(ws, row, ci, h,
                  Font(name="Calibri", size=8, bold=True, color=C_WHITE),
                  AL_C, _fill(C_TH_BG), _tb())
        row += 1

        for ri, (rcode, rdesc, runit, rqty, rrate) in enumerate(resources, 1):
            bg = C_ALT_ROW if ri % 2 == 0 else None
            amt = round(rqty * rrate, 2)
            vals = [ri, rcode, rdesc, runit, rqty, rrate, amt, ""]
            for ci, v in enumerate(vals, 1):
                num = "#,##0.00" if ci in (5, 6, 7) else None
                _cell(ws, row, ci, v,
                      Font(name="Calibri", size=8),
                      AL_R if ci in (5, 6, 7) else AL_L,
                      _fill(bg) if bg else PatternFill(fill_type=None),
                      _tb(), num)
            row += 1

        W    = round(sum(rqty * rrate for _, _, _, rqty, rrate in resources), 2)
        water = round(W * 0.01, 2)
        X    = round(W + water, 2)
        gst  = round(X * 0.1405, 2)
        Y    = round(X + gst, 2)
        cpoh = round(Y * 0.15, 2)
        Z    = round(Y + cpoh, 2)
        cess = round(Z * 0.01, 2)
        cost = round(Z + cess, 2)
        rate_calc = round(cost / basis, 4)

        markup_rows = [
            ("W",  "Direct cost (materials + labour + sundries)", W),
            ("X1", "+1% Water charges", water),
            ("X",  "Subtotal X", X),
            ("Y1", "+GST @14.05% on X", gst),
            ("Y",  "Subtotal Y", Y),
            ("Z1", "+15% CPOH on Y", cpoh),
            ("Z",  "Subtotal Z", Z),
            ("Z2", "+1% Labour Welfare Cess on Z", cess),
            ("",   f"Cost of {basis:.0f} {unit}", cost),
            ("",   f"Rate per 1 {unit}", round(rate_calc, 2)),
        ]
        for step, lbl, amt_val in markup_rows:
            _cell(ws, row, 1, step, Font(name="Calibri", size=8, bold=True), AL_C, _fill("EBF3FB"), _tb())
            ws.merge_cells(f"B{row}:F{row}")
            _cell(ws, row, 2, lbl, Font(name="Calibri", size=8), AL_L, _fill("EBF3FB"), _tb())
            c7 = ws.cell(row=row, column=7, value=amt_val)
            c7.number_format = "#,##0.00"
            c7.font = Font(name="Calibri", size=8)
            c7.alignment = AL_R
            c7.fill = _fill("EBF3FB")
            c7.border = _tb()
            row += 1

        ws.merge_cells(f"A{row}:F{row}")
        _cell(ws, row, 1,
              f"SAY (MROUND to Rs 0.05)  -  PDF published rate: Rs {say:.2f} / {unit}",
              Font(name="Calibri", size=9, bold=True, color="1F4E79"),
              AL_C, _fill(C_SAY_BG), _tb())
        _cell(ws, row, 7, say,
              Font(name="Calibri", size=9, bold=True, color="1F4E79"),
              AL_R, _fill(C_SAY_BG), _tb(), "#,##0.00")
        match = "MATCH" if abs(rate_calc - say) < 2.0 else f"DIFF {rate_calc:.2f} vs {say:.2f}"
        _cell(ws, row, 8, match,
              Font(name="Calibri", size=8, bold=True,
                   color="375623" if match == "MATCH" else "9C0006"),
              AL_C, _fill(C_SAY_BG), _tb())
        row += 2

    return len(items)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--chapter", type=int)
    ap.add_argument("--all", action="store_true")
    args = ap.parse_args()

    chapters = list(range(14, 27)) if args.all else ([args.chapter] if args.chapter else [])
    if not chapters:
        ap.print_help(); sys.exit(1)

    wb_path = Path(WB_VOL2_FILE)
    print(f"Opening {wb_path.name} ...")
    wb = openpyxl.load_workbook(str(wb_path))

    for ch in chapters:
        if ch not in CHAPTER_SHEETS:
            print(f"  Ch {ch}: unknown, skipping"); continue
        try:
            n = build_chapter_sheet(wb, ch)
            print(f"  Ch {ch} ({CHAPTER_SHEETS[ch]}): {n} items written")
        except FileNotFoundError as e:
            print(f"  Ch {ch}: SKIP – {e}")

    wb.save(str(wb_path))
    print(f"Saved -> {wb_path}")


if __name__ == "__main__":
    main()
