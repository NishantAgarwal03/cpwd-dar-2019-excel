"""vol1_chapter_extractor.py
Extract Ch.03–12 resource rows from CivilDAR_2019_Vol_1_Converted.xlsx and
write data/reference_json/ch03_items.json … ch12_items.json in the same
{code, unit, basis, say, resources:[{code,desc,unit,qty,rate}]} format as
the Vol 2 JSON files produced by vol2_chapter_extractor.py.

This makes the Vol 1 registry pipeline identical to the Vol 2 pipeline:
    Converted XLSX → ch0N_items.json → cpwd_item_parser.py → registry

Usage:
    python scripts/vol1_chapter_extractor.py            # all chapters 03-12
    python scripts/vol1_chapter_extractor.py --chapter 3
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from scripts.paths import VOL1_CONVERTED_XLSX

OUT_DIR = Path(__file__).resolve().parent.parent / "data" / "reference_json"

CHAPTER_SHEETS = {
    3:  "03_Mortars",
    4:  "04_Concrete_Work",
    5:  "05_RCC_Work",
    6:  "06_Masonry_Work",
    7:  "07_Stone_Work",
    8:  "08_Cladding_Work",
    9:  "09_Wood_and_PVC_Work",
    10: "10_Steel_Work",
    11: "11_Flooring",
    12: "12_Roofing",
}

_BASIS_PAT = re.compile(r"([0-9]+(?:\.[0-9]+)?)\s*(cum|sqm|m\b|metre|kg|nos|tonne|quintal|bag|litre|pair|set|each|running)", re.I)
_SKIP_CODES = {"TOTAL", "Add", "Deduct", "Add 1", "GRAND", "SAY", "Say"}


def _is_item_header(code: str, qty, rate) -> bool:
    """Row is an item header when code has a dot, qty and rate are blank."""
    return bool(code and "." in code and code[0].isdigit() and qty is None and rate is None)


def _is_resource_row(code: str, qty, rate) -> bool:
    """Row is a resource when code is a number (or known pattern) and qty is populated."""
    if not code or qty is None:
        return False
    for skip in _SKIP_CODES:
        if code.startswith(skip):
            return False
    return True


def _parse_basis(desc: str) -> tuple[float, str]:
    """Extract (quantity, unit) from a 'Details of cost for X cum' string."""
    if not desc:
        return 1.0, "nos"
    m = _BASIS_PAT.search(desc)
    if m:
        return float(m.group(1)), m.group(2).lower().rstrip(".")
    return 1.0, "nos"


def extract_chapter(ch: int, wb: openpyxl.Workbook) -> list[dict]:
    sheet_name = CHAPTER_SHEETS[ch]
    if sheet_name not in wb.sheetnames:
        print(f"  [SKIP] sheet '{sheet_name}' not found in workbook")
        return []

    ws = wb[sheet_name]
    items: list[dict] = []
    current: dict | None = None
    pending_basis_desc: str | None = None   # description row immediately after item header

    for row in ws.iter_rows(values_only=True):
        # Normalise the 8-column row
        page, code_raw, desc_raw, unit_raw, qty_raw, rate_raw, amt_raw, *_ = (list(row) + [None] * 8)[:8]
        code = str(code_raw or "").strip()
        desc = str(desc_raw or "").strip()
        unit = str(unit_raw or "").strip()

        # Parse qty and rate as floats (they arrive as strings in the XLSX)
        qty: float | None = None
        rate: float | None = None
        try:
            if qty_raw is not None and str(qty_raw).strip() not in ("", "-"):
                qty = float(str(qty_raw).replace(",", ""))
        except (ValueError, TypeError):
            pass
        try:
            if rate_raw is not None and str(rate_raw).strip() not in ("", "-"):
                rate = float(str(rate_raw).replace(",", "").split()[0])
        except (ValueError, TypeError):
            pass

        # ── New item header ──────────────────────────────────────────────────
        if _is_item_header(code, qty_raw, rate_raw):
            if current and current.get("resources"):
                items.append(current)
            current = {"code": code, "unit": "nos", "basis": 1.0, "say": None, "resources": []}
            # The *next* non-blank description row gives the basis
            pending_basis_desc = desc or None
            continue

        if current is None:
            continue

        # ── Capture basis from the first description line after item header ──
        if pending_basis_desc is not None and desc and not code:
            b, u = _parse_basis(pending_basis_desc + " " + desc)
            current["basis"] = b
            current["unit"] = u
            pending_basis_desc = None
            continue
        if pending_basis_desc is not None and desc:
            b, u = _parse_basis(pending_basis_desc)
            current["basis"] = b
            current["unit"] = u
            pending_basis_desc = None
            # fall through — this row may itself be a resource

        # ── Say / total rate ─────────────────────────────────────────────────
        if code.upper() in ("SAY", "SAY RS.", "SAY RS") or (
            not code and desc and re.search(r"\bsay\b", desc, re.I) and amt_raw
        ):
            try:
                current["say"] = float(str(amt_raw).replace(",", "").split()[0])
            except (ValueError, TypeError, AttributeError):
                pass
            continue

        # ── Resource row ──────────────────────────────────────────────────────
        if _is_resource_row(code, qty_raw, rate_raw) and qty is not None and rate is not None:
            current["resources"].append({
                "code": code,
                "desc": desc or code,
                "unit": unit or "day",
                "qty":  qty,
                "rate": rate,
            })

    # Commit last item
    if current and current.get("resources"):
        items.append(current)

    return items


def run(chapters: list[int]) -> None:
    wb = openpyxl.load_workbook(str(VOL1_CONVERTED_XLSX), read_only=True, data_only=True)
    for ch in chapters:
        print(f"Ch {ch:02d} ({CHAPTER_SHEETS.get(ch, '?')}) …", end=" ", flush=True)
        items = extract_chapter(ch, wb)
        out_path = OUT_DIR / f"ch{ch:02d}_items.json"
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(items, f, indent=2, ensure_ascii=False)
        print(f"{len(items)} items → {out_path.name}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--chapter", type=int, help="Single chapter number (3-12)")
    args = ap.parse_args()
    if args.chapter:
        if args.chapter not in CHAPTER_SHEETS:
            sys.exit(f"Chapter {args.chapter} not in range 3-12")
        run([args.chapter])
    else:
        run(sorted(CHAPTER_SHEETS.keys()))


if __name__ == "__main__":
    main()
