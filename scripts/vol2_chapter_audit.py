"""
vol2_chapter_audit.py
Audit any Vol 2 chapter: compares SAY rates in the built worksheet
against the PDF primary source.  Tolerance Rs 0.05.

Usage:
    python scripts/vol2_chapter_audit.py --chapter 14
    python scripts/vol2_chapter_audit.py --all
    python scripts/vol2_chapter_audit.py --all --export reports/vol2_audit.md
"""

import sys, re, json, argparse
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl, pdfplumber
from scripts.paths import WB_VOL2_FILE, PDF_CIVIL_DAR_VOL2

TOLERANCE = 0.05

CHAPTER_PAGES = {
    13: (9,   109),
    14: (109, 179),
    15: (179, 227),
    16: (227, 371),
    17: (371, 485),
    18: (485, 753),
    19: (753, 855),
    20: (855, 887),
    21: (887, 919),
    22: (919, 945),
    23: (945, 971),
    24: (971, 979),
    25: (979, 991),
    26: (991, 1110),
}

CHAPTER_SHEETS = {
    13: "13_Finishing",
    14: "14_Repairs_to_Buildings",
    15: "15_Dismantling_Demolishing",
    16: "16_Road_Work",
    17: "17_Sanitary_Installations",
    18: "18_Water_Supply",
    19: "19_Drainage",
    20: "20_Pile_Work",
    21: "21_Aluminium_Work",
    22: "22_Water_Proofing",
    23: "23_Rain_Water_Harvesting",
    24: "24_Heritage_Buildings",
    25: "25_Structural_Glazing",
    26: "26_New_Technologies",
}


def read_sheet_rates(wb_path, ch):
    wb = openpyxl.load_workbook(str(wb_path), read_only=True, data_only=True)
    sn = CHAPTER_SHEETS[ch]
    if sn not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[sn]
    rates = {}
    code_pat = re.compile(rf"^({re.escape(str(ch))}\.\d+(?:\.\d+)?)\s+")
    current_code = None
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        cell_a = str(row[0]).strip()
        m = code_pat.match(cell_a)
        if m:
            current_code = m.group(1)
        if current_code and "SAY" in cell_a and "PDF published rate" in cell_a:
            try:
                rates[current_code] = float(row[6])
            except (TypeError, ValueError):
                pass
    wb.close()
    return rates


def read_pdf_rates(pdf_path, ch):
    start, end = CHAPTER_PAGES[ch]
    with pdfplumber.open(str(pdf_path)) as pdf:
        pages = [pdf.pages[i].extract_text() or "" for i in range(start, end)]
    full = "\n".join(pages)
    ch_s = str(ch)
    pat = re.compile(rf"(?m)^({re.escape(ch_s)}\.\d+(?:\.\d+)?)\s+(.+?)$")
    segments = list(pat.finditer(full))
    rates = {}
    for idx, m in enumerate(segments):
        code = m.group(1)
        start_pos = m.end()
        end_pos = segments[idx + 1].start() if idx + 1 < len(segments) else len(full)
        block = full[start_pos:end_pos]
        say_m = re.search(r"\bSay\s+([\d,]+\.?\d*)", block)
        if say_m:
            rates[code] = float(say_m.group(1).replace(",", ""))
    return rates


def audit_chapter(ch, wb_path, pdf_path):
    sheet_rates = read_sheet_rates(wb_path, ch)
    pdf_rates   = read_pdf_rates(pdf_path, ch)

    if not sheet_rates:
        return {"ch": ch, "status": "NO_SHEET", "passed": [], "failed": [],
                "missing_sheet": [], "missing_pdf": [],
                "total_pdf": len(pdf_rates), "total_sheet": 0}

    def sort_key(c):
        try: return [int(x) for x in c.split(".")]
        except: return [0]

    all_codes = sorted(set(list(sheet_rates) + list(pdf_rates)), key=sort_key)
    passed, failed, miss_sheet, miss_pdf = [], [], [], []

    for code in all_codes:
        in_s = code in sheet_rates
        in_p = code in pdf_rates
        if not in_s:
            miss_sheet.append(code)
        elif not in_p:
            miss_pdf.append(code)
        else:
            diff = abs(sheet_rates[code] - pdf_rates[code])
            row  = (code, sheet_rates[code], pdf_rates[code], diff)
            (passed if diff <= TOLERANCE else failed).append(row)

    return {"ch": ch, "status": "OK",
            "passed": passed, "failed": failed,
            "missing_sheet": miss_sheet, "missing_pdf": miss_pdf,
            "total_pdf": len(pdf_rates), "total_sheet": len(sheet_rates)}


def format_chapter(res):
    ch = res["ch"]
    W  = 104
    lines = []
    lines.append("=" * W)
    lines.append(f"  CH.{ch}  {CHAPTER_SHEETS.get(ch, '')}  |  "
                 f"PDF:{res['total_pdf']}  Sheet:{res['total_sheet']}  "
                 f"PASS:{len(res['passed'])}  FAIL:{len(res['failed'])}  "
                 f"MissSheet:{len(res['missing_sheet'])}  MissPDF:{len(res['missing_pdf'])}")
    lines.append("=" * W)

    if res["status"] == "NO_SHEET":
        lines.append("  *** SHEET NOT FOUND IN WORKBOOK – builder has not been run ***")
        return "\n".join(lines)

    total = len(res["passed"]) + len(res["failed"])
    pct   = 100 * len(res["passed"]) / total if total else 0
    lines.append(f"  RESULT: {len(res['passed'])}/{total} PASS  ({pct:.1f}%)")

    if res["failed"]:
        lines.append("")
        lines.append("  FAILURES:")
        lines.append(f"  {'CODE':<16} {'SHEET':>10} {'PDF':>10} {'DIFF':>8}")
        lines.append("  " + "-" * 48)
        for code, sr, pr, diff in res["failed"]:
            lines.append(f"  {code:<16} {sr:>10.2f} {pr:>10.2f} {diff:>8.2f}")

    if res["missing_sheet"]:
        lines.append("")
        lines.append(f"  MISSING IN SHEET ({len(res['missing_sheet'])} items):")
        lines.append("  " + ", ".join(res["missing_sheet"][:20])
                     + (" ..." if len(res["missing_sheet"]) > 20 else ""))

    if res["missing_pdf"]:
        lines.append("")
        lines.append(f"  EXTRA IN SHEET (not in PDF, {len(res['missing_pdf'])} items):")
        lines.append("  " + ", ".join(res["missing_pdf"][:20]))

    return "\n".join(lines)


def format_summary(results):
    W = 104
    lines = ["", "=" * W,
             "  VOL 2 AUDIT SUMMARY  –  All chapters",
             "=" * W,
             f"  {'CH':<5} {'Sheet':<35} {'PDF':>5} {'Sheet':>6} {'PASS':>6} {'FAIL':>6} {'%':>6}",
             "  " + "-" * 68]
    for res in results:
        ch   = res["ch"]
        sn   = CHAPTER_SHEETS.get(ch, "")
        tot  = len(res["passed"]) + len(res["failed"])
        pct  = f"{100*len(res['passed'])/tot:.0f}%" if tot else "—"
        stat = "NO SHEET" if res["status"] == "NO_SHEET" else ""
        lines.append(
            f"  {ch:<5} {sn:<35} {res['total_pdf']:>5} {res['total_sheet']:>6} "
            f"{len(res['passed']):>6} {len(res['failed']):>6} {pct:>6}  {stat}")
    lines.append("=" * W)
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--chapter", type=int)
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--export", metavar="FILE")
    args = ap.parse_args()

    chapters = list(range(13, 27)) if args.all else ([args.chapter] if args.chapter else [])
    if not chapters:
        ap.print_help(); sys.exit(1)

    wb  = Path(WB_VOL2_FILE)
    pdf = Path(PDF_CIVIL_DAR_VOL2)

    results = []
    for ch in chapters:
        print(f"Auditing Ch.{ch} ...", end=" ", flush=True)
        res = audit_chapter(ch, wb, pdf)
        tot = len(res["passed"]) + len(res["failed"])
        print(f"{len(res['passed'])}/{tot} pass")
        results.append(res)

    lines = []
    for res in results:
        lines.append(format_chapter(res))
        lines.append("")
    if len(results) > 1:
        lines.append(format_summary(results))

    report = "\n".join(lines)

    if args.export:
        out = Path(args.export)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(report, encoding="utf-8")
        print(f"Report saved -> {out}")
    else:
        print(report)


if __name__ == "__main__":
    main()
