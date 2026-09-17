"""
support_builder_wood_pvc.py
Rebuilds the existing '09_Wood_and_PVC_Work' worksheet with First-Principles Resource,
Productivity, Gang, Machinery and Material Analysis per CPWD DAR 2019 Vol 1:
  Item Code | Labour / Machine / Material | Work done | Condition / When used | Category | Productivity | Quantity
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
import os, sys, shutil

# Styling definitions matching standard
C_TITLE_BG   = "1A1A2E"   # Navy dark
C_SUBTITLE_BG= "E8F4FD"   # Soft blue
C_HDR_PARENT = "1F4E79"   # Dark blue
C_HDR_SUB    = "2E75B6"   # Mid blue
C_TH_BG      = "1F4E79"   # Table header dark blue
C_WHITE      = "FFFFFF"
C_ALT_ROW    = "F9FBFD"

def fill(hex_col):
    if not hex_col:
        return PatternFill(fill_type=None)
    return PatternFill(fill_type="solid", fgColor=hex_col)

def thin_border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

AL_C = Alignment(horizontal="center", vertical="center")
AL_L = Alignment(horizontal="left",   vertical="center", wrap_text=True)
AL_R = Alignment(horizontal="right",  vertical="center")

WOOD_PVC_ITEMS = [
    # ─── 9.1 WOOD WORK IN FRAMES (CHOWKHATS) ────────────────────────────────────
    {
        "id": "9.1.1",
        "parent_title": "9.1 WOOD WORK IN FRAMES OF DOORS, WINDOWS & CLERESTORY WINDOWS",
        "title": "9.1.1 Second class teak wood wrought framed and fixed in position with hold fast lugs (1 cum basis)",
        "unit": "cum", "base_qty": 1.0,
        "rows": [
            ("Second class teak wood scantling", "Selected timber scantling (Dandeli, Balarshah or Malabar) wrought, planed, rebated and mortise-tenon framed (1.055 cum incl. 5% sawing/planing wastage)", "Premium timber conforming to IS:1003 & IS:4021; moisture content <= 12%, straight grained", "Material", "1.055 cum (1.055 cum/cum)", "1 cum"),
            ("Carriage of Timber", "Haulage and stacking of timber scantlings from yard to carpentry workshop", "Carriage allowance", "Material", "1.055 cum", "1 cum"),
            ("Carpenter (average)", "Marking, sawing, planing true to size, cutting rebates, making mortise & tenon joints, assembling chowkhats and plumbing in position", "Skilled wood craftsman; 20.00 days = 160.00 man-hrs per cum", "Labour", "160.00 man-hrs (0.006 cum/man-hr)", "1 cum"),
            ("Beldar", "Assisting carpenter in holding heavy timbers, lifting assembled chowkhats into wall openings, and temporary strutting", "Labour crew; 1.94 days = 15.56 man-hrs per cum", "Labour", "15.56 man-hrs (0.064 cum/man-hr)", "1 cum"),
            ("Sundries", "Glue, wire nails, wooden plugs, sandpaper, and temporary bamboo braces for holding frames plumb", "Sundries allowance (L.S. 28.00 x cost index)", "Equipment", "L.S. allowance", "1 cum"),
            ("REF#9.53 (Hold fasts)", "Fixing to masonry with 40x5mm flat iron hold fasts embedded in concrete (6 nos. per standard door frame)", "Referenced fixing scope per CPWD REF#9.53 specification", "Reference", "6 nos / chowkhat", "1 cum"),
        ]
    },
    {
        "id": "9.1.2",
        "parent_title": "9.1 WOOD WORK IN FRAMES OF DOORS, WINDOWS & CLERESTORY WINDOWS",
        "title": "9.1.2 Sal wood wrought framed and fixed in position with hold fast lugs (1 cum basis)",
        "unit": "cum", "base_qty": 1.0,
        "rows": [
            ("Sal wood scantling", "Selected hard sal timber scantling, seasoned, wrought, rebated and jointed (1.055 cum incl. 5% conversion wastage)", "Dense hardwood conforming to IS:4021; high load-bearing capacity for heavy frames", "Material", "1.055 cum (1.055 cum/cum)", "1 cum"),
            ("Carriage of Timber", "Haulage of dense sal timber from store to carpentry workshop", "Carriage allowance", "Material", "1.055 cum", "1 cum"),
            ("Carpenter (average)", "Sawing, planing, mortise-and-tenon jointing with hardwood pins, squaring frames, and fixing plumb", "Skilled wood craftsman; 20.00 days = 160.00 man-hrs per cum", "Labour", "160.00 man-hrs (0.006 cum/man-hr)", "1 cum"),
            ("Beldar", "Lifting heavy sal wood frames, positioning in reveals, bracing with wooden props", "Labour crew; 1.94 days = 15.56 man-hrs per cum", "Labour", "15.56 man-hrs (0.064 cum/man-hr)", "1 cum"),
            ("Sundries", "Synthetic resin adhesive, nails, wedges, PoP temporary dots, and sandpaper", "Sundries allowance (L.S. 28.00 x cost index)", "Equipment", "L.S. allowance", "1 cum"),
            ("REF#9.53 (Hold fasts)", "Fixing to masonry with 40x5mm flat iron hold fasts embedded in concrete (6 nos. per door frame)", "Referenced fixing scope per CPWD REF#9.53 specification", "Reference", "6 nos / chowkhat", "1 cum"),
        ]
    },
    {
        "id": "9.1.3",
        "parent_title": "9.1 WOOD WORK IN FRAMES OF DOORS, WINDOWS & CLERESTORY WINDOWS",
        "title": "9.1.3 Kiln seasoned and chemically treated hollock wood framed and fixed in position (1 cum basis)",
        "unit": "cum", "base_qty": 1.0,
        "rows": [
            ("Kiln seasoned treated Hollock wood scantling", "Factory seasoned (IS:1141) and vacuum-pressure treated with copper-chrome-arsenic / borate wood preservative (IS:401) timber scantling (1.055 cum incl. wastage)", "Chemically immunized timber; moisture content 8-12%, resistant to termites and borers", "Material", "1.055 cum (1.055 cum/cum)", "1 cum"),
            ("Carriage of Timber", "Haulage of seasoned timber from kiln store to workshop", "Carriage allowance", "Material", "1.055 cum", "1 cum"),
            ("Carpenter (average)", "Planing, jointing, rebating, frame assembly, corner gusseting, and fixing in position", "Skilled craftsman; 20.00 days = 160.00 man-hrs per cum", "Labour", "160.00 man-hrs (0.006 cum/man-hr)", "1 cum"),
            ("Beldar", "Handling, shifting, and assisting in plumbing and fixing frames into openings", "Labour crew; 1.94 days = 15.56 man-hrs per cum", "Labour", "15.56 man-hrs (0.064 cum/man-hr)", "1 cum"),
            ("Sundries", "Waterproof adhesive, screws, wooden plugs, and bracing stays", "Sundries allowance", "Equipment", "L.S. allowance", "1 cum"),
            ("REF#9.53 (Hold fasts)", "Fixing with 40x5mm flat iron hold fasts embedded in concrete", "Referenced fixing scope per CPWD REF#9.53 specification", "Reference", "6 nos / chowkhat", "1 cum"),
        ]
    },

    # ─── 9.3 WOOD WORK IN FALSE CEILINGS & PARTITIONS ───────────────────────────
    {
        "id": "9.3.1",
        "parent_title": "9.3 WOOD WORK IN FRAMES OF FALSE CEILING, PARTITIONS ETC.",
        "title": "9.3.1 Sal wood in frames of false ceiling, partitions etc. sawn and fixed in position (1 cum basis)",
        "unit": "cum", "base_qty": 1.0,
        "rows": [
            ("Sal wood scantling", "Selected sal wood battens and runners (50x50mm or 75x50mm) sawn to section and jointed (1.10 cum incl. 10% wastage)", "Structural framing timber conforming to IS:4021; cut to grid framework", "Material", "1.100 cum (1.100 cum/cum)", "1 cum"),
            ("Carriage of Timber", "Haulage and vertical shifting to ceiling/staging level", "Carriage allowance", "Material", "1.100 cum", "1 cum"),
            ("M.S. brackets, suspenders & screws", "Galvanized mild steel suspenders (6mm rods or 25x3mm flats) with expansion plugs and wood screws", "Hardware connecting ceiling frame to RCC slab soffit", "Material", "25.00 kg (25.00 kg/cum)", "1 cum"),
            ("Carpenter (average)", "Laying out grid, halving/lap joints at intersections, leveling grid framework, and anchoring suspenders", "Skilled carpenter; 18.00 days = 144.00 man-hrs per cum", "Labour", "144.00 man-hrs (0.007 cum/man-hr)", "1 cum"),
            ("Beldar", "Erecting scaffolding, hoisting runners, holding framing members in ceiling height", "Labour crew; 9.00 days = 72.00 man-hrs per cum", "Labour", "72.00 man-hrs (0.014 cum/man-hr)", "1 cum"),
            ("Sundries", "Rawl plugs, ceiling fasteners, wood glue, and alignment strings", "Sundries allowance", "Equipment", "L.S. allowance", "1 cum"),
        ]
    },

    # ─── 9.5 PANELLED & GLAZED SHUTTERS ─────────────────────────────────────────
    {
        "id": "9.5.1.1",
        "parent_title": "9.5 PANELLED OR PANELLED AND GLAZED SHUTTERS FOR DOORS & WINDOWS",
        "title": "9.5.1.1 Second class teak wood 35 mm thick panelled shutters (10 sqm basis)",
        "unit": "10 sqm", "base_qty": 10.0,
        "rows": [
            ("Second class teak wood planks 35mm", "Selected timber planks for styles, top rail, lock rail, bottom rail and raised panels (0.28 cum net timber incl. 10% conversion wastage)", "Conforming to IS:1003 (Part 1); moisture content <= 12%, knot-free styles and rails", "Material", "0.280 cum (0.028 cum/sqm)", "10 sqm"),
            ("Carriage of Timber", "Transport of planks and finished shutters to site", "Carriage allowance", "Material", "0.280 cum", "10 sqm"),
            ("M.S. / Stainless steel butt hinges 100mm", "Heavy-duty butt hinges (100x58x1.9mm) fixed to styles with matching wood screws (3 nos. per shutter leaf)", "IS:12817 marked hinges", "Material", "15.00 nos (1.50 nos/sqm)", "10 sqm"),
            ("Wood screws 40mm", "Bright finished or brass/SS screws for hinges and fittings", "Hardware", "Material", "120.00 nos (12.00 nos/sqm)", "10 sqm"),
            ("Carpenter (average)", "Planing styles/rails, mortise and tenon jointing, grooving, assembling panels with waterproof glue, wedging, hanging shutter to frame", "Master joinery carpenter; 10.00 days = 80.00 man-hrs per 10 sqm", "Labour", "80.00 man-hrs (0.125 sqm/man-hr)", "10 sqm"),
            ("Beldar", "Assisting carpenter in holding, trimming reveals, sandpapering, and cleaning face", "Labour crew; 4.15 days = 33.20 man-hrs per 10 sqm", "Labour", "33.20 man-hrs (0.301 sqm/man-hr)", "10 sqm"),
            ("Sundries", "Synthetic resin adhesive (IS:851), sandpaper 80/120 grit, PoP setting dots", "Sundries allowance (L.S. 35.88 x cost index)", "Equipment", "L.S. allowance", "10 sqm"),
        ]
    },
    {
        "id": "9.5.1.2",
        "parent_title": "9.5 PANELLED OR PANELLED AND GLAZED SHUTTERS FOR DOORS & WINDOWS",
        "title": "9.5.1.2 Second class teak wood 30 mm thick panelled shutters (10 sqm basis)",
        "unit": "10 sqm", "base_qty": 10.0,
        "rows": [
            ("Second class teak wood planks 30mm", "Selected timber planks for styles, rails and solid wood panels (0.24 cum net timber incl. 10% conversion wastage)", "Conforming to IS:1003 (Part 1); kiln seasoned teak", "Material", "0.240 cum (0.024 cum/sqm)", "10 sqm"),
            ("Carriage of Timber", "Haulage of timber planks to joinery shop", "Carriage allowance", "Material", "0.240 cum", "10 sqm"),
            ("M.S. / Stainless steel butt hinges 100mm", "Butt hinges (100x58x1.9mm) with matching wood screws", "IS:12817 marked hinges", "Material", "15.00 nos (1.50 nos/sqm)", "10 sqm"),
            ("Wood screws 40mm", "Screws for hinges and hardware fixing", "Hardware", "Material", "120.00 nos (12.00 nos/sqm)", "10 sqm"),
            ("Carpenter (average)", "Marking, machining, mortise-tenon joinery, panel fitting, hanging and easing shutters", "Skilled craftsman; 9.50 days = 76.00 man-hrs per 10 sqm", "Labour", "76.00 man-hrs (0.132 sqm/man-hr)", "10 sqm"),
            ("Beldar", "Handling shutters, assisting hanging and finishing", "Labour crew; 4.00 days = 32.00 man-hrs per 10 sqm", "Labour", "32.00 man-hrs (0.313 sqm/man-hr)", "10 sqm"),
            ("Sundries", "Waterproof adhesive, sandpaper, temporary stops", "Sundries allowance", "Equipment", "L.S. allowance", "10 sqm"),
        ]
    },

    # ─── 9.20 & 9.21 FLUSH DOOR SHUTTERS ────────────────────────────────────────
    {
        "id": "9.20.1",
        "parent_title": "9.20 FACTORY MADE FLUSH DOOR SHUTTERS (DECORATIVE VENEER)",
        "title": "9.20.1 35 mm thick decorative flush door shutter (teak veneer both faces) with SS butt hinges (10 sqm basis)",
        "unit": "10 sqm", "base_qty": 10.0,
        "rows": [
            ("Decorative Flush door shutter 35mm", "Solid core blockboard construction flush door shutter with decorative teak ply on both faces conforming to IS:2202 (Part 1) (10 sqm basis)", "Factory made solid blockboard core, preservative treated timber lipping, phenol formaldehyde synthetic resin bonded", "Material", "10.00 sqm (1.000 sqm/sqm)", "10 sqm"),
            ("Carriage of door shutters", "Careful carriage, edge protection and vertical hoisting of flush doors to floor levels", "Carriage allowance", "Material", "10.00 sqm", "10 sqm"),
            ("Stainless steel butt hinges 100x60x2.5mm", "Heavy-duty SS-304 butt hinges conforming to IS:12817 (3 nos. per shutter leaf, approx. 27 nos. per 10 sqm)", "Non-corrosive stainless steel hinges", "Material", "27.00 nos (2.70 nos/sqm)", "10 sqm"),
            ("Stainless steel screws 40mm", "SS counter-sunk screws for fixing hinges into timber chowkhat and shutter edge", "Hardware", "Material", "218.00 nos (21.80 nos/sqm)", "10 sqm"),
            ("Carpenter (average)", "Rebating edge, chiseling hinge recesses in chowkhat & shutter, hanging shutter, easing clearance (3mm all round), fixing latch/stops", "Skilled carpenter; 2.50 days = 20.00 man-hrs per 10 sqm", "Labour", "20.00 man-hrs (0.500 sqm/man-hr)", "10 sqm"),
            ("Beldar", "Carrying shutters, holding leaf in position during screw driving, and surface cleaning", "Labour crew; 2.50 days = 20.00 man-hrs per 10 sqm", "Labour", "20.00 man-hrs (0.500 sqm/man-hr)", "10 sqm"),
            ("Sundries", "Wooden wedges, sandpaper, masking tape, and temporary door stops", "Sundries allowance (L.S. allowance)", "Equipment", "L.S. allowance", "10 sqm"),
        ]
    },
    {
        "id": "9.21.1",
        "parent_title": "9.21 FACTORY MADE FLUSH DOOR SHUTTERS (COMMERCIAL FACE)",
        "title": "9.21.1 35 mm thick commercial flush door shutter (commercial ply both faces) with SS butt hinges (10 sqm basis)",
        "unit": "10 sqm", "base_qty": 10.0,
        "rows": [
            ("Commercial Flush door shutter 35mm", "Solid core blockboard flush door shutter with commercial ply on both faces conforming to IS:2202 (Part 1) (10 sqm basis)", "Factory made solid blockboard core, preservative treated internal frame, hot-press bonded", "Material", "10.00 sqm (1.000 sqm/sqm)", "10 sqm"),
            ("Carriage of door shutters", "Transport and handling of door shutters to work locations", "Carriage allowance", "Material", "10.00 sqm", "10 sqm"),
            ("Stainless steel butt hinges 100x60x2.5mm", "Heavy-duty SS butt hinges conforming to IS:12817 (3 nos. per shutter leaf)", "Non-corrosive hinges", "Material", "27.00 nos (2.70 nos/sqm)", "10 sqm"),
            ("Stainless steel screws 40mm", "SS wood screws for secure hinge attachment", "Hardware", "Material", "218.00 nos (21.80 nos/sqm)", "10 sqm"),
            ("Carpenter (average)", "Trimming, recessing hinges, hanging shutters, aligning latch bolt and strike plate", "Skilled carpenter; 2.50 days = 20.00 man-hrs per 10 sqm", "Labour", "20.00 man-hrs (0.500 sqm/man-hr)", "10 sqm"),
            ("Beldar", "Assisting carpenter in hoisting and positioning leaves", "Labour crew; 2.50 days = 20.00 man-hrs per 10 sqm", "Labour", "20.00 man-hrs (0.500 sqm/man-hr)", "10 sqm"),
            ("Sundries", "Sandpaper, temporary packing, and drill bits", "Sundries allowance", "Equipment", "L.S. allowance", "10 sqm"),
        ]
    },

    # ─── 9.53 HOLD FASTS FOR FRAMES ─────────────────────────────────────────────
    {
        "id": "9.53",
        "parent_title": "9.53 FLAT IRON HOLD FASTS EMBEDDED IN CEMENT CONCRETE BLOCKS",
        "title": "9.53 Providing 40x5 mm flat iron hold fast 40 cm long with 10mm bolts & concrete block 1:3:6 (10 nos basis)",
        "unit": "10 nos", "base_qty": 10.0,
        "rows": [
            ("M.S. flat iron 40x5mm", "Mild steel flat section 40x5mm cut to 40 cm length with split/fishtail anchor ends (6.72 kg for 10 nos)", "Conforming to IS:2062 Grade E250; 0.672 kg per hold fast", "Material", "0.067 quintal (6.72 kg / 10 nos)", "10 nos"),
            ("M.S. bolts, nuts & washers 10mm", "Galvanized 10mm dia bolts 50mm long with washers for fixing hold fast to timber chowkhat", "Hardware", "Material", "20.00 nos (2.00 nos/holdfast)", "10 nos"),
            ("Portland Cement", "Hydraulic binder decomposed from 0.045 cum of concrete 1:3:6 (REF#4.1.8) (0.045 cum x 0.220 t/cum = 0.010 t)", "OPC-43 Grade cement; foundation block embedding anchor fish tail", "Material", "0.010 tonne (0.001 t/holdfast)", "10 nos"),
            ("Coarse sand", "Zone III sand decomposed from 0.045 cum of concrete 1:3:6 (REF#4.1.8) (0.045 cum x 0.46 = 0.021 cum)", "Clean natural river sand", "Material", "0.021 cum (0.0021 cum/holdfast)", "10 nos"),
            ("Graded stone aggregate 20mm", "20mm graded coarse aggregate decomposed from 0.045 cum of concrete 1:3:6 (REF#4.1.8) (0.045 cum x 0.92 = 0.041 cum)", "Hard broken granite/trap aggregate", "Material", "0.041 cum (0.0041 cum/holdfast)", "10 nos"),
            ("Carriage of materials", "Haulage of steel flats, cement, sand and aggregate to frame fixing locations", "Carriage allowance", "Material", "0.05 tonne/cum equiv", "10 nos"),
            ("Mason", "Excavating pocket in brickwork, bolting hold fast to frame, pouring concrete block 30x10x15cm, consolidating and trowel finish", "Skilled mason; 0.20 day = 1.60 man-hrs per 10 nos", "Labour", "1.60 man-hrs (6.25 nos/man-hr)", "10 nos"),
            ("Beldar", "Mixing 1:3:6 concrete batch, carrying mortar pan, and assisting in placing concrete into anchor pockets", "Labour crew; 0.25 day = 2.00 man-hrs per 10 nos", "Labour", "2.00 man-hrs (5.00 nos/man-hr)", "10 nos"),
            ("Bhisti", "Water curing concrete anchor blocks for 7 days", "Watering crew; 0.10 day = 0.80 man-hrs per 10 nos", "Labour", "0.80 man-hrs (12.50 nos/man-hr)", "10 nos"),
            ("REF#4.1.8 (Cement concrete 1:3:6)", "Referenced embedding concrete: 0.045 cum CC 1:3:6 for 10 hold fast blocks (30x10x15cm each)", "Concrete scope executed per CPWD REF#4.1.8 specification", "Reference", "0.045 cum (0.0045 cum/holdfast)", "10 nos"),
        ]
    },

    # ─── 9.117 & 9.119 FACTORY MADE uPVC & PVC DOOR FRAMES ─────────────────────
    {
        "id": "9.117",
        "parent_title": "9.117 FACTORY MADE uPVC DOOR FRAMES",
        "title": "9.117 Factory made uPVC door frame extruded sections with GI tube reinforcement & SS hinges (10 m run basis)",
        "unit": "10 m", "base_qty": 10.0,
        "rows": [
            ("uPVC extruded door frame profile", "Multi-chambered uPVC hollow profile (wall thickness 2.0mm) mitred and plastic welded at corners (10 m run)", "UV-stabilized unplasticized polyvinyl chloride conforming to IS:15388", "Material", "10.00 metre (1.000 m/m)", "10 m"),
            ("Galvanized steel reinforcement tube", "GI square tube 19x19x1.0mm inserted into hinge-side vertical frame for structural rigidity", "Galvanized mild steel conforming to IS:277", "Material", "4.50 metre (0.450 m/m)", "10 m"),
            ("Galvanized brackets & SS screws", "Corner joint brackets, galvanized anchor brackets, and stainless steel fasteners", "Hardware", "Material", "10.00 nos (1.00 no/m)", "10 m"),
            ("Stainless steel hinges", "SS-304 butt hinges factory fitted to uPVC frame (3 nos. per door frame)", "Hardware", "Material", "6.00 nos (0.60 no/m)", "10 m"),
            ("Carriage of uPVC frames", "Careful transport and handling of welded frames with protective film", "Carriage allowance", "Material", "10.00 metre", "10 m"),
            ("Carpenter (average)", "Positioning frame in masonry opening, drilling holes in jambs, driving anchor dash fasteners, plumbing, and applying silicone seal", "Skilled installer; 0.30 day = 2.40 man-hrs per 10 m", "Labour", "2.40 man-hrs (4.17 m/man-hr)", "10 m"),
            ("Beldar", "Assisting in holding frame, mixing mortar for bracket pockets, and cleaning reveals", "Labour crew; 0.30 day = 2.40 man-hrs per 10 m", "Labour", "2.40 man-hrs (4.17 m/man-hr)", "10 m"),
            ("Sundries", "Dash fasteners 8x75mm, expanding polyurethane foam / silicone sealant, and masking tape", "Sundries allowance", "Equipment", "L.S. allowance", "10 m"),
        ]
    },
    {
        "id": "9.119",
        "parent_title": "9.119 FACTORY MADE PVC DOOR FRAMES",
        "title": "9.119 Factory made PVC door frame of size 50x47 mm with 5mm rigid PVC foam sheet (10 m run basis)",
        "unit": "10 m", "base_qty": 10.0,
        "rows": [
            ("PVC door frame profile 50x47mm", "Extruded 5mm rigid PVC foam sheet profile mitred at corners with 150mm MS square tube brackets (10 m run)", "Rigid PVC foam conforming to CPWD specifications, moisture and termite proof", "Material", "10.00 metre (1.000 m/m)", "10 m"),
            ("M.S. reinforcement tube 19x19mm", "19 gauge galvanized MS square tube inserted inside jambs", "Reinforcement", "Material", "4.40 metre (0.440 m/m)", "10 m"),
            ("EPDM rubber weather gasket", "Continuous perimeter EPDM rubber gasket inserted into frame groove", "Sealing gasket", "Material", "10.00 metre (1.000 m/m)", "10 m"),
            ("Carriage of PVC frames", "Haulage and handling to installation points", "Carriage allowance", "Material", "10.00 metre", "10 m"),
            ("Carpenter (average)", "Fixing frame to masonry wall using 65/100mm screws and rawl plugs, checking squareness and diagonal true", "Skilled installer; 0.30 day = 2.40 man-hrs per 10 m", "Labour", "2.40 man-hrs (4.17 m/man-hr)", "10 m"),
            ("Beldar", "Assisting installer, holding frame plumb, and cleaning reveals", "Labour crew; 0.30 day = 2.40 man-hrs per 10 m", "Labour", "2.40 man-hrs (4.17 m/man-hr)", "10 m"),
            ("Sundries", "M.S. screws 65/100mm, plastic wall plugs, and silicone sealant", "Sundries allowance (L.S. 15.60 x cost index)", "Equipment", "L.S. allowance", "10 m"),
        ]
    },

    # ─── 9.118 & 9.120 FACTORY MADE PVC DOOR SHUTTERS ───────────────────────────
    {
        "id": "9.118.2",
        "parent_title": "9.118 FACTORY MADE PVC DOOR SHUTTERS (HOLLOW SECTION)",
        "title": "9.118.2 30 mm thick factory made PVC door shutter with uPVC hollow styles & rails (10 sqm basis)",
        "unit": "10 sqm", "base_qty": 10.0,
        "rows": [
            ("PVC door shutter 30mm thick", "Factory made PVC shutter with styles/rails of 60x30mm uPVC hollow section (wall 2mm), lock rail 100x30mm, multi-chambered infill panel (10 sqm)", "Pre-fabricated waterproof shutter ideal for bathrooms, toilets and wet areas", "Material", "10.00 sqm (1.000 sqm/sqm)", "10 sqm"),
            ("Galvanized MS tie rods & cleats", "6mm GI tie rods inserted horizontally with nuts/washers and MS 'U' cleats for rigid jointing", "Internal bracing hardware", "Material", "10.00 nos (1.00 no/sqm)", "10 sqm"),
            ("Stainless steel / powder coated hinges", "Heavy-duty hinges with matching SS screws (3 nos. per shutter)", "Hardware", "Material", "13.00 nos (1.30 nos/sqm)", "10 sqm"),
            ("Carriage of door shutters", "Transport and hoisting of PVC shutters with protective packaging", "Carriage allowance", "Material", "10.00 sqm", "10 sqm"),
            ("Carpenter (average)", "Fitting hinges, mortising latch/lock into PVC frame, hanging shutter, and adjusting swing", "Skilled craftsman; 1.68 days = 13.45 man-hrs per 10 sqm", "Labour", "13.45 man-hrs (0.743 sqm/man-hr)", "10 sqm"),
            ("Beldar", "Assisting in lifting and positioning shutters during hanging", "Labour crew; 1.68 days = 13.45 man-hrs per 10 sqm", "Labour", "13.45 man-hrs (0.743 sqm/man-hr)", "10 sqm"),
            ("Sundries", "Fasteners, rubber buffers, masking tape, and cleaning solvent", "Sundries allowance (L.S. allowance)", "Equipment", "L.S. allowance", "10 sqm"),
        ]
    },
    {
        "id": "9.120.1",
        "parent_title": "9.120 FACTORY MADE SOLID PVC FOAM SHUTTERS",
        "title": "9.120.1 30 mm thick plain PVC door shutters with powder coated MS butt hinges (10 sqm basis)",
        "unit": "10 sqm", "base_qty": 10.0,
        "rows": [
            ("PVC rigid foam panelled shutter 30mm", "Factory made rigid PVC foam panelled door shutter of approved shade and finish (10 sqm basis)", "High-density solid PVC foam conforming to CPWD specifications; termite, rot and moisture resistant", "Material", "10.00 sqm (1.000 sqm/sqm)", "10 sqm"),
            ("Carriage of door shutters", "Careful carriage and vertical shifting of PVC shutters", "Carriage allowance", "Material", "10.00 sqm", "10 sqm"),
            ("Powder coated MS butt hinges 100mm", "Powder coated mild steel butt hinges 100x58x1.9mm (3 nos. per shutter leaf, approx. 17 nos. per 10 sqm)", "Hardware", "Material", "17.00 nos (1.70 nos/sqm)", "10 sqm"),
            ("Mild steel screws 40mm & 20mm", "Black enamelled / bright finished screws for hinges and fittings (235 nos. per 10 sqm)", "Hardware", "Material", "235.00 nos (23.50 nos/sqm)", "10 sqm"),
            ("Carpenter (average)", "Fixing hinges to shutter and frame, hanging door leaf, aligning latch bolt and stops", "Skilled carpenter; 1.68 days = 13.45 man-hrs per 10 sqm", "Labour", "13.45 man-hrs (0.743 sqm/man-hr)", "10 sqm"),
            ("Beldar", "Carrying and holding shutters during installation", "Labour crew; 1.68 days = 13.45 man-hrs per 10 sqm", "Labour", "13.45 man-hrs (0.743 sqm/man-hr)", "10 sqm"),
            ("Sundries", "Drill bits, spirit level, wedges, and protective tape", "Sundries allowance", "Equipment", "L.S. allowance", "10 sqm"),
        ]
    },
]

def write_item_block(ws, item, start_row):
    r = start_row

    # Parent category banner
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c_parent = ws.cell(row=r, column=1, value=item["parent_title"])
    c_parent.font = Font(name="Calibri", size=11, bold=True, color=C_WHITE)
    c_parent.fill = fill(C_HDR_PARENT)
    c_parent.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r].height = 24
    for col in range(1, 8):
        ws.cell(row=r, column=col).border = thin_border()
    r += 1

    # Item title banner
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    c_sub = ws.cell(row=r, column=1, value=f"{item['id']}  —  {item['title']}")
    c_sub.font = Font(name="Calibri", size=10, bold=True, color=C_WHITE)
    c_sub.fill = fill(C_HDR_SUB)
    c_sub.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r].height = 22
    for col in range(1, 8):
        ws.cell(row=r, column=col).border = thin_border()
    r += 1

    # Table Header Row (7 standard columns)
    headers = [
        "Item Code",
        "Labour / Machine / Material",
        "Work done",
        "Condition / When used",
        "Category",
        "Productivity",
        "Quantity"
    ]
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=r, column=col_idx, value=h)
        c.font = Font(name="Calibri", size=9, bold=True, color=C_WHITE)
        c.fill = fill(C_TH_BG)
        c.alignment = AL_C
        c.border = thin_border()
    ws.row_dimensions[r].height = 20
    r += 1

    # Data rows
    for row_idx, data in enumerate(item["rows"]):
        bg = C_ALT_ROW if (row_idx % 2 == 1) else C_WHITE
        row_fill = fill(bg)

        c1 = ws.cell(row=r, column=1, value=item["id"])
        c1.alignment = AL_C
        c1.font = Font(name="Calibri", size=9)
        c1.fill = row_fill
        c1.border = thin_border()

        c2 = ws.cell(row=r, column=2, value=data[0])
        c2.alignment = AL_L
        c2.font = Font(name="Calibri", size=9, bold=(data[3] in ["Labour", "Machinery"]))
        c2.fill = row_fill
        c2.border = thin_border()

        c3 = ws.cell(row=r, column=3, value=data[1])
        c3.alignment = AL_L
        c3.font = Font(name="Calibri", size=8.5)
        c3.fill = row_fill
        c3.border = thin_border()

        c4 = ws.cell(row=r, column=4, value=data[2])
        c4.alignment = AL_L
        c4.font = Font(name="Calibri", size=8.5)
        c4.fill = row_fill
        c4.border = thin_border()

        c5 = ws.cell(row=r, column=5, value=data[3])
        c5.alignment = AL_C
        c5.font = Font(name="Calibri", size=9)
        c5.fill = row_fill
        c5.border = thin_border()

        c6 = ws.cell(row=r, column=6, value=data[4])
        c6.alignment = AL_L
        c6.font = Font(name="Calibri", size=8.5)
        c6.fill = row_fill
        c6.border = thin_border()

        c7 = ws.cell(row=r, column=7, value=data[5])
        c7.alignment = AL_C
        c7.font = Font(name="Calibri", size=9)
        c7.fill = row_fill
        c7.border = thin_border()

        ws.row_dimensions[r].height = 22
        r += 1

    ws.row_dimensions[r].height = 8
    r += 1
    return r

def build_wood_pvc_sheet(ws):
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True

    # Sheet title spanning A1:G1
    ws.merge_cells("A1:G1")
    t1 = ws.cell(row=1, column=1, value="Sub-Head 9.0 — WOOD AND PVC WORK  |  First-Principles Resource, Work & Gang Analysis")
    t1.font = Font(name="Calibri", size=13, bold=True, color=C_WHITE)
    t1.fill = fill(C_TITLE_BG)
    t1.alignment = AL_C
    ws.row_dimensions[1].height = 28
    for col in range(1, 8):
        ws.cell(row=1, column=col).border = thin_border()

    # Subtitle spanning A2:G2
    ws.merge_cells("A2:G2")
    t2 = ws.cell(row=2, column=1, value="Evidence: CPWD DAR 2019 Vol 1  |  Format: Item Code • Labour / Machine / Material • Work done • Condition / When used • Category • Productivity • Standard Batch Quantity")
    t2.font = Font(name="Calibri", size=9, italic=True, color="333333")
    t2.fill = fill(C_SUBTITLE_BG)
    t2.alignment = AL_L
    ws.row_dimensions[2].height = 18
    for col in range(1, 8):
        ws.cell(row=2, column=col).border = thin_border()

    # Spacing row
    ws.row_dimensions[3].height = 8

    current_row = 4
    for item in WOOD_PVC_ITEMS:
        current_row = write_item_block(ws, item, current_row)

    # Column widths for 7 columns (A to G)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 54
    ws.column_dimensions["D"].width = 52
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 28
    ws.column_dimensions["G"].width = 16

    ws.sheet_view.showGridLines = True
    ws.freeze_panes = "A4"

def rebuild_wood_pvc_in_workbook(file_path):
    print(f"Opening {file_path} ...")
    wb = openpyxl.load_workbook(file_path)
    if "09_Wood_and_PVC_Work" not in wb.sheetnames:
        raise ValueError(f"'09_Wood_and_PVC_Work' sheet not found in {file_path}")

    pos = wb.sheetnames.index("09_Wood_and_PVC_Work")
    del wb["09_Wood_and_PVC_Work"]
    ws = wb.create_sheet("09_Wood_and_PVC_Work", pos)
    ws.sheet_properties.tabColor = "2E75B6"
    build_wood_pvc_sheet(ws)

    temp_path = file_path.replace(".xlsx", "_TMP_WOOD_PVC.xlsx")
    wb.save(temp_path)
    wb.close()
    os.replace(temp_path, file_path)
    print(f"Successfully updated {file_path} -> sheet '09_Wood_and_PVC_Work' ({len(WOOD_PVC_ITEMS)} items).")

def main():
    repo_root = Path(__file__).resolve().parents[1]
    main_wb = repo_root / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"
    output_wb = repo_root / "outputs" / "earthwork-custom-rate-composer" / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"

    if main_wb.exists():
        rebuild_wood_pvc_in_workbook(str(main_wb))
    if output_wb.exists():
        rebuild_wood_pvc_in_workbook(str(output_wb))

if __name__ == "__main__":
    main()
