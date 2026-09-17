# -*- coding: utf-8 -*-
"""
scripts.keywords.report
=======================
Reporting and serialization module for DSR keyword analysis results.
Outputs:
  - Formatted multi-tab Excel workbook with KPI Summary and auto-filters.
  - Machine-readable reference JSON for downstream trade builders and rate estimators.
  - Formatted terminal tables and summaries.
"""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import List, Dict, Any

from scripts.keywords.engine import ScheduleItem


# ─────────────────────────────────────────────────────────────────────────────
# 1. EXCEL EXPORTER
# ─────────────────────────────────────────────────────────────────────────────

def write_excel(items: List[ScheduleItem], out_path: str) -> None:
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  [WARN] openpyxl not available – skipping Excel output.")
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Color definitions (Hex)
    HDR_FILL = PatternFill("solid", fgColor="1F4E79")
    UNIQUE_FILL = PatternFill("solid", fgColor="E2EFDA")
    STRUC_FILL = PatternFill("solid", fgColor="FCE4D6")
    MORE_FILL = PatternFill("solid", fgColor="FFF2CC")
    ALT_FILL = PatternFill("solid", fgColor="F9F9F9")

    HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
    BODY_FONT = Font(size=9)
    WRAP = Alignment(wrap_text=True, vertical="top")
    CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
    thin = Side(style="thin", color="D0D0D0")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

    COLUMNS = [
        ("Code", 11),
        ("Chapter", 18),
        ("Parent Description", 45),
        ("Child Qualifier", 28),
        ("Unit", 8),
        ("Rate (₹)", 10),
        ("# Kw", 6),
        ("Global Identifier", 32),
        ("Chapter Identifier", 25),
        ("Status", 20),
        ("Ambiguous With", 22),
    ]

    by_sheet: Dict[str, List[ScheduleItem]] = defaultdict(list)
    for it in items:
        by_sheet[it.sheet].append(it)

    STATUS_FILL = {
        "UNIQUE": UNIQUE_FILL,
        "STRUCTURALLY_IMPOSSIBLE": STRUC_FILL,
        "NEEDS_MORE_KEYWORDS": MORE_FILL,
        "NOT_UNIQUELY_IDENTIFIABLE": STRUC_FILL,
    }

    # Generate individual chapter sheets
    for sheet_name, sheet_items in by_sheet.items():
        clean_title = sheet_name[:31].replace(":", "_").replace("/", "_")
        ws = wb.create_sheet(title=clean_title)
        ws.row_dimensions[1].height = 28

        for ci, (col_name, col_w) in enumerate(COLUMNS, 1):
            c = ws.cell(row=1, column=ci, value=col_name)
            c.fill = HDR_FILL
            c.font = HDR_FONT
            c.alignment = CENTER
            c.border = BORDER
            ws.column_dimensions[get_column_letter(ci)].width = col_w

        ws.freeze_panes = "A2"

        for ri, item in enumerate(sheet_items, 2):
            g_id = " + ".join(item.identifier) if item.identifier else "—"
            c_id = " + ".join(item.chapter_identifier) if item.chapter_identifier else "—"
            n_kw = len(item.identifier) if item.identifier else "—"
            ambig = ", ".join(item.ambiguous_with[:3]) if item.ambiguous_with else (
                ", ".join(item.partial_matches[:3]) if item.partial_matches else "—"
            )

            row_vals = [
                item.code,
                item.sheet,
                item.parent_desc or "—",
                item.child_desc,
                item.unit or "—",
                item.rate,
                n_kw,
                g_id,
                c_id,
                item.identifier_status,
                ambig,
            ]
            fill = STATUS_FILL.get(item.identifier_status, ALT_FILL)
            if ri % 2 == 0 and item.identifier_status == "UNIQUE":
                fill = ALT_FILL

            for ci, val in enumerate(row_vals, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.font = BODY_FONT
                c.alignment = WRAP
                c.border = BORDER
                c.fill = fill

            ws.row_dimensions[ri].height = 36

        ws.auto_filter.ref = ws.dimensions

    # ── Summary sheet ────────────────────────────────────────────────────────
    ws_s = wb.create_sheet(title="📊 Summary", index=0)
    ws_s.column_dimensions["A"].width = 28
    for col, w in [("B", 10), ("C", 10), ("D", 16), ("E", 16), ("F", 16)]:
        ws_s.column_dimensions[col].width = w

    hdr = ["Chapter", "Total", "Global Unique", "Chapter Unique", "Struct. Impossible", "Needs More Kw"]
    for ci, h in enumerate(hdr, 1):
        c = ws_s.cell(row=1, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = CENTER

    tot_tot = tot_g_uniq = tot_c_uniq = tot_si = tot_nm = 0
    for ri, (sname, sitems) in enumerate(by_sheet.items(), 2):
        total = len(sitems)
        g_uniq = sum(1 for x in sitems if x.identifier_status == "UNIQUE")
        c_uniq = sum(1 for x in sitems if getattr(x, "chapter_identifier_status", "") == "UNIQUE")
        si = sum(1 for x in sitems if x.identifier_status == "STRUCTURALLY_IMPOSSIBLE")
        nm = sum(1 for x in sitems if x.identifier_status in ("NEEDS_MORE_KEYWORDS", "NOT_UNIQUELY_IDENTIFIABLE"))

        tot_tot += total
        tot_g_uniq += g_uniq
        tot_c_uniq += c_uniq
        tot_si += si
        tot_nm += nm

        for ci, v in enumerate([sname, total, g_uniq, c_uniq, si, nm], 1):
            c = ws_s.cell(row=ri, column=ci, value=v)
            c.font = BODY_FONT
            c.alignment = CENTER
            if ci in (3, 4):
                c.fill = UNIQUE_FILL
            elif ci == 5 and si > 0:
                c.fill = STRUC_FILL
            elif ci == 6 and nm > 0:
                c.fill = MORE_FILL

    tr = len(by_sheet) + 2
    for ci, v in enumerate(["TOTAL", tot_tot, tot_g_uniq, tot_c_uniq, tot_si, tot_nm], 1):
        c = ws_s.cell(row=tr, column=ci, value=v)
        c.font = Font(bold=True, size=10)
        c.alignment = CENTER

    # ── Legend ───────────────────────────────────────────────────────────────
    start = tr + 3
    legends = [
        ("✅ UNIQUE", "1E8449", "Minimum keyword set found – uniquely identifies this item"),
        ("🔴 STRUCT. IMPOSSIBLE", "C0392B", "Item's full keyword set is a subset of another item's keywords (ambiguous without negation/context)"),
        ("🟡 NEEDS MORE KEYWORDS", "D68910", "No unique combo found within standard word limit; search depth expanded"),
    ]
    ws_s.cell(row=start - 1, column=1, value="Legend").font = Font(bold=True)
    for i, (label, col, note) in enumerate(legends):
        ws_s.cell(row=start + i, column=1, value=label).font = Font(bold=True, color=col)
        ws_s.cell(row=start + i, column=2, value=note)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"  Excel Analysis -> {out_path}")


# ─────────────────────────────────────────────────────────────────────────────
# 2. JSON EXPORTER
# ─────────────────────────────────────────────────────────────────────────────

def write_json(items: List[ScheduleItem], out_path: str) -> None:
    records: List[Dict[str, Any]] = []
    lookup_dict: Dict[str, Dict[str, Any]] = {}

    for item in items:
        rec = {
            "code": item.code,
            "sheet": item.sheet,
            "parent_code": item.parent_code,
            "parent_desc": item.parent_desc,
            "child_desc": item.child_desc,
            "complete_desc": item.complete_desc,
            "unit": item.unit,
            "rate": item.rate,
            "extracted_keywords": item.keywords,
            "global_identifier": list(item.identifier) if item.identifier else None,
            "global_n_keywords": len(item.identifier) if item.identifier else None,
            "global_status": item.identifier_status,
            "chapter_identifier": list(item.chapter_identifier) if item.chapter_identifier else None,
            "chapter_status": item.chapter_identifier_status,
            "best_partial": list(item.best_partial) if item.best_partial else [],
            "ambiguous_with": item.ambiguous_with or item.partial_matches,
        }
        records.append(rec)
        lookup_dict[item.code] = {
            "sheet": item.sheet,
            "global_identifier": rec["global_identifier"],
            "chapter_identifier": rec["chapter_identifier"],
            "status": item.identifier_status,
            "unit": item.unit,
            "rate": item.rate,
        }

    output_payload = {
        "metadata": {
            "total_items": len(items),
            "global_unique_items": sum(1 for x in items if x.identifier_status == "UNIQUE"),
            "chapter_unique_items": sum(1 for x in items if getattr(x, "chapter_identifier_status", "") == "UNIQUE"),
        },
        "items": records,
        "quick_lookup": lookup_dict,
    }

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(output_payload, f, indent=2, ensure_ascii=False)
    print(f"  Reference JSON -> {out_path}")


# ─────────────────────────────────────────────────────────────────────────────
# 3. CONSOLE OUTPUT
# ─────────────────────────────────────────────────────────────────────────────

STATUS_TAG = {
    "UNIQUE": "[UNIQUE] ",
    "STRUCTURALLY_IMPOSSIBLE": "[IMPOSS] ",
    "NEEDS_MORE_KEYWORDS": "[NEED_KW]",
    "NOT_UNIQUELY_IDENTIFIABLE": "[NOT_UNI]",
}


def print_sample(items: List[ScheduleItem], n: int = 30) -> None:
    print("\n" + "=" * 105)
    print(f"{'CODE':<12} {'STATUS':<12} {'GLOBAL IDENTIFIER':<38} {'CHAPTER IDENTIFIER':<25} SHEET")
    print("-" * 105)
    for item in items[:n]:
        tag = STATUS_TAG.get(item.identifier_status, "[?]")
        g_kw = " + ".join(item.identifier) if item.identifier else "—"
        c_kw = " + ".join(item.chapter_identifier) if item.chapter_identifier else "—"
        print(f"{item.code:<12} {tag:<12} {g_kw:<38} {c_kw:<25} {item.sheet}")
    if len(items) > n:
        print(f"  ... and {len(items) - n} more items.")
    print("=" * 105)


def print_summary(items: List[ScheduleItem]) -> None:
    total = len(items)
    g_unique = sum(1 for x in items if x.identifier_status == "UNIQUE")
    c_unique = sum(1 for x in items if getattr(x, "chapter_identifier_status", "") == "UNIQUE")
    si = sum(1 for x in items if x.identifier_status == "STRUCTURALLY_IMPOSSIBLE")
    nm = sum(1 for x in items if x.identifier_status in ("NEEDS_MORE_KEYWORDS", "NOT_UNIQUELY_IDENTIFIABLE"))
    print(f"\n{'-' * 55}")
    print(f"  TOTAL SCHEDULE ITEMS      : {total}")
    print(f"  [OK] GLOBAL UNIQUE        : {g_unique} ({100*g_unique/total:.1f}%)")
    print(f"  [OK] CHAPTER-LOCAL UNIQUE : {c_unique} ({100*c_unique/total:.1f}%)")
    print(f"  [!]  STRUCTURALLY IMPOSS. : {si} ({100*si/total:.1f}%)")
    print(f"  [?]  NEEDS MORE KEYWORDS  : {nm} ({100*nm/total:.1f}%)")
    print(f"{'-' * 55}\n")
