# -*- coding: utf-8 -*-
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def get_workbook_styles():
    return {
        # Fonts
        'font_title': Font(name='Segoe UI', size=13, bold=True, color='FFFFFF'),
        'font_header': Font(name='Segoe UI', size=10, bold=True, color='FFFFFF'),
        'font_section': Font(name='Segoe UI', size=11, bold=True, color='1E3A8A'),
        'font_regular': Font(name='Segoe UI', size=10, color='0F172A'),
        'font_bold': Font(name='Segoe UI', size=10, bold=True, color='0F172A'),
        'font_note': Font(name='Segoe UI', size=9, italic=True, color='475569'),
        'font_legend': Font(name='Segoe UI', size=9, bold=True, color='1E293B'),
        'font_result': Font(name='Segoe UI', size=11, bold=True, color='065F46'),
        'font_say': Font(name='Segoe UI', size=13, bold=True, color='1E3A8A'),
        'font_white_bold': Font(name='Segoe UI', size=10, bold=True, color='FFFFFF'),
        
        # Fills
        'fill_title': PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid'),    # Deep Navy
        'fill_header': PatternFill(start_color='2B4C7E', end_color='2B4C7E', fill_type='solid'),   # Slate Blue
        'fill_section': PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid'), # Cool Gray
        'fill_input': PatternFill(start_color='FEF9C3', end_color='FEF9C3', fill_type='solid'),     # Soft Yellow (User Input)
        'fill_lookup': PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid'),   # Pale Blue (Reference / Lookup)
        'fill_calc': PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid'),     # Clean White (Formula Calc)
        'fill_subtotal': PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid'), # Very Light Slate (Subtotal)
        'fill_result': PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid'),   # Soft Mint Green (Result)
        'fill_say': PatternFill(start_color='FEF08A', end_color='FEF08A', fill_type='solid'),      # Prominent Gold Yellow (Say Rate)
        'fill_note': PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid'),     # Off-white panel
        
        # Borders
        'border_thin': Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        ),
        'border_header': Border(
            left=Side(style='thin', color='475569'),
            right=Side(style='thin', color='475569'),
            top=Side(style='thin', color='475569'),
            bottom=Side(style='thin', color='475569')
        ),
        'border_double_bottom': Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='double', color='1E3A8A')
        ),
        
        # Alignments
        'align_left': Alignment(horizontal='left', vertical='center'),
        'align_right': Alignment(horizontal='right', vertical='center'),
        'align_center': Alignment(horizontal='center', vertical='center'),
        'align_wrap': Alignment(horizontal='left', vertical='center', wrap_text=True),
        
        # Number Formats
        'fmt_currency': '₹ #,##0.00',
        'fmt_qty': '#,##0.000',
        'fmt_percent': '0.00%',
        'fmt_integer': '#,##0'
    }
