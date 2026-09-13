"""Render the visible Custom Rate Composer above the earthwork schedule.

The panel deliberately separates a controlled, student-facing scope selection
from the CPWD schedule below.  It is a workbook view of the resolver's current
selection; a later task will attach the resolved items to resource schedules
and calculations.
"""

from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from earthwork_composer_catalogue import BASE_WORK_FAMILIES, apply_difficult_condition_reference
from earthwork_composer_resolver import resolve_selected_keywords


COMPOSER_PANEL_ROWS = 28
COMPOSER_FIRST_SCHEDULE_ROW = COMPOSER_PANEL_ROWS + 1

_DARK_BLUE = "1F4E79"
_MID_BLUE = "2E75B6"
_LIGHT_BLUE = "D9EAF7"
_INPUT_AMBER = "FFF2CC"
_LIGHT_GREEN = "E2F0D9"
_LIGHT_RED = "FCE4D6"
_GREY = "F2F2F2"
_THIN_GREY = Side(style="thin", color="A6A6A6")
_BORDER = Border(left=_THIN_GREY, right=_THIN_GREY, top=_THIN_GREY, bottom=_THIN_GREY)

_CONTROLLED_KEYWORDS = (
    "Surface excavation",
    "All kinds of soil",
    "Rough excavation + banking",
    "Banking excavated earth",
    "No power roller",
    "No watering",
    "Clearing grass",
    "Water / liquid mud",
    "Foul position",
)

_DEFAULT_SELECTION = (
    "Banking excavated earth",
    "All kinds of soil",
    "No power roller",
    "No watering",
)

_RESOLVER_ALIASES = {
    "Rough excavation + banking": "rough excavation and banking",
    "Water / liquid mud": "water or liquid mud",
}


def build_custom_rate_composer_output(source_path: str | Path, output_path: str | Path) -> Path:
    """Create the safe production export used by the Custom Rate Composer.

    The source workbook remains untouched.  The export first corrects the two
    official difficult-condition reference rows, then inserts the visible
    composer panel above the existing support schedule.
    """
    source = Path(source_path)
    output = Path(output_path)
    workbook = load_workbook(source, data_only=False)
    sheet = workbook["02_support_earth_work"]
    apply_difficult_condition_reference(sheet)
    insert_custom_rate_composer_panel(sheet)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def insert_custom_rate_composer_panel(sheet) -> None:
    """Insert the composer while explicitly moving schedule-native features.

    ``openpyxl.insert_rows`` moves cell values and styles but does not translate
    formulas, merge ranges, or row dimensions.  Capture and restore those
    structures so the existing CPWD support schedule remains one intact block.
    """
    formulas = _captured_formulas(sheet)
    merges = _captured_merges(sheet)
    row_dimensions = _captured_row_dimensions(sheet)
    _unmerge_all(sheet, merges)
    sheet.insert_rows(1, COMPOSER_PANEL_ROWS)
    _translate_shifted_formulas(sheet, formulas)
    _restore_row_dimensions(sheet, row_dimensions)
    _restore_merges(sheet, merges, insertion_row=1)
    _write_panel(sheet, _DEFAULT_SELECTION)


def _captured_formulas(sheet) -> list[tuple[int, int, str]]:
    """Capture formulas because openpyxl does not update them on row insertion."""
    return [
        (cell.row, cell.column, cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]


def _translate_shifted_formulas(sheet, formulas: list[tuple[int, int, str]]) -> None:
    """Translate only formulas moving with the existing schedule block."""
    for original_row, column, formula in formulas:
        origin = sheet.cell(original_row, column).coordinate
        destination = sheet.cell(original_row + COMPOSER_PANEL_ROWS, column).coordinate
        sheet[destination].value = Translator(formula, origin=origin).translate_formula(destination)


def _captured_merges(sheet) -> tuple[tuple[int, int, int, int], ...]:
    return tuple(
        (merged.min_col, merged.min_row, merged.max_col, merged.max_row)
        for merged in sheet.merged_cells.ranges
    )


def _unmerge_all(sheet, merges: tuple[tuple[int, int, int, int], ...]) -> None:
    for min_col, min_row, max_col, max_row in merges:
        sheet.unmerge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)


def _restore_merges(sheet, merges: tuple[tuple[int, int, int, int], ...], insertion_row: int) -> None:
    for min_col, min_row, max_col, max_row in merges:
        if min_row >= insertion_row:
            min_row += COMPOSER_PANEL_ROWS
            max_row += COMPOSER_PANEL_ROWS
        elif max_row >= insertion_row:
            max_row += COMPOSER_PANEL_ROWS
        sheet.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)


def _captured_row_dimensions(sheet) -> tuple[tuple[int, object], ...]:
    return tuple((row, copy(dimension)) for row, dimension in sheet.row_dimensions.items())


def _restore_row_dimensions(sheet, dimensions: tuple[tuple[int, object], ...]) -> None:
    for row in list(sheet.row_dimensions):
        del sheet.row_dimensions[row]
    for original_row, dimension in dimensions:
        destination = original_row + COMPOSER_PANEL_ROWS
        dimension.index = destination
        sheet.row_dimensions[destination] = dimension


def _write_panel(sheet, selected_keywords: tuple[str, ...]) -> None:
    result = resolve_selected_keywords(_resolver_keywords(selected_keywords))
    _merge(sheet, "A1:H1", "Custom Rate Composer — Earth Work", _DARK_BLUE, 13, True, "FFFFFF")
    _merge(
        sheet,
        "A2:H2",
        "Choose controlled scope keywords in order. The composer identifies the base work and linked CPWD deductions, additions or conditional extras. Resource-based pricing is shown separately below.",
        _LIGHT_BLUE,
        9,
        False,
        "1F1F1F",
    )
    _merge(sheet, "A4:H4", "1. Select the work scope", _MID_BLUE, 10, True, "FFFFFF")
    _merge(
        sheet,
        "A5:H5",
        "Use the drop-down choices only. Select one base work and then add only the operations or conditions that are genuinely part of the amended description.",
        _GREY,
        9,
        False,
        "1F1F1F",
    )

    for row, slot in zip(range(8, 15), range(1, 8), strict=True):
        sheet.cell(row, 1).value = f"Keyword {slot}"
        sheet.cell(row, 1).font = Font(bold=True, size=9)
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=_LIGHT_BLUE)
        sheet.cell(row, 1).border = _BORDER
        sheet.cell(row, 2).value = selected_keywords[slot - 1] if slot <= len(selected_keywords) else None
        sheet.cell(row, 2).fill = PatternFill("solid", fgColor=_INPUT_AMBER)
        sheet.cell(row, 2).border = _BORDER
        sheet.cell(row, 2).font = Font(size=9)
        sheet.cell(row, 2).alignment = Alignment(horizontal="left", vertical="center")
    _merge(
        sheet,
        "C8:H14",
        f"Base-work taxonomy visible in this CPWD Earth Work support sheet: {'; '.join(BASE_WORK_FAMILIES)}. Controlled selections keep a custom description auditable: start with one supported base work, then select omissions, separately measured additions, or difficult conditions. Productivity and gang norms are intentionally not created in this panel; they are later extension points for the resource schedule.",
        "FFFFFF",
        9,
        False,
        "595959",
    )
    sheet["C8"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    _add_keyword_validation(sheet)

    _merge(sheet, "A16:H16", "2. Current resolved composition", _MID_BLUE, 10, True, "FFFFFF")
    _output_row(sheet, 17, "Selected base", _component_text(result["base"]), _LIGHT_GREEN)
    _output_row(sheet, 18, "Deductions", _component_list(result["deductions"], "None selected"), _LIGHT_GREEN)
    _output_row(sheet, 19, "Additions", _component_list(result["additions"], "None selected"), _LIGHT_GREEN)
    _output_row(sheet, 20, "Conditional extras", _conditional_text(result["conditional_extras"]), _LIGHT_GREEN)
    _output_row(sheet, 21, "Overlap / measurement warning", _warnings_text(result), _LIGHT_RED)

    _merge(sheet, "A23:H23", "3. Rate position", _MID_BLUE, 10, True, "FFFFFF")
    _output_row(
        sheet,
        25,
        "CPWD benchmark",
        "Use the selected base item as a benchmark only. It is not the custom rate when scope has been amended.",
        _LIGHT_BLUE,
    )
    _output_row(
        sheet,
        26,
        "Custom calculated rate",
        "Not yet calculated here. The next step creates a resource schedule from the resolved scope, then calculates the custom unit rate from first principles.",
        _INPUT_AMBER,
    )
    _merge(
        sheet,
        "A28:H28",
        "Teaching note: Item 2.24 is a percentage extra only on the qualifying quantity and measured depth. It is not a flat resource item and must not be applied to the full work quantity by default.",
        _GREY,
        8,
        False,
        "595959",
    )

    _set_panel_dimensions(sheet)


def _resolver_keywords(selected_keywords: tuple[str, ...]) -> list[str]:
    return [_RESOLVER_ALIASES.get(keyword, keyword) for keyword in selected_keywords]


def _add_keyword_validation(sheet) -> None:
    literal_list = '"' + ",".join(_CONTROLLED_KEYWORDS) + '"'
    validation = DataValidation(type="list", formula1=literal_list, allow_blank=True)
    validation.error = "Choose a controlled earthwork keyword from the list."
    validation.errorTitle = "Controlled keyword required"
    validation.prompt = "Select the base work first, followed by any scope modifiers."
    validation.promptTitle = "Custom Rate Composer"
    sheet.add_data_validation(validation)
    validation.add("B8:B14")


def _merge(sheet, cell_range: str, value: str, colour: str, size: int, bold: bool, font_colour: str) -> None:
    sheet.merge_cells(cell_range)
    cell = sheet[cell_range.split(":", 1)[0]]
    cell.value = value
    cell.font = Font(bold=bold, size=size, color=font_colour)
    cell.fill = PatternFill("solid", fgColor=colour)
    cell.border = _BORDER
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _output_row(sheet, row: int, label: str, value: str, colour: str) -> None:
    sheet.cell(row, 1).value = label
    sheet.cell(row, 1).font = Font(bold=True, size=9)
    sheet.cell(row, 1).fill = PatternFill("solid", fgColor=colour)
    sheet.cell(row, 1).border = _BORDER
    sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    cell = sheet.cell(row, 2)
    cell.value = value
    cell.fill = PatternFill("solid", fgColor=colour)
    cell.font = Font(size=9)
    cell.border = _BORDER
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _component_text(component: dict) -> str:
    return f"Item {component['item_code']} — {component['description']}"


def _component_list(components: list[dict], empty_text: str) -> str:
    if not components:
        return empty_text
    return "  |  ".join(_component_text(component) for component in components)


def _conditional_text(components: list[dict]) -> str:
    if not components:
        return "None selected"
    return "  |  ".join(
        f"Item {component['item_code']} — {component['percent']}% extra only on the qualifying quantity of base item {component['qualifying_base_item']}; {component['measurement_basis']}."
        for component in components
    )


def _warnings_text(result: dict) -> str:
    messages = [*result["overlap_warnings"], *result["incompatibilities"]]
    return "  |  ".join(messages) if messages else "No overlap or measurement warning for the current selection."


def _set_panel_dimensions(sheet) -> None:
    widths = {"A": 28, "B": 46, "C": 14, "D": 14, "E": 14, "F": 14, "G": 14, "H": 14}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = max(sheet.column_dimensions[column].width or 0, width)
    for row in (2, 5, 8, 17, 18, 19, 20, 21, 25, 26, 28):
        sheet.row_dimensions[row].height = 34
    sheet.row_dimensions[8].height = 82
