"""verify_public_conversion_workbook.py
Zero-tolerance, exhaustive, re-runnable verification for the public
PDF-conversion workbooks (CPWD_DAR_2019_Vol{1,2}_PDF_Conversion.xlsx).

Checks EVERY item in EVERY sheet (no sampling) against the exact same
source data generate_public_conversion_workbook.py used: item presence
(no drops/duplicates), resource-row fields, Amount formulas, Resource
Total SUM ranges, SAY/percentage values, and structural label integrity.
Also scans the whole workbook for Excel error tokens.

Exits 0 if clean, 1 if ANY discrepancy is found (prints every one, no cap
unless --limit is given). Run this after every regeneration of either
workbook — do not rely on manual/agent spot-checks alone going forward.

Usage:
    python scripts/verify_public_conversion_workbook.py            # both volumes
    python scripts/verify_public_conversion_workbook.py --vol 1
    python scripts/verify_public_conversion_workbook.py --limit 50
"""
import sys, argparse
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl
from scripts.paths import PROJECT_ROOT
from scripts.generate_public_conversion_workbook import (
    VOL1_SHEETS, VOL2_SHEETS, _load_chapter_items, _load_ch02_items,
)

PROJECT_ROOT = Path(PROJECT_ROOT)

ERR_TOKENS = {"#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!"}


def _close(a, b, tol=0.01) -> bool:
    if a is None or b is None:
        return a == b
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return a == b


def verify_sheet(sheet_name: str, ws, source_items: list, errors: list) -> None:
    source_by_code: dict[str, list] = {}
    for it in source_items:
        source_by_code.setdefault(it["code"], []).append(it)

    r = 4
    maxr = ws.max_row
    while r <= maxr:
        code = ws.cell(row=r, column=1).value
        desc = ws.cell(row=r, column=2).value
        if code is None or desc is None:
            r += 1
            continue

        header_row = r
        basis_row = r + 1
        basis_val = ws.cell(row=basis_row, column=1).value
        if not basis_val or not str(basis_val).startswith("Rate basis:"):
            errors.append(f"[{sheet_name}] {code}: missing/garbled 'Rate basis:' note at row {basis_row}")

        bucket = source_by_code.get(code)
        if not bucket:
            errors.append(f"[{sheet_name}] {code}: present in workbook (row {header_row}) but not in source data")
            r = basis_row + 2
            continue
        src = bucket.pop(0)

        nxt_row = basis_row + 1
        nxt_val = ws.cell(row=nxt_row, column=1).value

        if src.get("is_pct"):
            if nxt_val != "Percentage-based item (no resource breakdown in the source document).":
                errors.append(f"[{sheet_name}] {code}: expected percentage note at row {nxt_row}, got {nxt_val!r}")
            pv_cell = ws.cell(row=nxt_row, column=6).value
            expected_pct = src.get("pct_value")
            if expected_pct is None:
                if pv_cell != "N/A (not given in source)":
                    errors.append(f"[{sheet_name}] {code}: expected N/A percentage text, got {pv_cell!r}")
            elif not (isinstance(pv_cell, (int, float)) and _close(pv_cell, expected_pct / 100.0, 1e-6)):
                errors.append(f"[{sheet_name}] {code}: percentage mismatch wb={pv_cell!r} expected={expected_pct/100.0}")
            r = nxt_row + 2
            continue

        if nxt_val != "Code":
            errors.append(f"[{sheet_name}] {code}: expected column-header row at {nxt_row}, got {nxt_val!r}")
            r = nxt_row + 1
            continue

        res_start = nxt_row + 1
        rr = res_start
        for i, sres in enumerate(src["resources"]):
            row_code = ws.cell(row=rr, column=1).value
            row_desc = ws.cell(row=rr, column=2).value
            row_unit = ws.cell(row=rr, column=3).value
            row_qty = ws.cell(row=rr, column=4).value
            row_rate = ws.cell(row=rr, column=5).value
            row_amt = ws.cell(row=rr, column=6).value

            if (row_code or None) != (sres["code"] or None):
                errors.append(f"[{sheet_name}] {code} resource#{i} (row {rr}): code mismatch wb={row_code!r} src={sres['code']!r}")
            if row_desc != sres["desc"]:
                errors.append(f"[{sheet_name}] {code} resource#{i} (row {rr}): desc mismatch wb={row_desc!r} src={sres['desc']!r}")
            if row_unit != sres["unit"]:
                errors.append(f"[{sheet_name}] {code} resource#{i} (row {rr}): unit mismatch wb={row_unit!r} src={sres['unit']!r}")
            if not _close(row_qty, sres["qty"], 1e-6):
                errors.append(f"[{sheet_name}] {code} resource#{i} (row {rr}): qty mismatch wb={row_qty!r} src={sres['qty']!r}")
            if not _close(row_rate, sres["rate"], 1e-6):
                errors.append(f"[{sheet_name}] {code} resource#{i} (row {rr}): rate mismatch wb={row_rate!r} src={sres['rate']!r}")
            if row_amt != f"=D{rr}*E{rr}":
                errors.append(f"[{sheet_name}] {code} resource#{i} (row {rr}): bad Amount formula {row_amt!r}")
            rr += 1
        res_end = rr - 1
        n_res = res_end - res_start + 1

        if n_res > 0:
            total_label = ws.cell(row=rr, column=1).value
            total_formula = ws.cell(row=rr, column=6).value
            expected_label = "Resource Total (Material + Labour + Machinery, as printed)"
            expected_formula = f"=SUM(F{res_start}:F{res_end})"
            if total_label != expected_label:
                errors.append(f"[{sheet_name}] {code}: expected Resource Total row at {rr}, got {total_label!r}")
            if total_formula != expected_formula:
                errors.append(f"[{sheet_name}] {code}: bad Resource Total formula {total_formula!r} expected {expected_formula!r} (row {rr})")
            rr += 1
        else:
            errors.append(f"[{sheet_name}] {code}: zero resource rows for a non-percentage item")

        say_row = rr
        say_label = ws.cell(row=say_row, column=1).value
        say_val = ws.cell(row=say_row, column=6).value
        if not say_label or "SAY" not in str(say_label):
            errors.append(f"[{sheet_name}] {code}: expected SAY row at {say_row}, got {say_label!r}")
        if not _close(say_val, src.get("say"), 0.01):
            errors.append(f"[{sheet_name}] {code}: SAY mismatch wb={say_val!r} src={src.get('say')!r} (row {say_row})")

        r = say_row + 2

    for c, remaining in source_by_code.items():
        if remaining:
            errors.append(f"[{sheet_name}] {c}: {len(remaining)} source occurrence(s) never found in workbook (dropped)")


def verify_workbook(vol_num: int, sheets_def: list) -> list:
    wb_path = PROJECT_ROOT / f"CPWD_DAR_2019_Vol{vol_num}_PDF_Conversion.xlsx"
    errors: list[str] = []
    if not wb_path.exists():
        return [f"{wb_path.name}: file not found"]

    wb = openpyxl.load_workbook(wb_path, data_only=False)

    for ch, sheet_name, _title in sheets_def:
        if sheet_name not in wb.sheetnames:
            errors.append(f"[{sheet_name}] sheet missing from workbook entirely")
            continue
        source_items = _load_ch02_items() if ch == 2 else _load_chapter_items(ch)
        verify_sheet(sheet_name, wb[sheet_name], source_items, errors)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() in ERR_TOKENS:
                    errors.append(f"[{ws.title}] {cell.coordinate}: Excel error token {cell.value!r}")

    return errors


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--vol", type=int, choices=[1, 2], help="Verify one volume only")
    ap.add_argument("--limit", type=int, default=0, help="Cap printed discrepancies (0 = no cap)")
    args = ap.parse_args()

    vols = [1, 2] if not args.vol else [args.vol]
    all_ok = True

    for v in vols:
        sheets_def = VOL1_SHEETS if v == 1 else VOL2_SHEETS
        print(f"Verifying CPWD_DAR_2019_Vol{v}_PDF_Conversion.xlsx ...")
        errs = verify_workbook(v, sheets_def)
        if errs:
            all_ok = False
            print(f"  FAIL: {len(errs)} discrepancies")
            shown = errs if not args.limit else errs[: args.limit]
            for e in shown:
                print("   -", e)
            if args.limit and len(errs) > args.limit:
                print(f"   ... and {len(errs) - args.limit} more")
        else:
            print("  PASS: zero discrepancies")

    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()
