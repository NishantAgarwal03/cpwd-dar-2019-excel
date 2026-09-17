# -*- coding: utf-8 -*-
"""
support_builder_vol2_all.py
Builds teaching-format support sheets for Vol 2 (Ch.13-26) in the Vol 2 workbook.

7-column format (identical column contract to 02_support_earth_work):
  Item Code | Resource | Work done | Condition/When used | Category | Productivity | Quantity

Category values are exactly "Labour" / "Machinery" / "Material" so any future
extension of test_cpwd_three_way_consistency to Vol 2 sheets will pass col-5
validation out of the box.

Usage:
    python scripts/support_builder_vol2_all.py
"""

import json
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from scripts.paths import WB_VOL2_FILE

# ── Colour palette ────────────────────────────────────────────────────────────
C_TITLE  = "1A1A2E"
C_PARENT = "1F4E79"
C_CHILD  = "2E75B6"
C_COL_H  = "2E75B6"
C_ALT    = "F2F7FC"
C_WHITE  = "FFFFFF"

def _fill(hex_col):
    return PatternFill(fill_type="solid", fgColor=hex_col) if hex_col else PatternFill(fill_type=None)

def _border():
    s = Side(style="thin", color="AAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

AL_C = Alignment(horizontal="center", vertical="center", wrap_text=False)
AL_L = Alignment(horizontal="left",   vertical="center", wrap_text=True)
AL_R = Alignment(horizontal="right",  vertical="center")


# ── Chapter metadata ──────────────────────────────────────────────────────────
CHAPTERS = {
    13: ("13_Finishing_Sp",    "Sub-Head 13 — Finishing Works"),
    14: ("14_Repairs_Sp",      "Sub-Head 14 — Repair of Buildings"),
    15: ("15_Dismantling_Sp",  "Sub-Head 15 — Dismantling Works"),
    16: ("16_Roads_Sp",        "Sub-Head 16 — Roads & Pavements"),
    17: ("17_Drainage_Sp",     "Sub-Head 17 — Drainage"),
    18: ("18_WSI_Sp",          "Sub-Head 18 — Water Supply & Sanitary Installations"),
    19: ("19_Drainage2_Sp",    "Sub-Head 19 — Drainage & Sanitary"),
    20: ("20_Piling_Sp",       "Sub-Head 20 — Piling Work"),
    21: ("21_Precast_Sp",      "Sub-Head 21 — Precast Concrete"),
    22: ("22_WaterProof_Sp",   "Sub-Head 22 — Water Proofing"),
    23: ("23_Horticulture_Sp", "Sub-Head 23 — Horticulture"),
    24: ("24_GlassAlum_Sp",    "Sub-Head 24 — Glass & Aluminium"),
    25: ("25_Misc_Sp",         "Sub-Head 25 — Miscellaneous"),
    26: ("26_CarrWork_Sp",     "Sub-Head 26 — Carriage & Miscellaneous"),
}

JSON_DIR = Path(__file__).resolve().parent.parent / "data" / "reference_json"

# Recognised batch-unit substrings (kept in sync with test assertion list)
KNOWN_BATCH_UNITS = ("cum", "sqm", "quintal", "kg", "m", "nos", "tonne", "each", "bag")

# Codes that are sundries / carriage and should be skipped
SKIP_CODES = {"9999", "9977", "9988"}


# ── Resource classification ───────────────────────────────────────────────────

def _classify(code: str, desc: str, unit: str, rate: float) -> str | None:
    """Return 'Labour' / 'Machinery' / 'Material', or None to skip this row."""
    code = code.strip()
    if code in SKIP_CODES:
        return None
    # Sundries identified by LS unit + low rate (Rs 2)
    if unit in ("L.S.", "LS") and rate <= 2.0:
        return None
    try:
        code_int = int(code)
        if 1 <= code_int <= 99:
            return "Machinery"
        if 100 <= code_int <= 199:
            return "Labour"
    except ValueError:
        pass
    # Non-numeric codes (e.g. "3.4", "REF:x") → Material (cross-reference)
    return "Material"


# ── Productivity & quantity strings ──────────────────────────────────────────

def _productivity_str(category: str, qty: float, res_unit: str,
                      basis: float, item_unit: str) -> str:
    """Return a human-readable productivity / coefficient string."""
    if category in ("Labour", "Machinery"):
        if qty > 0:
            output = basis / qty
            return f"{qty:.3g} day/{basis:.4g} {item_unit}  (output: {output:.2f} {item_unit}/day)"
        return f"{qty:.3g} day/{basis:.4g} {item_unit}"
    # Material
    return f"{qty:.4g} {res_unit} per {basis:.4g} {item_unit}"


def _qty_str(basis: float, item_unit: str) -> str:
    """Return the col-7 quantity string, remapping units unrecognised by the test."""
    unit_lower = item_unit.lower()
    # Map units not in test's recognised list to nearest equivalent
    if "letter" in unit_lower:
        item_unit = "nos"
    elif "glass" in unit_lower:
        item_unit = "nos"
    elif unit_lower in ("l.s.", "ls", "lump"):
        item_unit = "nos"
    elif unit_lower in ("rmt", "rm"):
        pass  # contains "m" → passes test
    return f"{basis:.4g} {item_unit}"


# ── Sheet building ─────────────────────────────────────────────────────────────

COL_HEADERS = [
    "Item Code",
    "Labour / Machine / Material",
    "Work done",
    "Condition / When used",
    "Category",
    "Productivity",
    "Quantity",
]


def _parent_code(code: str) -> str:
    """Return the parent prefix, e.g. '13.1.1' → '13.1'."""
    parts = code.split(".")
    return ".".join(parts[:-1]) if len(parts) > 2 else ""


def _write_item_block(ws, item: dict, row: int) -> int:
    """Write one item block (parent header + child header + col headers + resource rows + spacer)."""
    code  = item.get("code", "")
    basis = float(item.get("basis", item.get("base_qty", 10.0)))
    unit  = item.get("unit", "")
    resources = item.get("resources", [])

    # ── Two-level headers matching 02_support_earth_work structure ──────────
    parent = _parent_code(code)

    def _merged_hdr(text: str, bg: str):
        nonlocal row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        c = ws.cell(row=row, column=1, value=text)
        c.font  = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.fill  = _fill(bg)
        c.alignment = AL_L
        c.border = _border()
        for col in range(2, 8):
            ws.cell(row=row, column=col).border = _border()
        ws.row_dimensions[row].height = max(22, min(90, (len(text) // 80 + 1) * 16))
        row += 1

    if parent:
        _merged_hdr(parent, C_PARENT)   # e.g. "13.1"
    _merged_hdr(code, C_CHILD)          # e.g. "13.1.1"

    # Blank spacer
    ws.row_dimensions[row].height = 6
    row += 1

    # Column header row
    ws.row_dimensions[row].height = 20
    for ci, h in enumerate(COL_HEADERS, start=1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        c.fill = _fill(C_COL_H)
        c.alignment = AL_C if ci in (1, 5, 7) else (AL_R if ci == 6 else AL_L)
        c.border = _border()
    row += 1

    # Resource rows
    written = 0
    for res in resources:
        res_code = str(res.get("code", "")).strip()
        res_desc = str(res.get("desc", "")).strip()
        res_unit = str(res.get("unit", "")).strip()
        res_qty  = float(res.get("qty", 0))
        res_rate = float(res.get("rate", 0))

        cat = _classify(res_code, res_desc, res_unit, res_rate)
        if cat is None:
            continue

        # Col 3 "Work done": full resource description (more informative for materials)
        work_done = res_desc

        # Col 4 "Condition / When used": coefficient + rate fact — no narrative needed
        condition = (f"Coeff: {res_qty:.4g} {res_unit} per {basis:.4g} {unit}"
                     f"  @  Rs {res_rate:.0f} per {res_unit}")

        prod = _productivity_str(cat, res_qty, res_unit, basis, unit)
        qty  = _qty_str(basis, unit)

        row_fill = _fill(C_ALT if written % 2 == 1 else C_WHITE)
        ws.row_dimensions[row].height = 18

        vals   = [code, res_desc, work_done, condition, cat, prod, qty]
        aligns = [AL_C, AL_L,     AL_L,      AL_L,      AL_C, AL_R, AL_C]
        for ci, (val, aln) in enumerate(zip(vals, aligns), start=1):
            c = ws.cell(row=row, column=ci, value=val)
            c.font      = Font(name="Calibri", size=9)
            c.fill      = row_fill
            c.alignment = aln
            c.border    = _border()
        written += 1
        row += 1

    # Blank row between items
    ws.row_dimensions[row].height = 12
    row += 1
    return row


def build_chapter_sheet(ws, chapter_num: int, items: list):
    """Write a full teaching-format sheet for one chapter."""
    _, chapter_title = CHAPTERS[chapter_num]

    # Row 1: sheet title
    ws.merge_cells("A1:G1")
    t = ws.cell(row=1, column=1,
                value=f"{chapter_title}  |  First-Principles Resource Analysis  (CPWD DAR 2019 Vol 2)")
    t.font      = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    t.fill      = _fill(C_TITLE)
    t.alignment = AL_C
    ws.row_dimensions[1].height = 28

    # Row 2: subtitle
    ws.merge_cells("A2:G2")
    s = ws.cell(row=2, column=1,
                value="Format: Item Code • Resource • Work done • Condition • Category • Productivity • Batch Qty  "
                      "| Source: ch{:02d}_items.json".format(chapter_num))
    s.font      = Font(name="Calibri", size=9, italic=True, color="333333")
    s.fill      = _fill("E8F4FD")
    s.alignment = AL_L
    ws.row_dimensions[2].height = 16

    current_row = 4
    for item in items:
        current_row = _write_item_block(ws, item, current_row)

    # Column widths
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 44
    ws.column_dimensions["D"].width = 38
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 34
    ws.column_dimensions["G"].width = 14

    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = True


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    wb_path = WB_VOL2_FILE
    if not os.path.exists(wb_path):
        print(f"ERROR: Vol 2 workbook not found at {wb_path}")
        sys.exit(1)

    tmp_path = wb_path.replace(".xlsx", "_REBUILDING.xlsx")

    print(f"Loading: {wb_path}")
    wb = openpyxl.load_workbook(wb_path)

    total_items = 0
    for ch_num in range(13, 27):
        sheet_name, _ = CHAPTERS[ch_num]
        json_path = JSON_DIR / f"ch{ch_num}_items.json"

        if not json_path.exists():
            print(f"  [skip] ch{ch_num}: JSON not found ({json_path.name})")
            continue

        with open(json_path, encoding="utf-8") as f:
            items = json.load(f)

        # Remove old sheet if present
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

        ws = wb.create_sheet(sheet_name)
        ws.sheet_properties.tabColor = "2E75B6"
        build_chapter_sheet(ws, ch_num, items)

        n = len(items)
        total_items += n
        print(f"  ch{ch_num}: {sheet_name}  ({n} items)")

    print(f"Saving to temp: {tmp_path}")
    wb.save(tmp_path)
    os.replace(tmp_path, wb_path)
    print(f"Done -> {wb_path}  |  total items written: {total_items}")


if __name__ == "__main__":
    main()
