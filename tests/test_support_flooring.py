"""Tests for Sub-Head 11 Flooring First-Principles Sheet."""

import unittest
import openpyxl
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"

class TestSupportFlooring(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
        cls.ws = cls.wb["11_Flooring"]

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_sheet_exists(self):
        self.assertIn("11_Flooring", self.wb.sheetnames)

    def test_title_and_columns(self):
        title = self.ws.cell(1, 1).value
        self.assertIn("FLOORING", title)
        headers = [self.ws.cell(6, col).value for col in range(1, 8)]
        expected = [
            "Item Code",
            "Labour / Machine / Material",
            "Work done",
            "Condition / When used",
            "Category",
            "Productivity",
            "Quantity"
        ]
        self.assertEqual(headers, expected)

    def test_core_items_present(self):
        found_items = set()
        for r in range(7, self.ws.max_row + 1):
            val = self.ws.cell(r, 1).value
            if val and str(val).strip().startswith("11."):
                found_items.add(str(val).strip())

        key_items = {
            "11.1", "11.3.1", "11.3.2", "11.6", "11.9.1",
            "11.23.1", "11.26.1", "11.27", "11.37", "11.41.2", "11.56"
        }
        for ki in key_items:
            self.assertIn(ki, found_items, f"Missing key item {ki} in 11_Flooring")

    def test_no_four_digit_prefixes_in_column_b(self):
        code_prefixed = []
        for r in range(7, self.ws.max_row + 1):
            val = self.ws.cell(r, 2).value
            if val and re.match(r"^\d{4}\s+", str(val)):
                code_prefixed.append((r, val))
        self.assertEqual(code_prefixed, [], f"Found 4-digit code prefixes in Col B: {code_prefixed}")

    def test_direct_materials_and_references_present(self):
        references = []
        materials = []
        for r in range(7, self.ws.max_row + 1):
            cat = self.ws.cell(r, 5).value
            item = self.ws.cell(r, 2).value
            if cat == "Reference":
                references.append(item)
            elif cat == "Material":
                materials.append(item)

        self.assertTrue(len(references) > 0, "Expected referenced base items (e.g. REF#3.9, REF#3.8)")
        self.assertTrue(len(materials) > 0, "Expected direct materials (marble, tiles, Kota stone, cement, etc.)")

if __name__ == "__main__":
    unittest.main()
