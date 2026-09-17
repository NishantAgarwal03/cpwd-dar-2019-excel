"""Tests for Sub-Head 04 Concrete Work First-Principles Sheet."""

import unittest
import openpyxl
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"

class TestSupportConcrete(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
        cls.ws = cls.wb["04_Concrete_Work"]

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_sheet_exists(self):
        self.assertIn("04_Concrete_Work", self.wb.sheetnames)

    def test_title_and_columns(self):
        title = self.ws.cell(1, 1).value
        self.assertIn("CONCRETE WORK", title)
        headers = [self.ws.cell(7, col).value for col in range(1, 8)]
        expected = [
            "Item Code",
            "Labour / Machine / Material",
            "Work done",
            "Condition / When used",
            "Category",
            "Productivity",
            "Quantity"
        ]
        self.assertEqual(headers, expected)

    def test_core_items_present(self):
        found_items = set()
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 1).value
            if val and str(val).strip().startswith("4."):
                found_items.add(str(val).strip())

        # Check key items across different groups
        key_items = {
            "4.1.2", "4.1.3", "4.1.4", "4.1.5", "4.1.6", "4.1.8", "4.1.10", "4.1.12",
            "4.2.2", "4.2.3", "4.2.5",
            "4.3.1", "4.3.2", "4.3.3",
            "4.4.1", "4.5.1", "4.6.1", "4.7.1", "4.8.1", "4.9",
            "4.10", "4.11", "4.12", "4.13", "4.14", "4.15", "4.16", "4.17", "4.18",
            "4.19.1.1", "4.19.1.2", "4.19.2.1", "4.20.1.1", "4.20.2.2"
        }
        for ki in key_items:
            self.assertIn(ki, found_items, f"Missing key item {ki} in 04_Concrete_Work")

    def test_no_four_digit_prefixes_in_column_b(self):
        code_prefixed = []
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 2).value
            if val and re.match(r"^\d{4}\s+", str(val)):
                code_prefixed.append((r, val))
        self.assertEqual(code_prefixed, [], f"Found 4-digit code prefixes in Col B: {code_prefixed}")

    def test_item_categories_and_material_presence(self):
        item_categories = {}
        for r in range(8, self.ws.max_row + 1):
            code = self.ws.cell(r, 1).value
            cat = self.ws.cell(r, 5).value
            if code and str(code).strip().startswith("4."):
                item_id = str(code).strip()
                if item_id not in item_categories:
                    item_categories[item_id] = set()
                if cat:
                    item_categories[item_id].add(str(cat).strip())

        # Check that mix items have Material, Machine/Equipment, and Labour
        mix_items = ["4.1.2", "4.1.3", "4.1.5", "4.2.2", "4.19.1.1"]
        for mi in mix_items:
            cats = item_categories.get(mi, set())
            self.assertIn("Material", cats, f"Item {mi} missing Material")
            self.assertIn("Labour", cats, f"Item {mi} missing Labour")
            self.assertTrue("Machine" in cats or "Equipment" in cats, f"Item {mi} missing Machine/Equipment")

if __name__ == "__main__":
    unittest.main()
