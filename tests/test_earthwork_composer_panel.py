"""Workbook-facing tests for the student custom-rate composer panel."""

from __future__ import annotations

import sys
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from earthwork_composer_panel import (  # noqa: E402
    COMPOSER_PANEL_ROWS,
    COMPOSER_FIRST_SCHEDULE_ROW,
    build_custom_rate_composer_output,
    insert_custom_rate_composer_panel,
)
from earthwork_composer_catalogue import BASE_WORK_FAMILIES  # noqa: E402


class EarthworkComposerPanelTests(unittest.TestCase):
    def test_catalogue_exposes_the_visible_earthwork_base_family_taxonomy(self):
        self.assertEqual(
            BASE_WORK_FAMILIES[:6],
            (
                "Excavation",
                "Banking/Embankment",
                "Filling",
                "Timbering/Shoring/Planking",
                "Site Clearance/Surface Preparation",
                "Chemical Anti-Termite Treatment",
            ),
        )

    def _schedule_sheet(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "02_support_earth_work"
        sheet["A1"] = "Existing CPWD support schedule"
        sheet["A2"] = "Item 2.1.1"
        sheet["B2"] = "Existing description"
        sheet["D2"] = "=A2*2"
        return workbook, sheet

    def test_panel_sits_above_existing_schedule_and_keeps_existing_cells(self):
        _workbook, sheet = self._schedule_sheet()

        insert_custom_rate_composer_panel(sheet)

        self.assertEqual(sheet["A1"].value, "Custom Rate Composer — Earth Work")
        self.assertEqual(sheet.cell(COMPOSER_FIRST_SCHEDULE_ROW, 1).value, "Existing CPWD support schedule")
        self.assertEqual(sheet.cell(COMPOSER_FIRST_SCHEDULE_ROW + 1, 4).value, "=A30*2")

    def test_panel_provides_controlled_keyword_slots_with_excel_validation(self):
        _workbook, sheet = self._schedule_sheet()

        insert_custom_rate_composer_panel(sheet)

        validations = list(sheet.data_validations.dataValidation)
        self.assertEqual(len(validations), 1)
        self.assertIn("B8:B14", str(validations[0].sqref))
        self.assertEqual(validations[0].type, "list")
        self.assertIn("Banking excavated earth", validations[0].formula1)
        self.assertEqual(sheet["B8"].value, "Banking excavated earth")
        self.assertEqual(sheet["B10"].value, "No power roller")

    def test_panel_explains_base_work_taxonomy_without_claiming_productivity_data(self):
        _workbook, sheet = self._schedule_sheet()

        insert_custom_rate_composer_panel(sheet)

        self.assertIn("Excavation", sheet["C8"].value)
        self.assertIn("Chemical Anti-Termite Treatment", sheet["C8"].value)
        self.assertIn("productivity", sheet["C8"].value.casefold())

    def test_panel_renders_base_deductions_and_clear_rate_distinction(self):
        _workbook, sheet = self._schedule_sheet()

        insert_custom_rate_composer_panel(sheet)

        self.assertIn("2.3.1", sheet["B17"].value)
        self.assertIn("2.4", sheet["B18"].value)
        self.assertIn("2.5", sheet["B18"].value)
        self.assertIn("CPWD benchmark", sheet["A25"].value)
        self.assertIn("Custom calculated rate", sheet["A26"].value)
        self.assertIn("470.57", sheet["B25"].value)
        self.assertIn("Custom Unit Rate:", sheet["B26"].value)
        self.assertIn("433.25", sheet["B26"].value)
        self.assertIn("Variance & gang effect", sheet["A27"].value)
        self.assertIn("-37.32", sheet["B27"].value)

    def test_saved_panel_reopens_with_merged_title_and_wrap_ready_rows(self):
        workbook, sheet = self._schedule_sheet()
        insert_custom_rate_composer_panel(sheet)
        output = ROOT / "scratch" / "test_composer_panel.xlsx"
        output.parent.mkdir(exist_ok=True)
        workbook.save(output)

        reopened = load_workbook(output, data_only=False)
        panel = reopened["02_support_earth_work"]
        self.assertIn("A1:H1", {str(range_) for range_ in panel.merged_cells.ranges})
        self.assertTrue(panel["B20"].alignment.wrap_text)
        self.assertGreaterEqual(panel.row_dimensions[20].height, 30)

    def test_actual_support_sheet_can_be_saved_and_reopened_after_panel_insert(self):
        source = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"
        if not source.exists():
            self.skipTest("Source workbook is not available in this worktree")
        workbook = load_workbook(source, data_only=False)
        sheet = workbook["02_support_earth_work"]
        original_title = sheet["A1"].value

        insert_custom_rate_composer_panel(sheet)
        output = ROOT / "scratch" / "test_actual_composer_panel.xlsx"
        workbook.save(output)

        reopened = load_workbook(output, data_only=False)
        panel = reopened["02_support_earth_work"]
        self.assertEqual(panel["A1"].value, "Custom Rate Composer — Earth Work")
        self.assertEqual(panel.cell(COMPOSER_FIRST_SCHEDULE_ROW, 1).value, original_title)
        self.assertEqual(panel["A35"].value, "Item Code")
        self.assertEqual(panel["B35"].value, "Labour / Machine / Material")
        self.assertEqual(panel["A36"].value, "2.1.1")
        self.assertEqual(panel["B36"].value, "Beldar")
        self.assertTrue(panel["B26"].alignment.wrap_text)

    def test_actual_support_sheet_preserves_every_formula_merge_and_row_group(self):
        source = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"
        if not source.exists():
            self.skipTest("Source workbook is not available in this worktree")
        workbook = load_workbook(source, data_only=False)
        sheet = workbook["02_support_earth_work"]
        original_formulas = {
            cell.coordinate: cell.value
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
        original_merges = tuple(str(range_) for range_ in sheet.merged_cells.ranges)
        original_row_dimensions = {
            row: (dimension.height, dimension.hidden, dimension.outlineLevel, dimension.collapsed)
            for row, dimension in sheet.row_dimensions.items()
            if dimension.height is not None or dimension.hidden or dimension.outlineLevel or dimension.collapsed
        }

        insert_custom_rate_composer_panel(sheet)
        output = ROOT / "scratch" / "test_preserved_composer_panel.xlsx"
        workbook.save(output)
        panel = load_workbook(output, data_only=False)["02_support_earth_work"]

        panel_merges = {str(range_) for range_ in panel.merged_cells.ranges}
        for original_merge in original_merges:
            start, end = original_merge.split(":")
            expected = f"{start[0]}{int(start[1:]) + COMPOSER_PANEL_ROWS}:{end[0]}{int(end[1:]) + COMPOSER_PANEL_ROWS}"
            self.assertIn(expected, panel_merges)
        for coordinate, formula in original_formulas.items():
            column = "".join(filter(str.isalpha, coordinate))
            row = int("".join(filter(str.isdigit, coordinate)))
            destination = f"{column}{row + COMPOSER_PANEL_ROWS}"
            self.assertNotIsInstance(panel[destination], MergedCell)
            self.assertEqual(
                panel[destination].value,
                Translator(formula, origin=coordinate).translate_formula(destination),
            )
        for row, expected_dimension in original_row_dimensions.items():
            dimension = panel.row_dimensions[row + COMPOSER_PANEL_ROWS]
            self.assertEqual(
                (dimension.height, dimension.hidden, dimension.outlineLevel, dimension.collapsed),
                expected_dimension,
            )

    def test_production_export_path_writes_a_composer_ready_workbook(self):
        source = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"
        if not source.exists():
            self.skipTest("Source workbook is not available in this worktree")
        output = ROOT / "scratch" / "test_composer_production_export.xlsx"

        build_custom_rate_composer_output(source, output)

        self.assertTrue(output.exists())
        panel = load_workbook(output, data_only=False)["02_support_earth_work"]
        self.assertEqual(panel["A1"].value, "Custom Rate Composer — Earth Work")
        self.assertIn("20%", " ".join(str(cell.value or "") for cell in panel["A"]))
        with zipfile.ZipFile(output) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet8.xml"))
        formulas = [formula.text or "" for formula in root.findall(".//{*}f")]
        self.assertTrue(all(all(ord(character) <= 127 for character in formula) for formula in formulas))
        for formula in formulas:
            for literal in _formula_literals(formula):
                self.assertLessEqual(len(literal), 255)


def _formula_literals(formula: str) -> list[str]:
    literals: list[str] = []
    index = 0
    while index < len(formula):
        if formula[index] != '"':
            index += 1
            continue
        index += 1
        literal = []
        while index < len(formula):
            if formula[index] == '"' and index + 1 < len(formula) and formula[index + 1] == '"':
                literal.append('"')
                index += 2
            elif formula[index] == '"':
                index += 1
                break
            else:
                literal.append(formula[index])
                index += 1
        literals.append("".join(literal))
    return literals


if __name__ == "__main__":
    unittest.main()
