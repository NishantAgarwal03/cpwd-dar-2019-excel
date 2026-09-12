"""Tests for the source-backed earthwork productivity derivation model."""

from __future__ import annotations

import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASELINE_WORKBOOK = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"
sys.path.insert(0, str(ROOT / "scripts"))

from support_earth_learning_edit import (  # noqa: E402
    apply_first_principles_learning,
    apply_markup_outline,
    estimated_card_lines,
    build_derivation_catalog,
    derive_resource_norm,
    export_ascii_formula_workbook,
)


def sheet8_formulas(workbook_path: Path) -> list[str]:
    """Read formula nodes directly so Excel's XML compatibility is tested."""
    with zipfile.ZipFile(workbook_path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet8.xml"))
    return [node.text or "" for node in root.findall(".//{*}f")]


def source_sheet_with_verified_norms():
    """Build the same source-table shape used in 02_Earth_Work (rows 109+)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "02_Earth_Work"

    # A:D identify the source work item; E/F identify its resource; G is output.
    # H/I are the source activity and verbatim source calculation; K is the
    # published coefficient and L is the support-resource lookup key.
    sheet.append([None] * 17)
    sheet.append([
        "Surface Excavation (≤30 cm)", "All kinds of soil",
        "Manual labor (depth ≤1.5m)", "General surface cut", 1, "0114",
        14.70588235294118,
        "Manual pickaxe and phowrah loosening.",
        "Gang: 4 Beldars × 13.60 h = 54.40 man-hrs ÷ 8.00 h = 6.800 Beldar-days per 100 sqm.",
        "source-reference", 6.8,
        "2. Surface Preparation & Earthwork|General surface cut|Manual labor (depth ≤1.5m)|1",
        "Beldar", "day", 558, "W",
        "Surface Excavation (≤30 cm)|All kinds of soil|Manual labor (depth ≤1.5m)|General surface cut|1",
    ])
    sheet.append([
        "Embankment Banking & Dressing", "All kinds of soil",
        "Manual labor (depth ≤1.5m)", "Full cycle: rough excavation + banking + roll", 5, "0003",
        1250,
        "8-10 tonne diesel roller conducting compaction passes.",
        "Plant: 1 Roller × 0.064 h ÷ 8.00 h = 0.008 roller-day per 10 cum.",
        "source-reference", 0.008,
        "2. Surface Preparation & Earthwork|Full cycle: rough excavation + banking + roll|Manual labor (depth ≤1.5m)|5",
        "Hire – Diesel Road Roller 8-10t", "day", 3000, "W",
        "Embankment Banking & Dressing|All kinds of soil|Manual labor (depth ≤1.5m)|Full cycle: rough excavation + banking + roll|5",
    ])
    return sheet


class DerivationCatalogTests(unittest.TestCase):
    def setUp(self):
        sheet = source_sheet_with_verified_norms()
        # Excel commonly stores CPWD resource code 0114 as the number 114.
        sheet["F2"] = 114
        self.catalog = build_derivation_catalog(sheet)

    def test_derives_labour_norm_from_unique_source_record(self):
        result = derive_resource_norm(
            {"code": "0114", "description": "Beldar", "coefficient": 6.8, "unit": "day"},
            {"source_key": "2. Surface Preparation & Earthwork|General surface cut|Manual labor (depth ≤1.5m)|1", "batch_quantity": 100, "batch_unit": "sqm"},
            self.catalog,
        )

        self.assertEqual(result["role"], "Beldar")
        self.assertEqual(result["gang_or_machine"], 4)
        self.assertAlmostEqual(result["task_hours"], 13.6)
        self.assertEqual(result["shift_hours"], 8)
        self.assertAlmostEqual(result["coefficient"], 6.8)
        self.assertEqual(result["unit"], "day")
        self.assertEqual(result["batch_quantity"], 100)
        self.assertEqual(result["batch_unit"], "sqm")
        self.assertEqual(result["interpretation_label"], "CPWD fixed norm / source evidence")
        self.assertIn("Gang: 4 Beldars", result["source_norm"])

    def test_derives_machine_norm_from_unique_source_record(self):
        result = derive_resource_norm(
            {"code": "0003", "description": "Hire – Diesel Road Roller 8-10t", "coefficient": 0.008, "unit": "day"},
            {"source_key": "2. Surface Preparation & Earthwork|Full cycle: rough excavation + banking + roll|Manual labor (depth ≤1.5m)|5", "batch_quantity": 10, "batch_unit": "cum"},
            self.catalog,
        )

        self.assertEqual(result["role"], "Hire – Diesel Road Roller 8-10t")
        self.assertEqual(result["gang_or_machine"], 1)
        self.assertAlmostEqual(result["task_hours"], 0.064)
        self.assertAlmostEqual(result["machine_hours"], 0.064)
        self.assertAlmostEqual(result["coefficient"], 0.008)
        self.assertEqual(result["interpretation_label"], "CPWD fixed norm / source evidence")

    def test_returns_none_when_source_mapping_is_not_unique_or_not_found(self):
        result = derive_resource_norm(
            {"code": "9999", "description": "Unmatched resource", "coefficient": 1, "unit": "day"},
            {"source_key": "missing", "batch_quantity": 1, "batch_unit": "cum"},
            self.catalog,
        )

        self.assertIsNone(result)

    def test_keeps_final_coefficient_without_inventing_task_hours(self):
        sheet = source_sheet_with_verified_norms()
        sheet.append([
            "Other", "Soil", "Manual", "Final coefficient only", 1, "0114", None,
            "Published final norm only.", "Final coefficient: 1.25 Beldar-day per 10 cum.", "source-reference",
            1.25, "final-only", "Beldar", "day", 558, "W", "Other|Soil|Manual|Final coefficient only|1",
        ])
        result = derive_resource_norm(
            {"code": "0114", "description": "Beldar", "coefficient": 1.25, "unit": "day"},
            {"source_key": "final-only", "batch_quantity": 10, "batch_unit": "cum"},
            build_derivation_catalog(sheet),
        )

        self.assertIsNone(result["task_hours"])
        self.assertEqual(result["interpretation_label"], "Teaching interpretation, not a published CPWD rule")

    def test_matches_excel_numeric_resource_code_to_zero_padded_support_code(self):
        result = derive_resource_norm(
            {"code": "0114", "description": "Beldar", "coefficient": 6.8, "unit": "day"},
            {"source_key": "2. Surface Preparation & Earthwork|General surface cut|Manual labor (depth ≤1.5m)|1", "batch_quantity": 100, "batch_unit": "sqm"},
            self.catalog,
        )

        self.assertIsNotNone(result)


class FormulaCompatibilityTests(unittest.TestCase):
    def test_exported_sheet8_formula_nodes_are_ascii_only(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "sanitized.xlsx"
            export_ascii_formula_workbook(BASELINE_WORKBOOK, output)
            formulas = sheet8_formulas(output)

        self.assertEqual(len(formulas), 473)
        invalid = [formula for formula in formulas if any(ord(char) > 127 for char in formula)]
        self.assertEqual(invalid, [])

    def test_targeted_export_rewrites_only_formula_literals_in_sheet8(self):
        source_formulas = sheet8_formulas(BASELINE_WORKBOOK)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "sanitized.xlsx"
            export_ascii_formula_workbook(BASELINE_WORKBOOK, output)
            output_formulas = sheet8_formulas(output)

        self.assertEqual(len(output_formulas), 473)
        self.assertTrue(all(all(ord(char) <= 127 for char in formula) for formula in output_formulas))
        self.assertEqual(
            output_formulas,
            [formula.replace("—", "-").replace("→", "->").replace("›", ">") for formula in source_formulas],
        )


class FirstPrinciplesRenderingTests(unittest.TestCase):
    def _workbook(self):
        workbook = Workbook()
        support = workbook.active
        support.title = "02_support_earth_work"
        earth = workbook.create_sheet("02_Earth_Work")
        support["A3"] = "Item 2.1.1 | Surface dressing of ground"
        support["A4"] = "Details of cost for 100 sqm"
        support["A7"], support["B7"], support["C7"], support["D7"] = "0114", "Beldar", "day", 6.8
        support["A8"], support["B8"], support["C8"], support["D8"] = "0115", "Coolie", "day", 5.6
        support["A9"], support["B9"], support["C9"], support["D9"] = "0003", "Diesel Road Roller", "day", 0.008
        support["A10"], support["B10"], support["C10"], support["D10"] = "0999", "Published final-only resource", "day", 1.25
        support["F7"] = "=D7*1"
        earth.append([None] * 17)
        earth.append([
            "Surface Excavation", "All kinds of soil", "Manual", "General surface cut", 1, "0114", 0,
            "Loosen and trim soil.", "Gang: 4 Beldars × 13.60 h = 54.40 man-hrs ÷ 8.00 h = 6.800 Beldar-days per 100 sqm.",
            "source", 6.8, "support-2.1.1", "Beldar", "day", 0, "W", "key",
        ])
        earth.append([
            "Embankment", "Soil", "Rolling", "Roll", 3, "0003", 0,
            "Compact with roller.", "Plant: 1 Roller × 0.064 h ÷ 8.00 h = 0.008 roller-day per 100 sqm.",
            "source", 0.008, "support-2.1.1", "Diesel Road Roller", "day", 0, "W", "key",
        ])
        earth.append([
            "Surface Excavation", "All kinds of soil", "Manual", "General surface cut", 2, "0115", 0,
            "Carry spoil.", "Gang: 4 Coolies × 11.20 h = 44.80 man-hrs ÷ 8.00 h = 5.600 Coolie-days per 100 sqm.",
            "source", 5.6, "support-2.1.1", "Coolie", "day", 0, "W", "key",
        ])
        return workbook

    def test_renders_visible_shift_derivation_and_learning_note(self):
        workbook = self._workbook()
        apply_first_principles_learning(workbook, {"2.1.1": "support-2.1.1"})
        support = workbook["02_support_earth_work"]

        self.assertIn("CPWD fixed norm / source evidence", support["I7"].value)
        self.assertIn("4 × 13.6 task-hours ÷ 8-hour shift = 6.8 day per 100 sqm", support["I7"].value)
        self.assertIn("The inverse gives 14.705882352941 sqm per day", support["I7"].value)
        self.assertIn("8-hour shift", support["I8"].value)
        self.assertIn("1 × 0.064 task-hours ÷ 8-hour shift = 0.008 day per 100 sqm", support["I9"].value)
        self.assertIn("Engineering teaching reconstruction, not a published CPWD rule", support["I10"].value)
        self.assertIn("CPWD fixed norm: Beldar.", support["D7"].comment.text)
        self.assertIn("Calculation", support["D7"].comment.text)
        self.assertNotIn("Boundary conditions", support["D7"].comment.text)
        self.assertEqual(support["D7"].value, 6.8)
        self.assertEqual(support["F7"].value, "=D7*1")

    def test_all_338_resource_rows_have_labelled_shift_reconstructions(self):
        workbook = load_workbook(BASELINE_WORKBOOK)
        apply_first_principles_learning(workbook)
        support = workbook["02_support_earth_work"]
        resource_rows = [
            row for row in range(1, support.max_row + 1)
            if support.cell(row, 1).value not in (None, "")
            and isinstance(support.cell(row, 2).value, str)
            and isinstance(support.cell(row, 4).value, (int, float))
        ]
        self.assertEqual(len(resource_rows), 338)
        visible = [support.cell(row, 9).value for row in resource_rows]
        self.assertTrue(all("÷ 8-hour shift =" in text for text in visible))
        reconstructed = [text for text in visible if "Engineering teaching reconstruction" in text]
        self.assertEqual(len(reconstructed), 320)
        self.assertTrue(all("Engineering teaching reconstruction, not a published CPWD rule" in text for text in reconstructed))
        self.assertTrue(all("Engineering teaching reconstruction, not a published CPWD rule" in support.cell(row, 4).comment.text
                            for row in resource_rows if "Engineering teaching reconstruction" in support.cell(row, 9).value))

    def test_groups_only_repeated_markup_rows_on_first_middle_and_final_items(self):
        workbook = load_workbook(BASELINE_WORKBOOK)
        support = workbook["02_support_earth_work"]
        apply_markup_outline(support)
        item_starts = [r for r in range(1, support.max_row + 1) if str(support.cell(r, 1).value or "").startswith("Item ")]
        for start in (item_starts[0], item_starts[len(item_starts)//2], item_starts[-1]):
            end = next(r for r in range(start, support.max_row + 1)
                       if str(support.cell(r, 1).value or "").startswith("Total cost for"))
            water = next(r for r in range(start, end + 1)
                         if str(support.cell(r, 1).value or "").startswith("Add: Water charges"))
            self.assertEqual(support.row_dimensions[water].outlineLevel, 1)
            self.assertEqual(support.row_dimensions[end].outlineLevel, 1)
            self.assertEqual(support.row_dimensions[end + 1].outlineLevel, 0)
            self.assertEqual(support.row_dimensions[end + 2].outlineLevel, 0)

    def test_places_detailed_learning_card_in_i_and_only_concise_derivation_in_d_note(self):
        workbook = load_workbook(BASELINE_WORKBOOK)
        apply_first_principles_learning(workbook)
        support = workbook["02_support_earth_work"]

        self.assertIn("Engineering teaching reconstruction, not a published CPWD rule", support["I26"].value)
        self.assertIn("4-person labour gang", support["I26"].value)
        self.assertIn("Boundary conditions", support["I26"].value)
        self.assertIn("When the norm changes", support["I26"].value)
        self.assertIn("÷ 8-hour shift = 8.6 day per 100 sqm", support["D26"].comment.text)
        self.assertNotIn("Boundary conditions", support["D26"].comment.text)

    def test_sizes_resource_rows_for_their_actual_wrapped_i_card_only(self):
        workbook = load_workbook(BASELINE_WORKBOOK)
        support = workbook["02_support_earth_work"]
        original_non_resource_height = support.row_dimensions[10].height
        apply_first_principles_learning(workbook)

        longest_row = max(
            (row for row in range(1, support.max_row + 1) if support.cell(row, 4).comment),
            key=lambda row: estimated_card_lines(support.cell(row, 9).value, support.column_dimensions["I"].width),
        )
        for row in (26, longest_row):
            lines = estimated_card_lines(support.cell(row, 9).value, support.column_dimensions["I"].width)
            self.assertGreaterEqual(support.row_dimensions[row].height, lines * 15 + 8)
        self.assertEqual(support.row_dimensions[10].height, original_non_resource_height)


if __name__ == "__main__":
    unittest.main()
