"""Acceptance tests for difficult-condition additions in the earthwork composer."""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from earthwork_composer_catalogue import (  # noqa: E402
    EARTHWORK_COMPOSER_CATALOGUE,
    apply_difficult_condition_reference,
    calculate_difficult_condition_extra,
)
from support_earth_learning_edit import apply_first_principles_learning  # noqa: E402


class DifficultConditionCatalogueTests(unittest.TestCase):
    def test_water_extra_is_a_qualified_percent_addition_with_depth_measurement(self):
        item = EARTHWORK_COMPOSER_CATALOGUE["2.24.1"]

        self.assertEqual(item["relationship_type"], "conditional_extra")
        self.assertEqual(item["base_scope"], "each applicable earthwork item")
        self.assertEqual(item["percent"], 20)
        self.assertEqual(item["condition"], "in or under water and/or liquid mud, including pumping out water as required")
        self.assertEqual(item["measurement_basis"], "metre depth from sub-soil water level to the centre of gravity of qualifying work")

    def test_foul_position_extra_is_a_qualified_percent_addition_with_depth_measurement(self):
        item = EARTHWORK_COMPOSER_CATALOGUE["2.24.2"]

        self.assertEqual(item["relationship_type"], "conditional_extra")
        self.assertEqual(item["base_scope"], "each applicable earthwork item")
        self.assertEqual(item["percent"], 25)
        self.assertEqual(item["condition"], "in or under foul position, including pumping out water as required")
        self.assertEqual(item["measurement_basis"], "metre depth from sub-soil water level to the centre of gravity of qualifying work")

    def test_difficult_extra_is_a_percentage_of_the_selected_qualifying_base_rate(self):
        # Case 2: Item 2.2 selected as the base rate for qualifying excavation.
        water = calculate_difficult_condition_extra("2.24.1", selected_base_rate=746.80)
        foul = calculate_difficult_condition_extra("2.24.2", selected_base_rate=746.80)

        self.assertEqual(water["selected_base_rate"], 746.80)
        self.assertEqual(water["percent"], 20)
        self.assertAlmostEqual(water["extra_rate"], 149.36)
        self.assertEqual(foul["percent"], 25)
        self.assertAlmostEqual(foul["extra_rate"], 186.70)
        self.assertNotEqual(water["extra_rate"], 20)
        self.assertNotEqual(foul["extra_rate"], 25)

    def test_reference_sheet_shows_correct_scope_percent_and_qualifying_measurement(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "02_support_earth_work"

        apply_difficult_condition_reference(sheet)

        values = [str(row[0] or "") for row in sheet.iter_rows(min_col=1, max_col=1, values_only=True)]
        rendered = "\n".join(values)
        self.assertIn("2.24.1", rendered)
        self.assertIn("20% extra over each applicable earthwork item", rendered)
        self.assertIn("water and/or liquid mud", rendered)
        self.assertIn("2.24.2", rendered)
        self.assertIn("25% extra over each applicable earthwork item", rendered)
        self.assertIn("foul position", rendered)
        self.assertIn("sub-soil water level to the centre of gravity", rendered)
        self.assertNotIn("timbering items 2.16 to 2.23", rendered.lower())

    def test_production_learning_export_corrects_the_actual_224_rows_without_rebuilding_schedule(self):
        """The export path must correct the shipped support sheet, not just a blank sheet."""
        workbook = load_workbook(ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx")
        support = workbook["02_support_earth_work"]
        original_2241_style = support["A1194"].style_id
        original_2242_style = support["A1200"].style_id

        apply_first_principles_learning(workbook)
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "composer-output.xlsx"
            workbook.save(output)
            exported = load_workbook(output)
            try:
                support = exported["02_support_earth_work"]
                self.assertIn("20% extra over each applicable earthwork item", support["A1194"].value)
                self.assertIn("water and/or liquid mud", support["A1194"].value)
                self.assertIn("25% extra over each applicable earthwork item", support["A1200"].value)
                self.assertIn("foul position", support["A1200"].value)
                water_block = "\n".join(str(support.cell(row, col).value or "")
                                       for row in range(1194, 1200) for col in range(1, 7))
                foul_block = "\n".join(str(support.cell(row, col).value or "")
                                      for row in range(1200, 1206) for col in range(1, 7))
                self.assertNotIn("timbering", water_block.lower())
                self.assertNotIn("timbering", foul_block.lower())
                self.assertIn("20% of the qualifying selected base rate", water_block)
                self.assertIn("25% of the qualifying selected base rate", foul_block)
                self.assertIn("sub-soil water level to the centre of gravity", water_block)
                self.assertIn("sub-soil water level to the centre of gravity", foul_block)
                self.assertIn("not a flat Rs rate", water_block)
                self.assertIn("not a flat Rs rate", foul_block)
                self.assertEqual(support["A1194"].style_id, original_2241_style)
                self.assertEqual(support["A1200"].style_id, original_2242_style)
            finally:
                exported.close()


if __name__ == "__main__":
    unittest.main()
