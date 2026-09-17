"""Tests for Sub-Head 12 Roofing First-Principles Sheet."""

import unittest
import openpyxl
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"

class TestSupportRoofing(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
        cls.ws = cls.wb["12_Roofing"]

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_sheet_exists(self):
        self.assertIn("12_Roofing", self.wb.sheetnames)

    def test_title_and_columns(self):
        title = self.ws.cell(1, 1).value
        self.assertIn("ROOFING", title)
        headers = [self.ws.cell(6, col).value for col in range(1, 8)]
        expected = [
            "Item Code",
            "Labour / Machine / Material",
            "Work done",
            "Condition / When used",
            "Category",
            "Productivity",
            "Quantity"
        ]
        self.assertEqual(headers, expected)

    def test_core_items_present(self):
        found_items = set()
        for r in range(7, self.ws.max_row + 1):
            val = self.ws.cell(r, 1).value
            if val and str(val).strip().startswith("12."):
                found_items.add(str(val).strip())

        key_items = {
            "12.1.1", "12.4.1", "12.7.1", "12.16",
            "12.21", "12.22", "12.41.2", "12.50", "12.51.1"
        }
        for ki in key_items:
            self.assertIn(ki, found_items, f"Missing key item {ki} in 12_Roofing")

    def test_no_four_digit_prefixes_in_column_b(self):
        code_prefixed = []
        for r in range(7, self.ws.max_row + 1):
            val = self.ws.cell(r, 2).value
            if val and re.match(r"^\d{4}\s+", str(val)):
                code_prefixed.append((r, val))
        self.assertEqual(code_prefixed, [], f"Found 4-digit code prefixes in Col B: {code_prefixed}")

    def test_direct_materials_and_references_present(self):
        references = []
        materials = []
        for r in range(7, self.ws.max_row + 1):
            cat = self.ws.cell(r, 5).value
            item = self.ws.cell(r, 2).value
            if cat == "Reference":
                references.append(item)
            elif cat == "Material":
                materials.append(item)

        self.assertTrue(len(references) > 0, "Expected referenced base items (e.g. REF#4.1.3, REF#3.8)")
        self.assertTrue(len(materials) > 0, "Expected direct materials (CGS sheets, PPGI sheets, cement, etc.)")

if __name__ == "__main__":
    unittest.main()
