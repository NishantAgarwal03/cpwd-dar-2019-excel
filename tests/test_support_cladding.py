"""Tests for Sub-Head 08 Cladding Work First-Principles Sheet."""

import unittest
import openpyxl
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"

class TestSupportCladding(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
        cls.ws = cls.wb["08_Cladding_Work"]

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_sheet_exists(self):
        self.assertIn("08_Cladding_Work", self.wb.sheetnames)

    def test_title_and_columns(self):
        title = self.ws.cell(1, 1).value
        self.assertIn("CLADDING WORK", title)
        headers = [self.ws.cell(7, col).value for col in range(1, 8)]
        expected = [
            "Item Code",
            "Labour / Machine / Material",
            "Work done",
            "Condition / When used",
            "Category",
            "Productivity",
            "Quantity"
        ]
        self.assertEqual(headers, expected)

    def test_core_items_present(self):
        found_items = set()
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 1).value
            if val and str(val).strip().startswith("8."):
                found_items.add(str(val).strip())

        key_items = {
            "8.1.1.1", "8.1.1.2", "8.2.2.1", "8.2.2.2",
            "8.3.1", "8.3.2", "8.6", "8.7.2"
        }
        for ki in key_items:
            self.assertIn(ki, found_items, f"Missing key item {ki} in 08_Cladding_Work")

    def test_no_four_digit_prefixes_in_column_b(self):
        code_prefixed = []
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 2).value
            if val and re.match(r"^\d{4}\s+", str(val)):
                code_prefixed.append((r, val))
        self.assertEqual(code_prefixed, [], f"Found 4-digit code prefixes in Col B: {code_prefixed}")

    def test_direct_materials_and_references_present(self):
        references = []
        materials = []
        for r in range(8, self.ws.max_row + 1):
            cat = self.ws.cell(r, 5).value
            item = self.ws.cell(r, 2).value
            if cat == "Reference":
                references.append(item)
            elif cat == "Material":
                materials.append(item)

        self.assertTrue(len(references) > 0, "Expected referenced base items (e.g. REF#3.8)")
        self.assertTrue(len(materials) > 0, "Expected direct materials (marble, granite, cramps, cement)")

if __name__ == "__main__":
    unittest.main()
