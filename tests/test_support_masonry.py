"""Tests for Sub-Head 06 Masonry Work First-Principles Sheet."""

import unittest
import openpyxl
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"

class TestSupportMasonry(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
        cls.ws = cls.wb["06_Masonry_Work"]

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_sheet_exists(self):
        self.assertIn("06_Masonry_Work", self.wb.sheetnames)

    def test_title_and_columns(self):
        title = self.ws.cell(1, 1).value
        self.assertIn("MASONRY WORK", title)
        headers = [self.ws.cell(7, col).value for col in range(1, 8)]
        expected = [
            "Item Code",
            "Labour / Machine / Material",
            "Work done",
            "Condition / When used",
            "Category",
            "Productivity",
            "Quantity"
        ]
        self.assertEqual(headers, expected)

    def test_core_items_present(self):
        found_items = set()
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 1).value
            if val and str(val).strip().startswith("6."):
                found_items.add(str(val).strip())

        key_items = {
            "6.1.1", "6.1.2", "6.2.1", "6.2.2", "6.4.1", "6.4.2", "6.6",
            "6.12.1", "6.12.2", "6.13.1", "6.13.2", "6.15", "6.23",
            "6.38", "6.44", "6.47"
        }
        for ki in key_items:
            self.assertIn(ki, found_items, f"Missing key item {ki} in 06_Masonry_Work")

    def test_no_four_digit_prefixes_in_column_b(self):
        code_prefixed = []
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 2).value
            if val and re.match(r"^\d{4}\s+", str(val)):
                code_prefixed.append((r, val))
        self.assertEqual(code_prefixed, [], f"Found 4-digit code prefixes in Col B: {code_prefixed}")

    def test_mortar_decomposition_and_categories(self):
        item_categories = {}
        for r in range(8, self.ws.max_row + 1):
            code = self.ws.cell(r, 1).value
            cat = self.ws.cell(r, 5).value
            if code and str(code).strip().startswith("6."):
                item_id = str(code).strip()
                if item_id not in item_categories:
                    item_categories[item_id] = set()
                if cat:
                    item_categories[item_id].add(str(cat).strip())

        # Check standard brickwork items have Material, Labour, and Reference
        standard_items = ["6.1.1", "6.1.2", "6.4.1", "6.12.1", "6.38"]
        for si in standard_items:
            cats = item_categories.get(si, set())
            self.assertIn("Material", cats, f"Item {si} missing Material")
            self.assertIn("Labour", cats, f"Item {si} missing Labour")
            self.assertIn("Reference", cats, f"Item {si} missing Reference")

if __name__ == "__main__":
    unittest.main()
