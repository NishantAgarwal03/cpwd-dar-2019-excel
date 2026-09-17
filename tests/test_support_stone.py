"""Tests for Sub-Head 07 Stone Work First-Principles Sheet."""

import unittest
import openpyxl
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"

class TestSupportStone(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
        cls.ws = cls.wb["07_Stone_Work"]

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_sheet_exists(self):
        self.assertIn("07_Stone_Work", self.wb.sheetnames)

    def test_title_and_columns(self):
        title = self.ws.cell(1, 1).value
        self.assertIn("STONE WORK", title)
        headers = [self.ws.cell(7, col).value for col in range(1, 8)]
        expected = [
            "Item Code",
            "Labour / Machine / Material",
            "Work done",
            "Condition / When used",
            "Category",
            "Productivity",
            "Quantity"
        ]
        self.assertEqual(headers, expected)

    def test_core_items_present(self):
        found_items = set()
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 1).value
            if val and str(val).strip().startswith("7."):
                found_items.add(str(val).strip())

        key_items = {
            "7.1.1", "7.2.1", "7.6.1", "7.7.1", "7.8.1",
            "7.12.1.1", "7.12.1.2"
        }
        for ki in key_items:
            self.assertIn(ki, found_items, f"Missing key item {ki} in 07_Stone_Work")

    def test_no_four_digit_prefixes_in_column_b(self):
        code_prefixed = []
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 2).value
            if val and re.match(r"^\d{4}\s+", str(val)):
                code_prefixed.append((r, val))
        self.assertEqual(code_prefixed, [], f"Found 4-digit code prefixes in Col B: {code_prefixed}")

    def test_item_categories_and_mortar_reference(self):
        item_categories = {}
        for r in range(8, self.ws.max_row + 1):
            code = self.ws.cell(r, 1).value
            cat = self.ws.cell(r, 5).value
            if code and str(code).strip().startswith("7."):
                item_id = str(code).strip()
                if item_id not in item_categories:
                    item_categories[item_id] = set()
                if cat:
                    item_categories[item_id].add(str(cat).strip())

        for item_id in ["7.1.1", "7.2.1", "7.6.1", "7.7.1", "7.8.1"]:
            cats = item_categories.get(item_id, set())
            self.assertIn("Material", cats, f"Item {item_id} missing Material")
            self.assertIn("Labour", cats, f"Item {item_id} missing Labour")
            self.assertIn("Reference", cats, f"Item {item_id} missing Reference")

if __name__ == "__main__":
    unittest.main()
