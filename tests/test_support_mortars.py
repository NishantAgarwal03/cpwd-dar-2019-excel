"""Tests for Sub-Head 03 Mortars First-Principles Sheet."""

import unittest
import openpyxl
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT / "CPWD_DAR_2019_Custom_Rate_Analysis_Workbook_Vol_1_Latest.xlsx"

class TestSupportMortars(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
        cls.ws = cls.wb["03_Mortars"]

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()

    def test_sheet_exists(self):
        self.assertIn("03_Mortars", self.wb.sheetnames)

    def test_title_and_columns(self):
        title = self.ws.cell(1, 1).value
        self.assertIn("MORTARS", title)
        headers = [self.ws.cell(7, col).value for col in range(1, 8)]
        expected = [
            "Item Code",
            "Labour / Machine / Material",
            "Work done",
            "Condition / When used",
            "Category",
            "Productivity",
            "Quantity"
        ]
        self.assertEqual(headers, expected)

    def test_all_18_items_present(self):
        found_items = set()
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 1).value
            if val and re.match(r"^3\.\d+$", str(val).strip()):
                found_items.add(str(val).strip())

        expected_items = {f"3.{i}" for i in range(1, 19)}
        self.assertEqual(found_items, expected_items)

    def test_no_four_digit_prefixes_in_column_b(self):
        code_prefixed = []
        for r in range(8, self.ws.max_row + 1):
            val = self.ws.cell(r, 2).value
            if val and re.match(r"^\d{4}\s+", str(val)):
                code_prefixed.append((r, val))
        self.assertEqual(code_prefixed, [], f"Found 4-digit code prefixes in Col B: {code_prefixed}")

    def test_item_categories_and_batch_qty(self):
        current_item = None
        item_categories = {}
        for r in range(8, self.ws.max_row + 1):
            code = self.ws.cell(r, 1).value
            cat = self.ws.cell(r, 5).value
            qty = self.ws.cell(r, 7).value
            if code and re.match(r"^3\.\d+$", str(code).strip()):
                item_id = str(code).strip()
                if item_id not in item_categories:
                    item_categories[item_id] = set()
                if cat:
                    item_categories[item_id].add(str(cat).strip())
                self.assertEqual(qty, "1 cum", f"Row {r} item {item_id} has non-standard batch qty {qty}")

        for item_id, cats in item_categories.items():
            self.assertIn("Material", cats, f"Item {item_id} missing Material")
            self.assertIn("Labour", cats, f"Item {item_id} missing Labour")

if __name__ == "__main__":
    unittest.main()
